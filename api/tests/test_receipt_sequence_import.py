from __future__ import annotations

from datetime import date
from io import BytesIO
import uuid

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from main import app


ADMIN_EMAIL = "ai-tests-admin@example.com"
ADMIN_PASSWORD = "AiTestsAdmin2026!"


def _login(client: TestClient) -> None:
    response = client.post(
        "/auth/login",
        json={"correo": ADMIN_EMAIL, "password": ADMIN_PASSWORD},
    )
    assert response.status_code == 200, response.text


def _receipt_payload(cliente_id: int) -> dict:
    return {
        "fecha": date.today().isoformat(),
        "cliente_id": cliente_id,
        "cajuelas": 20,
        "cuartillos": 0,
        "precio_fanega": 100000,
        "estado": "recibido",
    }


def test_receipt_sequence_is_configurable_monotonic_and_unique():
    suffix = uuid.uuid4().hex[:8]
    prefix = f"T{suffix[:4].upper()}-"

    with TestClient(app) as client:
        _login(client)
        created_client = client.post(
            "/clientes",
            json={
                "codigo": f"CL-SEQ-{suffix}",
                "nombre_completo": "Productor de consecutivos",
                "tipo_persona": "personal",
                "categoria": "agro",
                "numero_identificacion": f"SEQ-{suffix}",
                "activo": True,
            },
        )
        assert created_client.status_code == 201, created_client.text
        client_id = created_client.json()["id"]

        configured = client.patch(
            "/recibos/consecutivo",
            json={"prefijo": prefix, "siguiente_numero": 7000, "ancho": 6},
        )
        assert configured.status_code == 200, configured.text
        assert configured.json()["proximo_recibo"] == f"{prefix}007000"

        first = client.post("/recibos", json=_receipt_payload(client_id))
        second = client.post("/recibos", json=_receipt_payload(client_id))
        assert first.status_code == 201, first.text
        assert second.status_code == 201, second.text
        assert first.json()["numero_recibo"] == f"{prefix}007000"
        assert second.json()["numero_recibo"] == f"{prefix}007001"

        backwards = client.patch(
            "/recibos/consecutivo",
            json={"prefijo": prefix, "siguiente_numero": 7001, "ancho": 6},
        )
        assert backwards.status_code == 422

        explicit_number = f"{prefix}009000"
        explicit = client.post(
            "/recibos",
            json={**_receipt_payload(client_id), "numero_recibo": explicit_number},
        )
        assert explicit.status_code == 201, explicit.text
        duplicate = client.post(
            "/recibos",
            json={**_receipt_payload(client_id), "numero_recibo": explicit_number},
        )
        assert duplicate.status_code in {409, 422}

        after_explicit = client.post("/recibos", json=_receipt_payload(client_id))
        assert after_explicit.status_code == 201, after_explicit.text
        assert after_explicit.json()["numero_recibo"] == f"{prefix}009001"


def test_historical_receipt_import_is_transactional_and_advances_sequence():
    suffix = uuid.uuid4().hex[:8]
    prefix = f"I{suffix[:4].upper()}-"
    identification = f"HIST-{suffix}"
    historical_number = f"{prefix}000250"

    with TestClient(app) as client:
        _login(client)
        configured = client.patch(
            "/recibos/consecutivo",
            json={"prefijo": prefix, "siguiente_numero": 100, "ancho": 6},
        )
        assert configured.status_code == 200, configured.text

        template = client.get("/recibos/importacion/plantilla")
        assert template.status_code == 200, template.text
        assert template.content.startswith(b"PK")

        workbook = load_workbook(BytesIO(template.content))
        sheet = workbook["Recibos"]
        headers = [cell.value for cell in sheet[1]]
        row = {header: None for header in headers}
        row.update({
            "numero_recibo": historical_number,
            "fecha": date(2022, 12, 5),
            "cliente_identificacion": identification,
            "cliente_nombre": "Productor histórico",
            "cliente_telefono": "88887777",
            "finca_codigo": f"FIN-H-{suffix}",
            "finca_nombre": "Finca histórica",
            "cosecha": "2022-2023",
            "cajuelas": 25,
            "cuartillos": 2,
            "precio_fanega": 95000,
            "estado": "recibido",
            "observaciones": "Migrado desde documento físico",
        })
        sheet.append([row[header] for header in headers])
        buffer = BytesIO()
        workbook.save(buffer)

        imported = client.post(
            "/recibos/importacion",
            files={
                "file": (
                    "recibos-historicos.xlsx",
                    buffer.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert imported.status_code == 200, imported.text
        assert imported.json()["recibos_creados"] == 1
        assert imported.json()["clientes_creados"] == 1
        assert imported.json()["fincas_creadas"] == 1

        sequence = client.get("/recibos/consecutivo")
        assert sequence.status_code == 200
        assert sequence.json()["siguiente_numero"] == 251

        clients = client.get("/clientes").json()
        imported_client = next(row for row in clients if row["numero_identificacion"] == identification)
        automatic = client.post("/recibos", json=_receipt_payload(imported_client["id"]))
        assert automatic.status_code == 201, automatic.text
        assert automatic.json()["numero_recibo"] == f"{prefix}000251"

        # Una fila nueva junto a otra duplicada debe hacer rollback completo.
        duplicate_book = load_workbook(BytesIO(template.content))
        duplicate_sheet = duplicate_book["Recibos"]
        duplicate_headers = [cell.value for cell in duplicate_sheet[1]]
        unique_number = f"{prefix}000400"
        for receipt_number in (unique_number, historical_number):
            duplicate_row = {header: None for header in duplicate_headers}
            duplicate_row.update({
                "numero_recibo": receipt_number,
                "fecha": date(2021, 11, 1),
                "cliente_identificacion": identification,
                "cliente_nombre": "Productor histórico",
                "cajuelas": 10,
            })
            duplicate_sheet.append([duplicate_row[header] for header in duplicate_headers])
        duplicate_buffer = BytesIO()
        duplicate_book.save(duplicate_buffer)

        rejected = client.post(
            "/recibos/importacion",
            files={"file": ("duplicados.xlsx", duplicate_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert rejected.status_code == 422
        assert rejected.json()["detail"]["errors"]
        assert all(row["numero_recibo"] != unique_number for row in client.get("/recibos").json())
