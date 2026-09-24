from __future__ import annotations

from io import BytesIO
from decimal import Decimal
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func
from sqlalchemy.orm import Session

from modules.mod_floresvolcan.models import (
    CropCatalogValue, FlowerCrop, GoodsReceipt, GoodsReceiptLine, HarvestRecord,
    InventoryLot, PlantingOrder, PlantingRecord, ProductionBatch, ProductionSpace,
    PurchaseOrder, RequisitionDelivery, RequisitionLine, SapProduct, Supplier,
    WorkCrewMember,
)

HEADERS = {
    "Entradas": ["COD","Variedad","Calibre","Comun","Fecha de Ingreso","Contenedor","Lote","Cosecha","Proveedor","Cajas","Cantidad por Caja","Total Bulbos"],
    "Salida Bulbo": ["COD","Variedad","Calibre","Comun","Contenedor","Lote","Proveedor","Stock","Cajas","Semana Salida","Fecha de Salida","Cantidad por Caja","Bulbos"],
    "Siembra": ["KEY","LOTE","INVERNADERO","LADO","CAMA","COMUN","PROVEEDOR","COLOR","VARIEDAD","CANTIDAD","CICLO","SEM. SIEMBRA","FECHA DE SIEMBRA","SEM. COSECHA","CALIBRE","DENSIDAD","FECHA SELLADO","SEMBRADOR","M2"],
    "Corta": ["KEY","LOTE","INV.","CAMA","COMUN","PROVEEDOR","COLOR","VARIEDAD","CANTIDAD","SEM. DE SIEMBRA","SEM. COSECHA PROG.","SEM. COSECHA REAL","FECHA DE CORTA","DES. CANTIDAD","DES. MOTIVO","1 BL","2 BL","3-5 BL","TOTAL VENTA","TOTAL "],
}

COMMAND_CATEGORIES = [
    ("INVERNADERO", "greenhouse"), ("COMÚN", "common"), ("PROVEEDOR", "supplier"),
    ("COLOR", "color"), ("VARIEDAD", "variety"), ("CALIBRE", "caliber"),
    ("DENSIDAD", "density"), ("LADO", "side"), ("SEMBRADOR", "legacy_planter"),
]


def _n(v) -> float:
    try:
        return float(v or 0)
    except Exception:
        return 0.0


def _legacy_code(variety, caliber, container, lot) -> str:
    return f"{variety or ''}{caliber or ''}{container or ''}{lot or ''}".replace(" ", "")


def _supplier_for_lot(db: Session, lot: InventoryLot) -> Supplier | None:
    line = db.query(GoodsReceiptLine).filter(GoodsReceiptLine.id == lot.goods_receipt_line_id).first()
    receipt = db.query(GoodsReceipt).filter(GoodsReceipt.id == line.goods_receipt_id).first() if line else None
    po = db.query(PurchaseOrder).filter(PurchaseOrder.id == receipt.purchase_order_id).first() if receipt else None
    return db.query(Supplier).filter(Supplier.id == po.supplier_id).first() if po else None


def _product(db: Session, lot: InventoryLot | None, product_id: int | None = None) -> SapProduct | None:
    pid = product_id or (lot.sap_product_id if lot else None)
    return db.query(SapProduct).filter(SapProduct.id == pid).first() if pid else None


def _space(db: Session, sid: int | None) -> ProductionSpace | None:
    return db.query(ProductionSpace).filter(ProductionSpace.id == sid).first() if sid else None


def _crew_label(db: Session, crew_id: int | None) -> str:
    if not crew_id:
        return ""
    members = db.query(WorkCrewMember).filter(WorkCrewMember.crew_id == crew_id).all()
    names = [m.worker.name for m in members if getattr(m, "worker", None)]
    return "-".join(names)


def _style_sheet(ws) -> None:
    fill = PatternFill("solid", fgColor="107C5C")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        width = min(35, max(10, max(len(str(c.value or "")) for c in list(col)[:120]) + 2))
        ws.column_dimensions[col[0].column_letter].width = width


def build_lirio_workbook(db: Session, title: str = "Lirio") -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    # Entradas
    ws = wb.create_sheet("Entradas"); ws.append(HEADERS["Entradas"])
    lots = db.query(InventoryLot).order_by(InventoryLot.received_at, InventoryLot.id).all()
    for lot in lots:
        product = _product(db, lot); supplier = _supplier_for_lot(db, lot)
        variety = lot.variety_snapshot or (product.variety if product else "")
        caliber = lot.caliber_snapshot or (product.caliber if product else "")
        common = lot.common_code_snapshot or (product.common_code if product else "")
        ws.append([
            _legacy_code(variety, caliber, lot.container_number, lot.supplier_lot_code), variety, caliber, common,
            lot.received_at.date() if lot.received_at else None, lot.container_number, lot.supplier_lot_code, lot.harvest_year,
            supplier.code if supplier else "", _n(lot.boxes_received), _n(lot.units_per_box), _n(lot.quantity_received),
        ])
    _style_sheet(ws)

    # Salida Bulbo = entregas parciales dentro de requisiciones.
    ws = wb.create_sheet("Salida Bulbo"); ws.append(HEADERS["Salida Bulbo"])
    deliveries = db.query(RequisitionDelivery).order_by(RequisitionDelivery.delivered_at, RequisitionDelivery.id).all()
    for d in deliveries:
        line = db.query(RequisitionLine).filter(RequisitionLine.id == d.requisition_line_id).first()
        lot = db.query(InventoryLot).filter(InventoryLot.id == d.inventory_lot_id).first()
        product = _product(db, lot, line.sap_product_id if line else None); supplier = _supplier_for_lot(db, lot) if lot else None
        variety = (lot.variety_snapshot if lot else None) or (product.variety if product else "")
        caliber = (lot.caliber_snapshot if lot else None) or (product.caliber if product else "")
        common = (lot.common_code_snapshot if lot else None) or (product.common_code if product else "")
        units_per_box = _n(d.units_per_box or (lot.units_per_box if lot else None))
        boxes = _n(d.boxes) or (_n(d.quantity) / units_per_box if units_per_box else 0)
        ws.append([
            _legacy_code(variety, caliber, lot.container_number if lot else "", lot.supplier_lot_code if lot else ""),
            variety, caliber, common, lot.container_number if lot else "", lot.supplier_lot_code if lot else "",
            supplier.code if supplier else "", _n(lot.boxes_received if lot else 0), boxes, d.delivery_week,
            d.delivered_at.date() if d.delivered_at else None, units_per_box, _n(d.quantity),
        ])
    _style_sheet(ws)

    # Siembra
    ws = wb.create_sheet("Siembra"); ws.append(HEADERS["Siembra"])
    plantings = db.query(PlantingRecord).order_by(PlantingRecord.planted_at, PlantingRecord.id).all()
    for pr in plantings:
        lot = db.query(InventoryLot).filter(InventoryLot.id == pr.inventory_lot_id).first()
        line = db.query(RequisitionLine).filter(RequisitionLine.id == pr.requisition_line_id).first()
        product = _product(db, lot, line.sap_product_id if line else None); supplier = _supplier_for_lot(db, lot) if lot else None
        greenhouse = _space(db, pr.greenhouse_space_id); bed = _space(db, pr.bed_space_id)
        batch = db.query(ProductionBatch).filter(ProductionBatch.id == pr.production_batch_id).first()
        variety = (lot.variety_snapshot if lot else None) or (product.variety if product else "")
        caliber = (lot.caliber_snapshot if lot else None) or (product.caliber if product else "")
        common = (lot.common_code_snapshot if lot else None) or (product.common_code if product else "")
        color = (lot.color_snapshot if lot else None) or (product.color if product else "")
        inv = greenhouse.name if greenhouse else ""; bed_label = (bed.name.replace("Cama ", "") if bed else "")
        key = f"{inv}{lot.supplier_lot_code if lot else ''}{pr.planting_week or ''}{bed_label}"
        ws.append([
            key, lot.supplier_lot_code if lot else "", inv, pr.side, bed_label, common, supplier.code if supplier else "",
            color, variety, _n(pr.quantity_planted), pr.cycle_weeks, pr.planting_week,
            pr.planted_at.date() if pr.planted_at else None, pr.projected_harvest_week, caliber, _n(pr.density),
            pr.sealed_at, _crew_label(db, pr.crew_id), _n(pr.area_m2),
        ])
    _style_sheet(ws)

    # Corta
    ws = wb.create_sheet("Corta"); ws.append(HEADERS["Corta"])
    cuts = db.query(HarvestRecord).order_by(HarvestRecord.harvested_at, HarvestRecord.id).all()
    for cut in cuts:
        batch = db.query(ProductionBatch).filter(ProductionBatch.id == cut.production_batch_id).first()
        lot = db.query(InventoryLot).filter(InventoryLot.id == batch.inventory_lot_id).first() if batch and batch.inventory_lot_id else None
        product = _product(db, lot, batch.sap_product_id if batch else None); supplier = _supplier_for_lot(db, lot) if lot else None
        greenhouse = _space(db, batch.greenhouse_space_id if batch else None); bed = _space(db, batch.bed_space_id if batch else None)
        planted = db.query(func.coalesce(func.sum(PlantingRecord.quantity_planted), 0)).filter(PlantingRecord.production_batch_id == cut.production_batch_id).scalar()
        variety = (lot.variety_snapshot if lot else None) or (product.variety if product else "")
        common = (lot.common_code_snapshot if lot else None) or (product.common_code if product else "")
        color = (lot.color_snapshot if lot else None) or (product.color if product else "")
        inv = greenhouse.name if greenhouse else ""; bed_label = bed.name.replace("Cama ", "") if bed else ""
        key = f"{inv}{lot.supplier_lot_code if lot else ''}{batch.planting_week if batch else ''}{bed_label}"
        ws.append([
            key, lot.supplier_lot_code if lot else "", inv, bed_label, common, supplier.code if supplier else "", color,
            variety, _n(planted), batch.planting_week if batch else None, batch.projected_harvest_week if batch else None,
            cut.actual_harvest_week, cut.harvested_at.date() if cut.harvested_at else None, _n(cut.discard_qty),
            cut.discard_reason, _n(cut.bloom_1_qty), _n(cut.bloom_2_qty), _n(cut.bloom_3_5_qty),
            _n(cut.quantity_harvested), _n(cut.total_output),
        ])
    _style_sheet(ws)

    # COMANDOS familiar: columnas separadas por una columna vacía.
    ws = wb.create_sheet("COMANDOS")
    crop = db.query(FlowerCrop).filter(FlowerCrop.code == "LIRIO").first()
    columns = []
    for label, category in COMMAND_CATEGORIES:
        if category == "greenhouse":
            values = [x.name for x in db.query(ProductionSpace).filter(ProductionSpace.type == "greenhouse", ProductionSpace.status == "active").order_by(ProductionSpace.name).all()]
        elif category == "supplier":
            values = [x.code for x in db.query(Supplier).filter(Supplier.active.is_(True)).order_by(Supplier.code).all()]
        else:
            values = [x.label for x in db.query(CropCatalogValue).filter(CropCatalogValue.crop_id == crop.id, CropCatalogValue.category == category, CropCatalogValue.active.is_(True)).order_by(CropCatalogValue.sort_order, CropCatalogValue.label).all()] if crop else []
        columns.append((label, values))
    max_len = max([len(v) for _, v in columns] + [1])
    for i, (label, values) in enumerate(columns):
        col = 2 + i * 2
        ws.cell(1, col, label)
        for row_idx, value in enumerate(values, start=2):
            ws.cell(row_idx, col, value)
    for cell in ws[1]:
        if cell.value:
            cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="107C5C")
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = 22
    ws.freeze_panes = "A2"

    bio = BytesIO(); wb.save(bio); return bio.getvalue()


def build_excel_bundle(db: Session) -> bytes:
    data = build_lirio_workbook(db)
    bio = BytesIO()
    with ZipFile(bio, "w", ZIP_DEFLATED) as zf:
        zf.writestr("Lirio.xlsx", data)
        zf.writestr("Liriodoc.xlsx", data)
    return bio.getvalue()
