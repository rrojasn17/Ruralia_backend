from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import or_
from sqlalchemy.orm import Session

from modules.mod_floresvolcan.models import Worker

HEADERS = ["employee_code", "identification", "name", "hire_date", "nationality", "birth_date", "daily_salary", "active"]


def build_worker_template() -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = "Trabajadores"
    ws.append(HEADERS)
    # El machote no incluye una fila ficticia para evitar importarla accidentalmente.
    info = wb.create_sheet("INSTRUCCIONES")
    info.append(["Campo", "Uso"])
    info.append(["employee_code", "Código único del trabajador. Obligatorio."])
    info.append(["identification", "Identificación única. Obligatorio."])
    info.append(["name", "Nombre completo. Obligatorio."])
    info.append(["hire_date", "Fecha de ingreso: YYYY-MM-DD o DD/MM/YYYY. Obligatorio."])
    info.append(["nationality", "Nacionalidad. Opcional."])
    info.append(["birth_date", "Fecha de nacimiento. Opcional."])
    info.append(["daily_salary", "Salario diario numérico. Opcional."])
    info.append(["active", "TRUE/FALSE. Si queda vacío se considera activo."])
    fill = PatternFill("solid", fgColor="107C5C")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = fill; c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    widths = [16, 20, 32, 16, 20, 16, 16, 12]
    for i, width in enumerate(widths, 1): ws.column_dimensions[ws.cell(1, i).column_letter].width = width
    bio = BytesIO(); wb.save(bio); return bio.getvalue()


def _date(value, field: str):
    if value in (None, ""): return None
    if isinstance(value, datetime): return value.date()
    if isinstance(value, date): return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try: return datetime.strptime(text, fmt).date()
        except ValueError: pass
    raise ValueError(f"{field}: fecha inválida {text!r}")


def _bool(value) -> bool:
    if isinstance(value, bool): return value
    if value in (None, ""): return True
    return str(value).strip().lower() not in {"0", "false", "no", "inactivo", "inactive"}


def import_workers(db: Session, content: bytes) -> dict:
    wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [str(x.value or "").strip() for x in ws[1]]
    missing = [h for h in HEADERS[:4] if h not in headers]
    if missing: raise ValueError("Faltan columnas obligatorias: " + ", ".join(missing))
    idx = {h: headers.index(h) for h in headers if h}
    created = updated = 0; errors=[]
    for excel_row, cells in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(v not in (None, "") for v in cells): continue
        try:
            code = str(cells[idx["employee_code"]] or "").strip()
            identification = str(cells[idx["identification"]] or "").strip()
            name = str(cells[idx["name"]] or "").strip()
            if not code or not identification or not name: raise ValueError("employee_code, identification y name son obligatorios")
            hire_date = _date(cells[idx["hire_date"]], "hire_date")
            if not hire_date: raise ValueError("hire_date es obligatoria")
            birth = _date(cells[idx.get("birth_date", -1)] if "birth_date" in idx else None, "birth_date")
            salary = cells[idx.get("daily_salary", -1)] if "daily_salary" in idx else None
            salary = Decimal(str(salary)) if salary not in (None, "") else None
            row = db.query(Worker).filter(or_(Worker.employee_code == code, Worker.identification == identification)).first()
            is_new = row is None
            row = row or Worker(employee_code=code, identification=identification, name=name, hire_date=hire_date)
            row.employee_code=code; row.identification=identification; row.name=name; row.hire_date=hire_date
            row.nationality=str(cells[idx["nationality"]]).strip() if "nationality" in idx and cells[idx["nationality"]] not in (None, "") else None
            row.birth_date=birth; row.daily_salary=salary
            row.active=_bool(cells[idx["active"]] if "active" in idx else True)
            db.add(row)
            if is_new: created += 1
            else: updated += 1
        except (ValueError, InvalidOperation) as exc:
            errors.append({"row": excel_row, "error": str(exc)})
    db.commit()
    return {"created":created,"updated":updated,"errors":errors,"processed":created+updated+len(errors)}
