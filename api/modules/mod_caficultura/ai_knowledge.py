from __future__ import annotations

import json
import re
import unicodedata
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy.orm import Session

from config import AI_KNOWLEDGE_MAX_CHARS
from modules.mod_caficultura.model_ai_knowledge import AIKnowledgeChunk, AIKnowledgeDocument


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".csv", ".json", ".md", ".txt"}
STOP_WORDS = {
    "para",
    "como",
    "con",
    "del",
    "desde",
    "donde",
    "esta",
    "este",
    "estos",
    "estas",
    "finca",
    "sensor",
    "sensores",
    "sobre",
    "entre",
    "porque",
    "cual",
    "cuales",
    "que",
    "una",
    "uno",
    "unos",
    "unas",
    "por",
    "los",
    "las",
    "sus",
    "han",
    "hay",
    "muy",
}


class KnowledgeExtractionError(ValueError):
    pass


def _cap(text: str) -> str:
    return text[:AI_KNOWLEDGE_MAX_CHARS]


def _plain_text(path: Path) -> str:
    return _cap(path.read_text(encoding="utf-8", errors="replace"))


def _json_text(path: Path) -> str:
    try:
        payload = json.loads(path.read_text(encoding="utf-8", errors="strict"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise KnowledgeExtractionError("El archivo JSON no es válido") from exc
    return _cap(json.dumps(payload, ensure_ascii=False, indent=2))


def _docx_text(path: Path) -> str:
    try:
        with zipfile.ZipFile(path) as archive:
            info = archive.getinfo("word/document.xml")
            if info.file_size > min(AI_KNOWLEDGE_MAX_CHARS * 8, 30_000_000):
                raise KnowledgeExtractionError(
                    "El documento Word es demasiado grande para indexarlo"
                )
            root = ElementTree.fromstring(archive.read(info))
    except KnowledgeExtractionError:
        raise
    except (OSError, KeyError, zipfile.BadZipFile, ElementTree.ParseError) as exc:
        raise KnowledgeExtractionError("No se pudo leer el documento Word") from exc
    paragraphs: list[str] = []
    current: list[str] = []
    for node in root.iter():
        if node.tag.endswith("}t") and node.text:
            current.append(node.text)
        elif node.tag.endswith("}p") and current:
            paragraphs.append(" ".join(current))
            current = []
        if sum(len(item) for item in paragraphs) >= AI_KNOWLEDGE_MAX_CHARS:
            break
    if current:
        paragraphs.append(" ".join(current))
    return _cap("\n\n".join(paragraphs))


def _xlsx_text(path: Path) -> str:
    lines: list[str] = []
    total = 0
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets[:12]:
                heading = f"[Hoja: {sheet.title}]"
                lines.append(heading)
                total += len(heading)
                for row_index, row in enumerate(
                    sheet.iter_rows(values_only=True), start=1
                ):
                    if row_index > 10_000:
                        break
                    values = [
                        str(value).strip()[:1_000]
                        for value in row[:80]
                        if value not in (None, "")
                    ]
                    if values:
                        line = " | ".join(values)
                        lines.append(line)
                        total += len(line)
                    if total >= AI_KNOWLEDGE_MAX_CHARS:
                        return _cap("\n".join(lines))
        finally:
            workbook.close()
    except (OSError, ValueError, zipfile.BadZipFile) as exc:
        raise KnowledgeExtractionError("No se pudo leer el archivo Excel") from exc
    return _cap("\n".join(lines))


def _pdf_text(path: Path) -> str:
    parts: list[str] = []
    total = 0
    try:
        reader = PdfReader(str(path), strict=False)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise KnowledgeExtractionError(
                    "El PDF está protegido con contraseña"
                ) from exc
        for page_number, page in enumerate(reader.pages, start=1):
            text = str(page.extract_text() or "").strip()
            if text:
                section = f"[Página {page_number}]\n{text}"
                parts.append(section)
                total += len(section)
            if total >= AI_KNOWLEDGE_MAX_CHARS:
                break
    except KnowledgeExtractionError:
        raise
    except Exception as exc:
        raise KnowledgeExtractionError("No se pudo extraer texto del PDF") from exc
    return _cap("\n\n".join(parts))


def extract_document_text(path: Path, original_name: str) -> str:
    suffix = Path(original_name).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise KnowledgeExtractionError("Use PDF, DOCX, XLSX, CSV, JSON, Markdown o TXT")
    if suffix == ".pdf":
        text = _pdf_text(path)
    elif suffix == ".docx":
        text = _docx_text(path)
    elif suffix == ".xlsx":
        text = _xlsx_text(path)
    elif suffix == ".json":
        text = _json_text(path)
    else:
        text = _plain_text(path)
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 20:
        raise KnowledgeExtractionError(
            "No se encontró texto suficiente para indexar; si es un PDF escaneado, conviértalo con OCR"
        )
    return text


def chunk_document_text(
    text: str, max_chars: int = 1_400, overlap: int = 180
) -> list[str]:
    paragraphs = [item.strip() for item in re.split(r"\n\s*\n", text) if item.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        pieces = [
            paragraph[i : i + max_chars] for i in range(0, len(paragraph), max_chars)
        ] or [paragraph]
        for piece in pieces:
            candidate = f"{current}\n\n{piece}".strip() if current else piece
            if current and len(candidate) > max_chars:
                chunks.append(current[:max_chars])
                prefix = current[-overlap:] if overlap else ""
                current = f"{prefix}\n{piece}".strip()[:max_chars]
            else:
                current = candidate[:max_chars]
    if current:
        chunks.append(current)
    return chunks[:5_000]


def index_document(db: Session, document: AIKnowledgeDocument, text: str) -> int:
    document.chunks.clear()
    chunks = chunk_document_text(text)
    for position, content in enumerate(chunks):
        document.chunks.append(AIKnowledgeChunk(position=position, content=content))
    document.character_count = len(text)
    db.flush()
    return len(chunks)


def _normalized_tokens(value: str) -> list[str]:
    text = unicodedata.normalize("NFD", str(value or "").lower())
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return [
        token for token in re.findall(r"[a-z0-9]{3,}", text) if token not in STOP_WORDS
    ][:40]


def search_knowledge_base(
    db: Session, agent_id: int, query: str, limit: int = 6
) -> dict:
    terms = list(dict.fromkeys(_normalized_tokens(query)))
    rows = (
        db.query(AIKnowledgeChunk, AIKnowledgeDocument)
        .join(
            AIKnowledgeDocument, AIKnowledgeDocument.id == AIKnowledgeChunk.document_id
        )
        .filter(
            AIKnowledgeDocument.agent_id == agent_id,
            AIKnowledgeDocument.active.is_(True),
        )
        .order_by(AIKnowledgeDocument.id.asc(), AIKnowledgeChunk.position.asc())
        .limit(5_000)
        .all()
    )
    scored: list[tuple[float, AIKnowledgeChunk, AIKnowledgeDocument]] = []
    for chunk, document in rows:
        haystack = unicodedata.normalize(
            "NFD", f"{document.title} {chunk.content}".lower()
        )
        haystack = "".join(ch for ch in haystack if unicodedata.category(ch) != "Mn")
        if terms:
            matches = sum(min(4, haystack.count(term)) for term in terms)
            coverage = sum(1 for term in terms if term in haystack) / len(terms)
            if matches == 0:
                continue
            score = matches + coverage * 5
        else:
            score = 1 / (1 + chunk.position)
        scored.append((score, chunk, document))
    scored.sort(key=lambda item: (-item[0], item[2].id, item[1].position))
    results = [
        {
            "document_id": document.id,
            "title": document.title,
            "original_name": document.original_name,
            "chunk": chunk.position,
            "score": round(score, 3),
            "content": chunk.content[:1_800],
        }
        for score, chunk, document in scored[: max(1, min(12, limit))]
    ]
    return {"ok": True, "query": query, "results": results, "count": len(results)}
