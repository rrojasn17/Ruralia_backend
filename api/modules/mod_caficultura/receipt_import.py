from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
import re
import unicodedata
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.orm import Session

from modules.mod_caficultura.models import Cliente, Finca, ReciboCafe
from core.models import Usuario
from modules.mod_caficultura.receipt_portal import normalize_identification
from modules.mod_caficultura.services import (
    advance_receipt_counter,
    cosecha_from_date,
    get_receipt_counter,
    next_code,
    numero_a_letras_es,
    recibo_payload,
)


MAX_IMPORT_BYTES = 8 * 1024 * 1024
MAX_IMPORT_ROWS = 5_000
REQUIRED_HEADERS = {
    "numero_recibo",
    "fecha",
    "cliente_identificacion",
    "cliente_nombre",
    "cajuelas",
}
TEMPLATE_HEADERS = [
    "numero_recibo",
    "fecha",
    "cliente_identificacion",
    "cliente_nombre",
    "cliente_telefono",
    "cliente_correo",
    "finca_codigo",
    "finca_nombre",
    "cosecha",
    "cajuelas",
    "cuartillos",
    "porcentaje_flote",
    "porcentaje_verde",
    "peso_promedio_cajuela",
    "precio_fanega",
    "precio_fanega_letras",
    "zona",
    "provincia",
    "canton",
    "distrito",
    "beneficio_recibe",
    "productor_entrega",
    "estado",
    "observaciones",
]


class ReceiptImportValidationError(ValueError):
    def __init__(self, errors: list[str]):
        super().__init__("El archivo contiene errores de validación")
        self.errors = errors


def build_receipt_import_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Recibos"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:X1"

    header_fill = PatternFill("solid", fgColor="163B57")
    required_fill = PatternFill("solid", fgColor="145A63")
    for column, header in enumerate(TEMPLATE_HEADERS, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = required_fill if header in REQUIRED_HEADERS else header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[cell.column_letter].width = max(16, min(30, len(header) + 4))
    sheet.row_dimensions[1].height = 28

    date_validation = DataValidation(
        type="date",
        operator="between",
        formula1="DATE(1900,1,1)",
        formula2="DATE(2100,12,31)",
        allow_blank=False,
    )
    state_validation = DataValidation(type="list", formula1='"recibido,en_lote,anulado"', allow_blank=True)
    sheet.add_data_validation(date_validation)
    sheet.add_data_validation(state_validation)
    date_validation.add("B2:B5001")
    state_validation.add("W2:W5001")
    for row in range(2, 5002):
        sheet.cell(row=row, column=2).number_format = "yyyy-mm-dd"

    instructions = workbook.create_sheet("Instrucciones")
    instructions.column_dimensions["A"].width = 30
    instructions.column_dimensions["B"].width = 100
    notes = [
        ("Regla", "Importación transaccional: si una fila tiene errores, no se guarda ninguna."),
        ("Obligatorios", "numero_recibo, fecha, cliente_identificacion, cliente_nombre y cajuelas."),
        ("Número", "Respete exactamente el consecutivo del recibo físico. Los duplicados se rechazan."),
        ("Fecha", "Use una fecha real de Excel o el formato AAAA-MM-DD."),
        ("Cliente", "Si la identificación ya existe se reutiliza el cliente; si no existe, se crea con el nombre indicado."),
        ("Finca", "Es opcional. Si informa un código nuevo debe indicar también finca_nombre."),
        ("Cuartillos", "Use un entero de 0 a 3."),
        ("Estado", "Valores permitidos: recibido, en_lote o anulado. Si se deja vacío se usa recibido."),
        ("Pagos", "La importación histórica no marca recibos como liquidados; los pagos requieren su flujo auditable y comprobante."),
        ("Límite", f"Máximo {MAX_IMPORT_ROWS:,} recibos por archivo."),
    ]
    for row, (label, description) in enumerate(notes, start=1):
        instructions.cell(row=row, column=1, value=label).font = Font(bold=True, color="145A63")
        instructions.cell(row=row, column=2, value=description).alignment = Alignment(wrap_text=True, vertical="top")

    example = workbook.create_sheet("Ejemplo (no importar)")
    example.append(TEMPLATE_HEADERS)
    example.append([
        "12548",
        date(2023, 11, 18),
        "1-0123-0456",
        "Productor de ejemplo",
        "88887777",
        "productor@ejemplo.com",
        "FIN-00125",
        "Finca El Cafetal",
        "2023-2024",
        18,
        2,
        3.5,
        1.2,
        12.8,
        120000,
        "ciento veinte mil colones",
        "A",
        "Cartago",
        "Central",
        "Corralillo",
        "Operario histórico",
        "Productor de ejemplo",
        "recibido",
        "Migrado desde recibo físico",
    ])
    example.sheet_state = "hidden"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def import_historical_receipts(
    db: Session,
    file_bytes: bytes,
    current: Usuario,
) -> dict[str, int | bool | list[str]]:
    rows = _read_rows(file_bytes)
    errors: list[str] = []
    normalized_rows: list[dict[str, object]] = []
    seen_numbers: set[str] = set()

    for row_number, item in rows:
        try:
            normalized = _normalize_row(row_number, item)
            receipt_number = str(normalized["numero_recibo"])
            if receipt_number in seen_numbers:
                raise ValueError(f"número de recibo repetido dentro del archivo: {receipt_number}")
            seen_numbers.add(receipt_number)
            normalized_rows.append(normalized)
        except ValueError as exc:
            errors.append(f"Fila {row_number}: {exc}")

    if not normalized_rows and not errors:
        errors.append("El archivo no contiene recibos para importar")
    if errors:
        raise ReceiptImportValidationError(errors[:100])

    existing_numbers = {
        value
        for (value,) in db.query(ReciboCafe.numero_recibo).filter(
            ReciboCafe.numero_recibo.in_(seen_numbers)
        ).all()
    }
    for normalized in normalized_rows:
        if normalized["numero_recibo"] in existing_numbers:
            errors.append(
                f"Fila {normalized['row_number']}: el recibo {normalized['numero_recibo']} ya existe"
            )

    clients_by_identification: dict[str, Cliente] = {}
    ambiguous_identifications: set[str] = set()
    for client in db.query(Cliente).all():
        key = normalize_identification(client.numero_identificacion)
        if not key:
            continue
        if key in clients_by_identification:
            ambiguous_identifications.add(key)
        else:
            clients_by_identification[key] = client

    planned_clients: dict[str, dict[str, object]] = {}
    for normalized in normalized_rows:
        identification = str(normalized["identification_key"])
        if identification in ambiguous_identifications:
            errors.append(
                f"Fila {normalized['row_number']}: hay más de un cliente existente con esa identificación normalizada"
            )
            continue
        if identification not in clients_by_identification:
            planned_clients.setdefault(identification, normalized)

    farms_by_code = {
        str(farm.codigo or "").strip().upper(): farm
        for farm in db.query(Finca).all()
        if str(farm.codigo or "").strip()
    }
    planned_farm_codes: dict[str, str] = {}
    for normalized in normalized_rows:
        farm_code = str(normalized.get("finca_codigo") or "").strip().upper()
        farm_name = str(normalized.get("finca_nombre") or "").strip()
        identification = str(normalized["identification_key"])
        if not farm_code:
            continue
        existing_farm = farms_by_code.get(farm_code)
        existing_client = clients_by_identification.get(identification)
        if existing_farm and (not existing_client or existing_farm.cliente_id != existing_client.id):
            errors.append(
                f"Fila {normalized['row_number']}: la finca {farm_code} pertenece a otro cliente"
            )
        elif not existing_farm:
            if not farm_name:
                errors.append(
                    f"Fila {normalized['row_number']}: indique finca_nombre para crear la finca {farm_code}"
                )
            previous_owner = planned_farm_codes.get(farm_code)
            if previous_owner and previous_owner != identification:
                errors.append(
                    f"Fila {normalized['row_number']}: el código de finca {farm_code} se asignó a dos clientes"
                )
            planned_farm_codes[farm_code] = identification

    if errors:
        raise ReceiptImportValidationError(errors[:100])

    # Serializa importaciones, cambios de configuración y emisiones normales.
    get_receipt_counter(db, lock=True)
    raced_numbers = {
        value
        for (value,) in db.query(ReciboCafe.numero_recibo).filter(
            ReciboCafe.numero_recibo.in_(seen_numbers)
        ).all()
    }
    if raced_numbers:
        raise ReceiptImportValidationError([
            f"Los recibos {', '.join(sorted(raced_numbers)[:10])} fueron registrados mientras se validaba el archivo. Vuelva a revisar la plantilla."
        ])

    created_clients = 0
    for identification, source in planned_clients.items():
        client = Cliente(
            codigo=next_code(db, Cliente, "codigo", "CL-", 5),
            nombre_completo=str(source["cliente_nombre"]),
            tipo_persona="personal",
            categoria="agro",
            numero_identificacion=str(source["cliente_identificacion"]),
            telefono=str(source.get("cliente_telefono") or "").strip() or None,
            correo=str(source.get("cliente_correo") or "").strip() or None,
            provincia=str(source.get("provincia") or "").strip() or None,
            canton=str(source.get("canton") or "").strip() or None,
            distrito=str(source.get("distrito") or "").strip() or None,
            activo=True,
        )
        db.add(client)
        db.flush()
        clients_by_identification[identification] = client
        created_clients += 1

    created_farms = 0
    planned_farms: dict[tuple[str, str], Finca] = {}
    for normalized in normalized_rows:
        identification = str(normalized["identification_key"])
        client = clients_by_identification[identification]
        farm_code = str(normalized.get("finca_codigo") or "").strip().upper()
        farm_name = str(normalized.get("finca_nombre") or "").strip()
        if not farm_code and not farm_name:
            continue

        key = (identification, farm_code or farm_name.casefold())
        farm = farms_by_code.get(farm_code) if farm_code else None
        if not farm and not farm_code:
            farm = db.query(Finca).filter(
                Finca.cliente_id == client.id,
                Finca.nombre.ilike(farm_name),
            ).first()
        if not farm:
            farm = planned_farms.get(key)
        if not farm:
            farm = Finca(
                cliente_id=client.id,
                codigo=farm_code or next_code(db, Finca, "codigo", "FIN-", 5),
                nombre=farm_name,
                propietario=client.nombre_completo,
                provincia=str(normalized.get("provincia") or "").strip() or client.provincia,
                canton=str(normalized.get("canton") or "").strip() or client.canton,
                distrito=str(normalized.get("distrito") or "").strip() or client.distrito,
                cultivo="Café",
                activa=True,
            )
            db.add(farm)
            db.flush()
            planned_farms[key] = farm
            if farm_code:
                farms_by_code[farm_code] = farm
            created_farms += 1
        normalized["finca_id"] = farm.id

    for normalized in normalized_rows:
        identification = str(normalized["identification_key"])
        client = clients_by_identification[identification]
        price = float(normalized.get("precio_fanega") or 0)
        price_words = str(normalized.get("precio_fanega_letras") or "").strip() or (
            numero_a_letras_es(price) if price > 0 else None
        )
        row = ReciboCafe(
            numero_recibo=str(normalized["numero_recibo"]),
            fecha=normalized["fecha"],
            cosecha=str(normalized.get("cosecha") or "").strip() or cosecha_from_date(normalized["fecha"]),
            cliente_id=client.id,
            finca_id=normalized.get("finca_id"),
            productor_nombre=client.nombre_completo,
            productor_cedula=client.numero_identificacion,
            provincia=str(normalized.get("provincia") or "").strip() or client.provincia,
            canton=str(normalized.get("canton") or "").strip() or client.canton,
            distrito=str(normalized.get("distrito") or "").strip() or client.distrito,
            zona=str(normalized.get("zona") or "").strip() or None,
            cajuelas=float(normalized["cajuelas"]),
            cuartillos=float(normalized.get("cuartillos") or 0),
            porcentaje_flote=normalized.get("porcentaje_flote"),
            porcentaje_verde=normalized.get("porcentaje_verde"),
            peso_promedio_cajuela=normalized.get("peso_promedio_cajuela"),
            precio_promedio_cajuela=normalized.get("peso_promedio_cajuela"),
            precio_fanega=price,
            precio_fanega_letras=price_words,
            precio_adelanto_letras=price_words,
            beneficio_recibe=str(normalized.get("beneficio_recibe") or "").strip() or None,
            productor_entrega=str(normalized.get("productor_entrega") or "").strip() or None,
            estado=str(normalized.get("estado") or "recibido"),
            observaciones=str(normalized.get("observaciones") or "").strip() or None,
            created_by_id=current.id,
        )
        row.qr_payload = recibo_payload(row)
        db.add(row)
        advance_receipt_counter(db, row.numero_recibo)

    db.flush()
    return {
        "ok": True,
        "recibos_creados": len(normalized_rows),
        "clientes_creados": created_clients,
        "fincas_creadas": created_farms,
        "errores": [],
    }


def _read_rows(file_bytes: bytes) -> list[tuple[int, dict[str, object]]]:
    if not file_bytes:
        raise ReceiptImportValidationError(["El archivo está vacío"])
    if len(file_bytes) > MAX_IMPORT_BYTES:
        raise ReceiptImportValidationError(["El archivo supera el límite de 8 MB"])
    try:
        with ZipFile(BytesIO(file_bytes)) as archive:
            members = archive.infolist()
            uncompressed_size = sum(member.file_size for member in members)
            if len(members) > 250 or uncompressed_size > 60 * 1024 * 1024:
                raise ReceiptImportValidationError(["El archivo Excel expandido es demasiado grande"])
    except BadZipFile as exc:
        raise ReceiptImportValidationError(["El archivo no es un .xlsx válido"]) from exc
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ReceiptImportValidationError(["No se pudo leer el archivo .xlsx"]) from exc

    sheet = workbook["Recibos"] if "Recibos" in workbook.sheetnames else workbook.active
    iterator = sheet.iter_rows(values_only=True)
    first = next(iterator, None)
    if not first:
        raise ReceiptImportValidationError(["La hoja Recibos no contiene encabezados"])
    headers = [_normalize_header(value) for value in first]
    duplicate_headers = {header for header in headers if header and headers.count(header) > 1}
    if duplicate_headers:
        raise ReceiptImportValidationError([
            f"Hay columnas duplicadas: {', '.join(sorted(duplicate_headers))}"
        ])
    missing = REQUIRED_HEADERS - set(headers)
    if missing:
        raise ReceiptImportValidationError([
            f"Faltan columnas obligatorias: {', '.join(sorted(missing))}"
        ])

    result: list[tuple[int, dict[str, object]]] = []
    for row_number, raw in enumerate(iterator, start=2):
        values = {
            header: raw[index] if index < len(raw) else None
            for index, header in enumerate(headers)
            if header
        }
        if not any(value not in (None, "") for value in values.values()):
            continue
        result.append((row_number, values))
        if len(result) > MAX_IMPORT_ROWS:
            raise ReceiptImportValidationError([
                f"El archivo supera el máximo de {MAX_IMPORT_ROWS:,} recibos"
            ])
    return result


def _normalize_row(row_number: int, item: dict[str, object]) -> dict[str, object]:
    receipt_number = _clean_text_number(item.get("numero_recibo"), 80)
    if not receipt_number:
        raise ValueError("falta numero_recibo")
    receipt_date = _parse_date(item.get("fecha"))
    identification = _clean_text_number(item.get("cliente_identificacion"), 80)
    identification_key = normalize_identification(identification)
    if len(identification_key) < 3:
        raise ValueError("cliente_identificacion no es válida")
    client_name = _clean_text(item.get("cliente_nombre"), 220)
    if len(client_name) < 2:
        raise ValueError("falta cliente_nombre")

    cajuelas = _parse_float(item.get("cajuelas"), "cajuelas", required=True, minimum=0)
    cuartillos = _parse_float(item.get("cuartillos"), "cuartillos", minimum=0, maximum=3) or 0
    if float(cuartillos).is_integer() is False:
        raise ValueError("cuartillos debe ser un entero de 0 a 3")
    if float(cajuelas or 0) + (float(cuartillos) / 4) <= 0:
        raise ValueError("la cantidad de café debe ser mayor que cero")

    state = _clean_text(item.get("estado"), 40).lower() or "recibido"
    if state not in {"recibido", "en_lote", "anulado"}:
        raise ValueError("estado debe ser recibido, en_lote o anulado")

    email = _clean_text(item.get("cliente_correo"), 180).lower()
    if email and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
        raise ValueError("cliente_correo no es válido")

    return {
        "row_number": row_number,
        "numero_recibo": receipt_number,
        "fecha": receipt_date,
        "cliente_identificacion": identification,
        "identification_key": identification_key,
        "cliente_nombre": client_name,
        "cliente_telefono": _clean_text_number(item.get("cliente_telefono"), 80),
        "cliente_correo": email,
        "finca_codigo": _clean_text_number(item.get("finca_codigo"), 80).upper(),
        "finca_nombre": _clean_text(item.get("finca_nombre"), 180),
        "cosecha": _clean_text_number(item.get("cosecha"), 30),
        "cajuelas": cajuelas,
        "cuartillos": int(cuartillos),
        "porcentaje_flote": _parse_float(item.get("porcentaje_flote"), "porcentaje_flote", minimum=0, maximum=100),
        "porcentaje_verde": _parse_float(item.get("porcentaje_verde"), "porcentaje_verde", minimum=0, maximum=100),
        "peso_promedio_cajuela": _parse_float(item.get("peso_promedio_cajuela"), "peso_promedio_cajuela", minimum=0),
        "precio_fanega": _parse_float(item.get("precio_fanega"), "precio_fanega", minimum=0) or 0,
        "precio_fanega_letras": _clean_text(item.get("precio_fanega_letras"), 255),
        "zona": _clean_text(item.get("zona"), 40),
        "provincia": _clean_text(item.get("provincia"), 80),
        "canton": _clean_text(item.get("canton"), 80),
        "distrito": _clean_text(item.get("distrito"), 80),
        "beneficio_recibe": _clean_text(item.get("beneficio_recibe"), 180),
        "productor_entrega": _clean_text(item.get("productor_entrega"), 180),
        "estado": state,
        "observaciones": _clean_text(item.get("observaciones"), 2_000),
    }


def _normalize_header(value: object) -> str:
    raw = unicodedata.normalize("NFKD", str(value or "").strip().lower())
    clean = "".join(character for character in raw if not unicodedata.combining(character))
    return re.sub(r"[^a-z0-9]+", "_", clean).strip("_")


def _clean_text(value: object, maximum: int) -> str:
    clean = " ".join(str(value or "").strip().split())
    if len(clean) > maximum:
        raise ValueError(f"un texto supera el máximo de {maximum} caracteres")
    return clean


def _clean_text_number(value: object, maximum: int) -> str:
    if isinstance(value, float) and value.is_integer():
        clean = str(int(value))
    elif isinstance(value, int):
        clean = str(value)
    else:
        clean = str(value or "").strip()
    if len(clean) > maximum:
        raise ValueError(f"un código supera el máximo de {maximum} caracteres")
    return clean


def _parse_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value or "").strip()
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw, pattern).date()
        except ValueError:
            continue
    raise ValueError("fecha inválida; use AAAA-MM-DD")


def _parse_float(
    value: object,
    label: str,
    *,
    required: bool = False,
    minimum: float | None = None,
    maximum: float | None = None,
) -> float | None:
    if value in (None, ""):
        if required:
            raise ValueError(f"falta {label}")
        return None
    try:
        if isinstance(value, str):
            raw = value.strip().replace(" ", "")
            if "," in raw and "." in raw:
                raw = raw.replace(".", "").replace(",", ".")
            else:
                raw = raw.replace(",", ".")
            number = float(raw)
        else:
            number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} no es numérico") from exc
    if minimum is not None and number < minimum:
        raise ValueError(f"{label} debe ser mayor o igual que {minimum:g}")
    if maximum is not None and number > maximum:
        raise ValueError(f"{label} debe ser menor o igual que {maximum:g}")
    return number
