from __future__ import annotations

from collections import defaultdict
import json
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Optional
import os
import uuid
from urllib.parse import quote

import qrcode
from PIL import Image, UnidentifiedImageError
from qrcode.image.svg import SvgImage
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status, Response, Request
from sqlalchemy import func, or_
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, selectinload

from config import ADMIN_EMAIL, APP_BASE_URL, APP_NAME, MAX_REQUEST_MB
from database import get_db
from modules.mod_caficultura.models import (
    ActividadFinca,
    Cliente,
    Finca,
    InsumoFinca,
    Proveedor,
    ProveedorContacto,
    CompraInsumoFactura,
    CompraInsumoLinea,
    OrdenTrabajo,
    OrdenTrabajoRecibo,
    OrdenTrabajoUnionOrigen,
    ReciboCafe,
    ReciboConsecutivo,
    RegistroFinca,
    RegistroFincaInsumo,
    RegistroFincaTrabajador,
    SeguimientoOT,
    TrabajadorFinca,
    DocumentoLote,
    ComentarioLote,
    CotizacionVenta,
    CotizacionVentaLinea,
    SolicitudSalidaVenta,
    SolicitudSalidaVentaLinea,
    VentaLeadLote,
    VentaSeguimiento,
    SolicitudVenta,
    SolicitudVentaLinea,
    SolicitudVentaLineaLiquidacion,
    SolicitudVentaSeguimiento,
    SolicitudVentaDocumento,
)
from core.models import ConfiguracionSistema, SesionUsuario, Usuario
from core.routers.auth import get_current_user, require_roles
from modules.mod_caficultura.schemas import (
    ClienteCreate,
    ClienteOut,
    ClienteUpdate,
    ConfiguracionOut,
    ConfiguracionUpdate,
    BrandingOut,
    DashboardOut,
    DashboardKPI,
    ChartPoint,
    ActividadFincaCreate,
    ActividadFincaOut,
    ActividadFincaUpdate,
    FincaCreate,
    FincaOut,
    FincaUpdate,
    MapLocationResolveIn,
    MapLocationResolveOut,
    InsumoFincaCreate,
    InsumoFincaOut,
    ProveedorCreate,
    ProveedorUpdate,
    ProveedorOut,
    InsumoFincaUpdate,
    CompraInsumoCreate,
    CompraInsumoOut,
    OTCreate,
    OTOut,
    OTUpdate,
    OTMergePayload,
    ReciboCreate,
    ReciboConsecutivoOut,
    ReciboConsecutivoUpdate,
    ReciboOut,
    ReciboUpdate,
    ReciboDisponibleOTOut,
    RegistroFincaCreate,
    RegistroFincaOut,
    RegistroFincaUpdate,
    SeguimientoCreate,
    SeguimientoRevisionPayload,
    EstadoOTPayload,
    ComentarioLoteCreate,
    CotizacionVentaCreate,
    CotizacionVentaUpdate,
    CotizacionVentaOut,
    SolicitudSalidaVentaCreate,
    SolicitudSalidaVentaOut,
    VentaLeadCreate,
    VentaLeadUpdate,
    VentaSeguimientoCreate,
    SolicitudVentaCreate,
    SolicitudVentaUpdate,
    SolicitudVentaSeguimientoCreate,
    SolicitudVentaLiquidarPayload,
    AssignSolicitudVentaPayload,
    SolicitudVentaOut,
    SyncPush,
    SyncResult,
    SyncResultItem,
    TrabajadorFincaCreate,
    TrabajadorFincaOut,
    TrabajadorFincaUpdate,
    UsuarioCreate,
    UsuarioOut,
    UsuarioUpdate,
    PublicLoteOut,
    normalize_role,
)
from security import hashear_contrasena, validar_fortaleza_contrasena
from services.map_location import MapLocationError, resolve_map_location
from modules.mod_caficultura.services import (
    create_ot_from_payload,
    create_recibo_from_payload,
    estimated_fanegas,
    get_config_int,
    next_code,
    ot_query,
    recibo_query,
    serialize_cliente,
    serialize_finca,
    serialize_ot,
    serialize_recibo,
    set_config,
    usuario_out,
    utcnow,
    lote_public_url,
    build_ot_qr_payload,
    serialize_public_lote,
    serialize_recibo_disponible_ot,
    recibo_disponible_cajuelas,
    serialize_lote_documento,
    serialize_lote_comentario,
    serialize_cotizacion_venta,
    serialize_solicitud_salida_venta,
    serialize_venta_lead,
    serialize_venta_seguimiento,
    serialize_solicitud_venta,
    serialize_solicitud_documento,
    serialize_solicitud_seguimiento,
    numero_a_letras_es,
    get_receipt_counter,
    last_receipt_number,
    preview_receipt_number,
)
from modules.mod_caficultura.receipt_import import (
    MAX_IMPORT_BYTES,
    ReceiptImportValidationError,
    build_receipt_import_template,
    import_historical_receipts,
)

router = APIRouter(tags=["navia"])

def make_qr_svg(data: str) -> str:
    img = qrcode.make(data, image_factory=SvgImage)
    buffer = BytesIO()
    img.save(buffer)
    return buffer.getvalue().decode("utf-8")


def ensure_superadmin(current: Usuario) -> None:
    if not bool(getattr(current, "is_superadmin", False)):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Solo un superadmin puede eliminar registros",
        )


def normalize_finca_map_location(data: dict) -> None:
    if "google_maps_url" not in data:
        return

    raw_value = str(data.get("google_maps_url") or "").strip()
    if not raw_value:
        data["google_maps_url"] = None
        return

    try:
        location = resolve_map_location(raw_value)
    except MapLocationError as exc:
        raise HTTPException(
            status_code=422,
            detail=f"No se pudo verificar la ubicación de la finca: {exc}",
        ) from exc

    data["google_maps_url"] = location.canonical_url


def is_bootstrap_superadmin(current: Usuario) -> bool:
    return bool(getattr(current, "is_superadmin", False)) and str(getattr(current, "correo", "")).strip().lower() == ADMIN_EMAIL


def ensure_bootstrap_superadmin(current: Usuario) -> None:
    if not is_bootstrap_superadmin(current):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Solo el superadmin creado por bootstrap puede ejecutar esta acción",
        )


def backup_allowed(current: Usuario) -> bool:
    return bool(getattr(current, "is_superadmin", False)) or str(getattr(current, "rol", "")).strip().lower() == "gerente"

def upload_root() -> Path:
    return Path(os.getenv("NAVIA_UPLOAD_DIR", "uploads"))




def ensure_immutable_code(current_value: str | None, incoming_value: str | None, label: str = "código") -> None:
    """Evita modificar códigos operativos una vez creados.

    El frontend puede reenviar el mismo código en un formulario de edición; eso se permite.
    Un valor distinto se rechaza para mantener trazabilidad e integridad histórica.
    """
    if incoming_value is None:
        return

    current_norm = str(current_value or "").strip()
    incoming_norm = str(incoming_value or "").strip()

    if incoming_norm and incoming_norm != current_norm:
        raise HTTPException(status_code=400, detail=f"El {label} no se puede modificar una vez creado")


def next_entity_code(db: Session, entidad: str) -> str:
    entity = (entidad or "").strip().lower()
    if entity == "recibo":
        return preview_receipt_number(db)
    mapping = {
        "cliente": (Cliente, "codigo", "CL-", 5),
        "finca": (Finca, "codigo", "FIN-", 5),
        "trabajador": (TrabajadorFinca, "codigo", "TR-", 3),
        "insumo": (InsumoFinca, "codigo", "INS-", 3),
        "ot": (OrdenTrabajo, "codigo_lote", "LOT-", 5),
    }

    if entity not in mapping:
        raise HTTPException(status_code=422, detail="Entidad de código no soportada")

    model, field, prefix, width = mapping[entity]
    return next_code(db, model, field, prefix, width)

def add_lote_evento(db: Session, ot_id: int, tipo: str, comentario: str, current: Usuario | None = None) -> None:
    db.add(ComentarioLote(
        ot_id=ot_id,
        tipo=(tipo or "nota").strip().lower(),
        comentario=(comentario or "").strip(),
        created_by_id=current.id if current else None,
    ))


def ensure_lote_not_unido(row: OrdenTrabajo, accion: str = "modificar") -> None:
    if (row.estado or "").lower() == "unido":
        raise HTTPException(
            status_code=409,
            detail=f"Este lote ya fue unido a otro lote. No se puede {accion}; se conserva únicamente como trazabilidad histórica."
        )


def get_config_value(db: Session, key: str, default: str = "") -> str:
    row = db.query(ConfiguracionSistema).filter(ConfiguracionSistema.clave == key).first()
    return str(row.valor or default) if row else default


def clean_phone(value: str | None) -> str:
    return "".join(ch for ch in str(value or "") if ch.isdigit())


def build_alert_links(db: Session, message: str) -> tuple[str | None, str | None]:
    phone = clean_phone(get_config_value(db, "alert_whatsapp", ""))
    email = get_config_value(db, "alert_email", "").strip()
    encoded = quote(message)
    whatsapp_url = f"https://wa.me/{phone}?text={encoded}" if phone else None
    correo_url = f"mailto:{email}?subject={quote('Alerta NAVIA')}&body={encoded}" if email else None
    return whatsapp_url, correo_url

def validate_signature_image(content: bytes) -> str:
    if not content:
        raise HTTPException(status_code=422, detail="La imagen de firma está vacía")
    try:
        with Image.open(BytesIO(content)) as image:
            if image.width * image.height > 25_000_000:
                raise HTTPException(status_code=422, detail="La firma tiene dimensiones demasiado grandes")
            image.verify()
            image_format = str(image.format or "").upper()
    except HTTPException:
        raise
    except (UnidentifiedImageError, OSError, ValueError) as exc:
        raise HTTPException(status_code=422, detail="El archivo no es una imagen válida") from exc

    allowed_formats = {"PNG": ".png", "JPEG": ".jpg", "WEBP": ".webp"}
    if image_format not in allowed_formats:
        raise HTTPException(status_code=422, detail="La firma debe ser PNG, JPG o WEBP")
    return allowed_formats[image_format]


@router.get("/codigos/siguiente")
def get_next_codigo(
    entidad: str,
    current: Usuario = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    return {"entidad": entidad, "codigo": next_entity_code(db, entidad)}


def build_configuracion_out(db: Session, current: Usuario) -> ConfiguracionOut:
    can_backup = backup_allowed(current)
    return ConfiguracionOut(
        session_days=get_config_int(db, "session_days", 30),
        app_name=get_config_value(db, "app_name", APP_NAME).strip() or APP_NAME,
        app_logo_url=get_config_value(db, "app_logo_url", "").strip() or None,
        alert_whatsapp=get_config_value(db, "alert_whatsapp", "") or None,
        alert_email=get_config_value(db, "alert_email", "") or None,
        can_export_backup=can_backup,
        can_import_backup=can_backup,
        can_full_reset=is_bootstrap_superadmin(current),
    )


# Core route moved to core.routers.admin: @router.get("/branding", response_model=BrandingOut)
def get_public_branding(db: Session = Depends(get_db)):
    """Configuración pública mínima necesaria antes de iniciar sesión."""
    return BrandingOut(
        app_name=get_config_value(db, "app_name", APP_NAME).strip() or APP_NAME,
        app_logo_url=get_config_value(db, "app_logo_url", "").strip() or None,
    )


# Core route moved to core.routers.admin: @router.get("/configuracion", response_model=ConfiguracionOut)
def get_configuracion(current: Usuario = Depends(require_roles("administrativo", "admin")), db: Session = Depends(get_db)):
    return build_configuracion_out(db, current)


# Core route moved to core.routers.admin: @router.patch("/configuracion", response_model=ConfiguracionOut)
def update_configuracion(payload: ConfiguracionUpdate, current: Usuario = Depends(require_roles("administrativo", "admin")), db: Session = Depends(get_db)):
    if payload.session_days is not None:
        set_config(db, "session_days", str(payload.session_days), "Duración de sesión en días")
    if payload.app_name is not None:
        set_config(db, "app_name", payload.app_name.strip(), "Nombre público de la plataforma")
    if payload.alert_whatsapp is not None:
        set_config(db, "alert_whatsapp", str(payload.alert_whatsapp or "").strip(), "Número WhatsApp para alertas operativas")
    if payload.alert_email is not None:
        set_config(db, "alert_email", str(payload.alert_email or "").strip(), "Correo para alertas operativas")
    db.commit()
    return build_configuracion_out(db, current)


# Core route moved to core.routers.admin: @router.post("/configuracion/branding/logo", response_model=ConfiguracionOut)
async def upload_brand_logo(
    file: UploadFile = File(...),
    current: Usuario = Depends(require_roles("administrativo", "admin")),
    db: Session = Depends(get_db),
):
    content = await file.read()
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="El logotipo no debe superar 5 MB")
    try:
        with Image.open(BytesIO(content)) as image:
            if image.width * image.height > 25_000_000:
                raise HTTPException(status_code=422, detail="El logotipo tiene dimensiones demasiado grandes")
            image.verify()
            image_format = str(image.format or "").upper()
    except HTTPException:
        raise
    except (UnidentifiedImageError, OSError, ValueError) as exc:
        raise HTTPException(status_code=422, detail="El archivo no es una imagen válida") from exc
    allowed_formats = {"PNG": ".png", "JPEG": ".jpg", "WEBP": ".webp"}
    extension = allowed_formats.get(image_format)
    if not extension:
        raise HTTPException(status_code=422, detail="El logotipo debe ser PNG, JPG o WEBP")

    root = upload_root().resolve()
    folder = (root / "branding").resolve()
    if root != folder and root not in folder.parents:
        raise HTTPException(status_code=500, detail="Ruta de logotipo inválida")
    folder.mkdir(parents=True, exist_ok=True)
    filename = f"logo_{uuid.uuid4().hex}{extension}"
    target = folder / filename
    target.write_bytes(content)

    previous = get_config_value(db, "app_logo_url", "")
    logo_url = f"{APP_BASE_URL}/uploads/branding/{filename}"
    set_config(db, "app_logo_url", logo_url, "Logotipo institucional de la plataforma")
    db.commit()

    if "/uploads/branding/" in previous:
        previous_path = (folder / Path(previous.split("?", 1)[0]).name).resolve()
        if previous_path != target and previous_path.parent == folder:
            previous_path.unlink(missing_ok=True)
    return build_configuracion_out(db, current)


# Core route moved to core.routers.admin: @router.delete("/configuracion/branding/logo", response_model=ConfiguracionOut)
def reset_brand_logo(
    current: Usuario = Depends(require_roles("administrativo", "admin")),
    db: Session = Depends(get_db),
):
    previous = get_config_value(db, "app_logo_url", "")
    set_config(db, "app_logo_url", "", "Logotipo institucional de la plataforma")
    db.commit()
    if "/uploads/branding/" in previous:
        folder = (upload_root().resolve() / "branding").resolve()
        previous_path = (folder / Path(previous.split("?", 1)[0]).name).resolve()
        if previous_path.parent == folder:
            previous_path.unlink(missing_ok=True)
    return build_configuracion_out(db, current)


# Core route moved to core.routers.admin: @router.get("/usuarios", response_model=list[UsuarioOut])
def list_usuarios(current: Usuario = Depends(require_roles("gerente")), db: Session = Depends(get_db)):
    return [usuario_out(u) for u in db.query(Usuario).order_by(Usuario.id.asc()).all()]


# Core route moved to core.routers.admin: @router.post("/usuarios", response_model=UsuarioOut, status_code=status.HTTP_201_CREATED)
def create_usuario(
    payload: UsuarioCreate,
    current: Usuario = Depends(require_roles("gerente")),
    db: Session = Depends(get_db)
):
    correo = payload.correo.strip().lower()

    if db.query(Usuario).filter(Usuario.correo == correo).first():
        raise HTTPException(status_code=409, detail="Ya existe un usuario con ese correo")

    requested_superadmin = bool(getattr(payload, "is_superadmin", False))

    if requested_superadmin and not bool(getattr(current, "is_superadmin", False)):
        raise HTTPException(
            status_code=403,
            detail="Solo un superadmin puede crear otro superadmin"
        )

    try:
        validar_fortaleza_contrasena(payload.password)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc

    user = Usuario(
        nombre=payload.nombre.strip(),
        correo=correo,
        hash_contrasena=hashear_contrasena(payload.password),
        rol="gerente" if requested_superadmin else normalize_role(payload.rol),
        activo=True if requested_superadmin else payload.activo,
        is_superadmin=requested_superadmin,
    )

    db.add(user)
    db.commit()
    db.refresh(user)

    return usuario_out(user)

# Core route moved to core.routers.admin: @router.get("/usuarios/select", response_model=list[UsuarioOut])
def list_usuarios_select(
    current: Usuario = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """
    Lista usuarios activos para selects operativos:
    P/Beneficio, asignación de OT, responsables, etc.
    """
    rows = (
        db.query(Usuario)
        .filter(Usuario.activo)
        .order_by(Usuario.nombre.asc())
        .all()
    )

    return [usuario_out(row) for row in rows]


# Core route moved to core.routers.admin: @router.patch("/usuarios/{usuario_id}", response_model=UsuarioOut)
def update_usuario(
    usuario_id: int,
    payload: UsuarioUpdate,
    current: Usuario = Depends(require_roles("gerente")),
    db: Session = Depends(get_db)
):
    user = db.query(Usuario).filter(Usuario.id == usuario_id).first()

    if not user:
        raise HTTPException(status_code=404, detail="Usuario no encontrado")

    current_is_superadmin = bool(getattr(current, "is_superadmin", False))
    target_is_superadmin = bool(getattr(user, "is_superadmin", False))

    data = payload.model_dump(exclude_unset=True)

    if target_is_superadmin and not current_is_superadmin:
        raise HTTPException(
            status_code=403,
            detail="Un gerente normal no puede modificar un superadmin"
        )

    if "is_superadmin" in data and not current_is_superadmin:
        raise HTTPException(
            status_code=403,
            detail="Solo un superadmin puede cambiar la bandera de superadmin"
        )

    if user.id == current.id:
        if data.get("activo") is False:
            raise HTTPException(
                status_code=400,
                detail="No puedes desactivar tu propio usuario"
            )

        if "rol" in data and normalize_role(data["rol"]) != "gerente" and target_is_superadmin:
            raise HTTPException(
                status_code=400,
                detail="No puedes quitarte permisos administrativos siendo superadmin"
            )

        if data.get("is_superadmin") is False:
            raise HTTPException(
                status_code=400,
                detail="No puedes quitarte la bandera de superadmin a ti mismo"
            )

    if target_is_superadmin:
        if "activo" in data and data["activo"] is False:
            raise HTTPException(
                status_code=400,
                detail="No se puede desactivar un superadmin"
            )

        if "rol" in data and normalize_role(data["rol"]) != "gerente":
            raise HTTPException(
                status_code=400,
                detail="Un superadmin siempre debe conservar rol gerente"
            )

    if "correo" in data and data["correo"]:
        correo = str(data["correo"]).strip().lower()

        exists = (
            db.query(Usuario)
            .filter(Usuario.correo == correo, Usuario.id != usuario_id)
            .first()
        )

        if exists:
            raise HTTPException(status_code=409, detail="Ese correo ya está en uso")

        user.correo = correo

    if "nombre" in data and data["nombre"]:
        user.nombre = data["nombre"].strip()

    if "is_superadmin" in data:
        nuevo_superadmin = bool(data["is_superadmin"])

        if target_is_superadmin and not nuevo_superadmin:
            activos = (
                db.query(Usuario)
                .filter(Usuario.is_superadmin, Usuario.activo)
                .count()
            )

            if activos <= 1:
                raise HTTPException(
                    status_code=400,
                    detail="Debe existir al menos un superadmin activo"
                )

        user.is_superadmin = nuevo_superadmin

        if nuevo_superadmin:
            user.rol = "gerente"
            user.activo = True

    if "rol" in data and data["rol"]:
        if bool(getattr(user, "is_superadmin", False)):
            user.rol = "gerente"
        else:
            user.rol = normalize_role(data["rol"])

    if "activo" in data:
        if bool(getattr(user, "is_superadmin", False)) and data["activo"] is False:
            raise HTTPException(
                status_code=400,
                detail="No se puede desactivar un superadmin"
            )

        user.activo = bool(data["activo"])

    if "firma_url" in data:
        user.firma_url = data["firma_url"] or None

    if data.get("password"):
        try:
            validar_fortaleza_contrasena(data["password"])
        except ValueError as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc
        user.hash_contrasena = hashear_contrasena(data["password"])

    db.commit()
    db.refresh(user)

    return usuario_out(user)

# Core route moved to core.routers.admin: @router.post("/usuarios/{usuario_id}/firma", response_model=UsuarioOut)
async def upload_usuario_firma(
    usuario_id: int,
    request: Request,
    file: UploadFile = File(...),
    current: Usuario = Depends(require_roles("gerente")),
    db: Session = Depends(get_db),
):
    user = db.query(Usuario).filter(Usuario.id == usuario_id).first()

    if not user:
        raise HTTPException(status_code=404, detail="Usuario no encontrado")

    current_is_superadmin = bool(getattr(current, "is_superadmin", False))
    target_is_superadmin = bool(getattr(user, "is_superadmin", False))

    if target_is_superadmin and not current_is_superadmin:
        raise HTTPException(status_code=403, detail="Solo un superadmin puede cambiar la firma de otro superadmin")

    content = await file.read()
    if len(content) > 3 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="La imagen de firma no debe superar 3 MB")

    ext = validate_signature_image(content)
    folder = upload_root() / "firmas"
    folder.mkdir(parents=True, exist_ok=True)

    filename = f"usuario_{usuario_id}_{uuid.uuid4().hex}{ext}"
    path = folder / filename
    path.write_bytes(content)

    user.firma_url = f"{APP_BASE_URL}/uploads/firmas/{filename}"

    db.commit()
    db.refresh(user)

    return usuario_out(user)


# Core route moved to core.routers.admin: @router.delete("/usuarios/{usuario_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_usuario(
    usuario_id: int,
    current: Usuario = Depends(require_roles("gerente")),
    db: Session = Depends(get_db)
):
    user = db.query(Usuario).filter(Usuario.id == usuario_id).first()

    if not user:
        raise HTTPException(status_code=404, detail="Usuario no encontrado")

    current_is_superadmin = bool(getattr(current, "is_superadmin", False))
    target_is_superadmin = bool(getattr(user, "is_superadmin", False))

    if user.id == current.id:
        raise HTTPException(
            status_code=400,
            detail="No puedes eliminar tu propio usuario"
        )

    if target_is_superadmin:
        raise HTTPException(
            status_code=400,
            detail="No se puede eliminar un superadmin"
        )

    if target_is_superadmin and not current_is_superadmin:
        raise HTTPException(
            status_code=403,
            detail="Un gerente normal no puede eliminar un superadmin"
        )

    dependencias = []

    ot_asignadas = db.query(OrdenTrabajo).filter(OrdenTrabajo.operario_id == usuario_id).count()
    if ot_asignadas:
        dependencias.append(f"{ot_asignadas} OT asignada(s) como operario")

    ot_creadas = db.query(OrdenTrabajo).filter(OrdenTrabajo.created_by_id == usuario_id).count()
    if ot_creadas:
        dependencias.append(f"{ot_creadas} OT creada(s)")

    ot_cerradas = db.query(OrdenTrabajo).filter(OrdenTrabajo.gerente_cierra_id == usuario_id).count()
    if ot_cerradas:
        dependencias.append(f"{ot_cerradas} OT cerrada(s) por este usuario")

    recibos_creados = db.query(ReciboCafe).filter(ReciboCafe.created_by_id == usuario_id).count()
    if recibos_creados:
        dependencias.append(f"{recibos_creados} recibo(s) creado(s)")

    seguimientos_creados = db.query(SeguimientoOT).filter(SeguimientoOT.created_by_id == usuario_id).count()
    if seguimientos_creados:
        dependencias.append(f"{seguimientos_creados} seguimiento(s) creado(s)")

    seguimientos_revisados = db.query(SeguimientoOT).filter(SeguimientoOT.revisado_por_id == usuario_id).count()
    if seguimientos_revisados:
        dependencias.append(f"{seguimientos_revisados} seguimiento(s) revisado(s)")

    if dependencias:
        raise HTTPException(
            status_code=409,
            detail=(
                "No se puede eliminar este usuario porque ya tiene información operativa asociada: "
                + "; ".join(dependencias)
            )
        )

    db.query(SesionUsuario).filter(SesionUsuario.usuario_id == usuario_id).delete(synchronize_session=False)

    db.delete(user)
    db.commit()

    return Response(status_code=status.HTTP_204_NO_CONTENT)

@router.get("/empleados", response_model=list[UsuarioOut])
def list_empleados(current: Usuario = Depends(get_current_user), db: Session = Depends(get_db)):
    rows = db.query(Usuario).filter(Usuario.rol.in_(["operario", "gerente"]), Usuario.activo).order_by(Usuario.rol.asc(), Usuario.nombre.asc()).all()
    return [usuario_out(u) for u in rows]





@router.get("/clientes", response_model=list[ClienteOut])
def list_clientes(current: Usuario = Depends(get_current_user), db: Session = Depends(get_db)):
    rows = db.query(Cliente).order_by(Cliente.nombre_completo.asc()).all()
    return [serialize_cliente(row) for row in rows]


@router.post("/clientes", response_model=ClienteOut, status_code=201)
def create_cliente(
    payload: ClienteCreate,
    current: Usuario = Depends(require_roles("admin")),
    db: Session = Depends(get_db),
):
    codigo = (payload.codigo or "").strip() or next_code(db, Cliente, "codigo", "CL-", 5)

    row = Cliente(
        codigo=codigo,
        nombre_completo=payload.nombre_completo.strip(),
        tipo_persona=payload.tipo_persona,
        categoria=payload.categoria,
        numero_identificacion=payload.numero_identificacion.strip(),
        direccion=payload.direccion,
        provincia=payload.provincia,
        canton=payload.canton,
        distrito=payload.distrito,
        telefono=(payload.telefono or "").strip() or None,
        correo=str(payload.correo).strip() if payload.correo else None,
        activo=payload.activo,
    )

    db.add(row)

    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Código o identificación de cliente duplicada")

    db.refresh(row)
    return serialize_cliente(row)


@router.get("/clientes/{cliente_id}", response_model=ClienteOut)
def get_cliente(
    cliente_id: int,
    current: Usuario = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    row = db.query(Cliente).filter(Cliente.id == cliente_id).first()

    if not row:
        raise HTTPException(status_code=404, detail="Cliente no encontrado")

    return serialize_cliente(row)


@router.patch("/clientes/{cliente_id}", response_model=ClienteOut)
def update_cliente(
    cliente_id: int,
    payload: ClienteUpdate,
    current: Usuario = Depends(require_roles("admin")),
    db: Session = Depends(get_db),
):
    row = db.query(Cliente).filter(Cliente.id == cliente_id).first()

    if not row:
        raise HTTPException(status_code=404, detail="Cliente no encontrado")

    data = payload.model_dump(exclude_unset=True)

    ensure_immutable_code(row.codigo, data.get("codigo"), "código del cliente")
    data.pop("codigo", None)

    for key, value in data.items():
        if isinstance(value, str):
            value = value.strip()
        setattr(row, key, value)

    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Código o identificación de cliente duplicada")

    db.refresh(row)
    return serialize_cliente(row)


@router.post("/maps/resolve", response_model=MapLocationResolveOut)
def resolve_google_maps_location(
    payload: MapLocationResolveIn,
    current: Usuario = Depends(require_roles("admin")),
):
    try:
        location = resolve_map_location(payload.value)
    except MapLocationError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc

    return MapLocationResolveOut(
        latitude=location.latitude,
        longitude=location.longitude,
        canonical_url=location.canonical_url,
        source=location.source,
    )


@router.delete("/clientes/{cliente_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_cliente(
    cliente_id: int,
    current: Usuario = Depends(require_roles("admin")),
    db: Session = Depends(get_db),
):
    row = db.query(Cliente).filter(Cliente.id == cliente_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Cliente no encontrado")
    blockers = []
    checks = [
        (Finca, Finca.cliente_id, "finca(s)"),
        (ReciboCafe, ReciboCafe.cliente_id, "recibo(s)"),
        (CotizacionVenta, CotizacionVenta.cliente_id, "cotización(es)"),
        (SolicitudSalidaVenta, SolicitudSalidaVenta.cliente_id, "salida(s) de venta"),
        (VentaLeadLote, VentaLeadLote.cliente_id, "lead(s) de venta"),
        (SolicitudVenta, SolicitudVenta.cliente_id, "solicitud(es) de venta"),
    ]
    for model, column, label in checks:
        count = db.query(model).filter(column == cliente_id).count()
        if count:
            blockers.append(f"{count} {label}")
    if blockers:
        raise HTTPException(status_code=409, detail="No se puede eliminar el cliente porque tiene relaciones: " + ", ".join(blockers) + ". Desactívelo si debe conservarse el historial.")
    db.delete(row)
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.get("/clientes/{cliente_id}/fincas", response_model=list[FincaOut])
def list_fincas_cliente(
    cliente_id: int,
    current: Usuario = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    cliente = db.query(Cliente).filter(Cliente.id == cliente_id).first()

    if not cliente:
        raise HTTPException(status_code=404, detail="Cliente no encontrado")

    rows = (
        db.query(Finca)
        .filter(Finca.cliente_id == cliente_id)
        .order_by(Finca.nombre.asc())
        .all()
    )

    return [serialize_finca(db, row) for row in rows]


@router.post("/clientes/{cliente_id}/fincas", response_model=FincaOut, status_code=201)
def create_finca_cliente(
    cliente_id: int,
    payload: FincaCreate,
    current: Usuario = Depends(require_roles("admin")),
    db: Session = Depends(get_db),
):
    cliente = db.query(Cliente).filter(Cliente.id == cliente_id, Cliente.activo).first()

    if not cliente:
        raise HTTPException(status_code=404, detail="Cliente no encontrado o inactivo")

    data = payload.model_dump()
    data["cliente_id"] = cliente_id
    data["codigo"] = (data.get("codigo") or "").strip() or next_code(db, Finca, "codigo", "FIN-", 5)
    normalize_finca_map_location(data)

    if not data.get("provincia"):
        data["provincia"] = cliente.provincia
    if not data.get("canton"):
        data["canton"] = cliente.canton
    if not data.get("distrito"):
        data["distrito"] = cliente.distrito
    if not data.get("propietario"):
        data["propietario"] = cliente.nombre_completo

    row = Finca(**data)
    db.add(row)

    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Código de finca duplicado")

    db.refresh(row)
    return serialize_finca(db, row)


@router.get("/fincas", response_model=list[FincaOut])
def list_fincas(
    cliente_id: Optional[int] = None,
    current: Usuario = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    query = db.query(Finca)

    if cliente_id:
        query = query.filter(Finca.cliente_id == cliente_id)

    rows = query.order_by(Finca.nombre.asc()).all()
    return [serialize_finca(db, row) for row in rows]




@router.get("/fincas/{finca_id}", response_model=FincaOut)
def get_finca(
    finca_id: int,
    current: Usuario = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    row = db.query(Finca).filter(Finca.id == finca_id).first()

    if not row:
        raise HTTPException(status_code=404, detail="Finca no encontrada")

    return serialize_finca(db, row)


@router.delete("/fincas/{finca_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_finca(
    finca_id: int,
    current: Usuario = Depends(require_roles("admin")),
    db: Session = Depends(get_db),
):
    row = db.query(Finca).filter(Finca.id == finca_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Finca no encontrada")
    # Las relaciones productivas e IoT deben impedir borrados accidentales.
    from modules.mod_caficultura.model_iot import IoTNode, IoTDashboardPreset
    from modules.mod_caficultura.model_ai_consulting import AIConversation
    blockers = []
    checks = [
        (ReciboCafe, ReciboCafe.finca_id, "recibo(s)"),
        (OrdenTrabajo, OrdenTrabajo.finca_id, "lote(s)/OT"),
        (RegistroFinca, RegistroFinca.finca_id, "registro(s) de finca"),
        (IoTNode, IoTNode.finca_id, "sensor(es) IoT"),
        (IoTDashboardPreset, IoTDashboardPreset.finca_id, "configuración(es) de tablero IoT"),
        (AIConversation, AIConversation.finca_id, "conversación(es) IA"),
    ]
    for model, column, label in checks:
        count = db.query(model).filter(column == finca_id).count()
        if count:
            blockers.append(f"{count} {label}")
    if blockers:
        raise HTTPException(status_code=409, detail="No se puede eliminar la finca porque tiene relaciones: " + ", ".join(blockers) + ". Desactívela para conservar el historial.")
    db.delete(row)
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.get("/gestion-fincas/fincas-disponibles", response_model=list[FincaOut])
def list_fincas_disponibles_gestion(
    current: Usuario = Depends(require_roles("gerente", "administrativo", "supervisor_finca", "admin")),
    db: Session = Depends(get_db),
):
    rows = (
        db.query(Finca)
        .filter(Finca.activa, Finca.gestion_fincas_habilitada)
        .order_by(Finca.nombre.asc())
        .all()
    )
    return [serialize_finca(db, row) for row in rows]

@router.post("/fincas", response_model=FincaOut, status_code=201)
def create_finca(
    payload: FincaCreate,
    current: Usuario = Depends(require_roles("admin")),
    db: Session = Depends(get_db),
):
    if not payload.cliente_id:
        raise HTTPException(status_code=422, detail="Debe indicar el cliente propietario de la finca")

    cliente = db.query(Cliente).filter(Cliente.id == payload.cliente_id, Cliente.activo).first()

    if not cliente:
        raise HTTPException(status_code=404, detail="Cliente no encontrado o inactivo")

    data = payload.model_dump()
    data["codigo"] = (data.get("codigo") or "").strip() or next_code(db, Finca, "codigo", "FIN-", 5)
    normalize_finca_map_location(data)

    if not data.get("provincia"):
        data["provincia"] = cliente.provincia
    if not data.get("canton"):
        data["canton"] = cliente.canton
    if not data.get("distrito"):
        data["distrito"] = cliente.distrito
    if not data.get("propietario"):
        data["propietario"] = cliente.nombre_completo

    row = Finca(**data)
    db.add(row)

    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Código de finca duplicado")

    db.refresh(row)
    return serialize_finca(db, row)


@router.patch("/fincas/{finca_id}", response_model=FincaOut)
def update_finca(
    finca_id: int,
    payload: FincaUpdate,
    current: Usuario = Depends(require_roles("admin")),
    db: Session = Depends(get_db),
):
    row = db.query(Finca).filter(Finca.id == finca_id).first()

    if not row:
        raise HTTPException(status_code=404, detail="Finca no encontrada")

    data = payload.model_dump(exclude_unset=True)
    normalize_finca_map_location(data)

    ensure_immutable_code(row.codigo, data.get("codigo"), "código de la finca")
    data.pop("codigo", None)

    if "cliente_id" in data and data["cliente_id"]:
        cliente = db.query(Cliente).filter(Cliente.id == data["cliente_id"], Cliente.activo).first()
        if not cliente:
            raise HTTPException(status_code=404, detail="Cliente no encontrado o inactivo")

    for key, value in data.items():
        setattr(row, key, value)

    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Código de finca duplicado")

    db.refresh(row)
    return serialize_finca(db, row)






@router.get("/recibos", response_model=list[ReciboOut])
def list_recibos(current: Usuario = Depends(get_current_user), db: Session = Depends(get_db)):
    rows = recibo_query(db).order_by(ReciboCafe.fecha.desc(), ReciboCafe.id.desc()).all()
    return [serialize_recibo(r) for r in rows]


@router.post("/recibos", response_model=ReciboOut, status_code=201)
def create_recibo(payload: ReciboCreate, current: Usuario = Depends(require_roles("gerente", "operario", "administrativo")), db: Session = Depends(get_db)):
    try:
        row = create_recibo_from_payload(db, payload, current)
        db.commit()
    except ValueError as exc:
        db.rollback()
        raise HTTPException(status_code=422, detail=str(exc))
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Número de recibo duplicado")
    row = recibo_query(db).filter(ReciboCafe.id == row.id).first()
    return serialize_recibo(row)


def receipt_counter_out(db: Session, counter) -> ReciboConsecutivoOut:
    last_used = last_receipt_number(db, str(counter.prefijo or ""))
    return ReciboConsecutivoOut(
        prefijo=str(counter.prefijo or ""),
        siguiente_numero=int(counter.siguiente_numero or 1),
        ancho=int(counter.ancho or 5),
        proximo_recibo=preview_receipt_number(db),
        ultimo_numero_usado=last_used,
        updated_at=counter.updated_at,
    )


@router.get("/recibos/consecutivo", response_model=ReciboConsecutivoOut)
def get_receipt_sequence(
    current: Usuario = Depends(require_roles("administrativo", "admin")),
    db: Session = Depends(get_db),
):
    del current
    return receipt_counter_out(db, get_receipt_counter(db))


@router.patch("/recibos/consecutivo", response_model=ReciboConsecutivoOut)
def update_receipt_sequence(
    payload: ReciboConsecutivoUpdate,
    current: Usuario = Depends(require_roles("administrativo", "admin")),
    db: Session = Depends(get_db),
):
    counter = get_receipt_counter(db, lock=True)
    last_used = last_receipt_number(db, payload.prefijo)
    if last_used is not None and payload.siguiente_numero <= last_used:
        raise HTTPException(
            status_code=422,
            detail=(
                f"El siguiente número debe ser mayor que {last_used}, que es el mayor consecutivo "
                f"existente para la serie {payload.prefijo or '(sin prefijo)'}."
            ),
        )
    counter.prefijo = payload.prefijo
    counter.siguiente_numero = payload.siguiente_numero
    counter.ancho = payload.ancho
    counter.updated_by_id = current.id
    db.commit()
    db.refresh(counter)
    return receipt_counter_out(db, counter)


@router.get("/recibos/importacion/plantilla")
def download_receipt_import_template(
    current: Usuario = Depends(require_roles("gerente")),
):
    del current
    return Response(
        content=build_receipt_import_template(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": 'attachment; filename="plantilla-importacion-recibos-historicos.xlsx"',
            "Cache-Control": "private, no-store",
        },
    )


@router.post("/recibos/importacion")
async def import_historical_receipt_file(
    file: UploadFile = File(...),
    current: Usuario = Depends(require_roles("gerente")),
    db: Session = Depends(get_db),
):
    filename = str(file.filename or "").lower()
    if not filename.endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=422, detail="Debe seleccionar un archivo Excel .xlsx")
    content = await file.read(MAX_IMPORT_BYTES + 1)
    try:
        result = import_historical_receipts(db, content, current)
        db.commit()
        return result
    except ReceiptImportValidationError as exc:
        db.rollback()
        raise HTTPException(
            status_code=422,
            detail={
                "message": "No se importó ningún recibo. Corrija las filas indicadas y vuelva a cargar el archivo.",
                "errors": exc.errors,
            },
        ) from exc
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="La importación coincidió con un número de recibo, cliente o finca existente. No se guardó ningún registro.",
        ) from exc


@router.get("/recibos/{recibo_id}", response_model=ReciboOut)
def get_recibo(
    recibo_id: int,
    current: Usuario = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    row = recibo_query(db).filter(ReciboCafe.id == recibo_id).first()

    if not row:
        raise HTTPException(status_code=404, detail="Recibo no encontrado")

    return serialize_recibo(row)


@router.patch("/recibos/{recibo_id}", response_model=ReciboOut)
def update_recibo(recibo_id: int, payload: ReciboUpdate, current: Usuario = Depends(require_roles("gerente", "administrativo")), db: Session = Depends(get_db)):
    row = recibo_query(db).filter(ReciboCafe.id == recibo_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Recibo no encontrado")
    if bool(getattr(row, "liquidado", False)):
        raise HTTPException(
            status_code=409,
            detail="Un recibo liquidado no se puede modificar; su información de pago es auditable",
        )

    data = payload.model_dump(exclude_unset=True)

    ensure_immutable_code(row.numero_recibo, data.get("numero_recibo"), "número de recibo")
    data.pop("numero_recibo", None)

    # Compatibilidad: el campo antiguo de precio promedio en realidad era peso promedio por cajuela.
    if "precio_promedio_cajuela" in data and "peso_promedio_cajuela" not in data:
        data["peso_promedio_cajuela"] = data.get("precio_promedio_cajuela")

    if data.get("peso_promedio_cajuela") is not None:
        data["precio_promedio_cajuela"] = data.get("peso_promedio_cajuela")

    if data.get("precio_fanega") and not data.get("precio_fanega_letras"):
        data["precio_fanega_letras"] = numero_a_letras_es(data.get("precio_fanega"))

    if "beneficio_recibe_usuario_id" in data:
        usuario_id = data.get("beneficio_recibe_usuario_id")

        if usuario_id:
            usuario = db.query(Usuario).filter(Usuario.id == usuario_id, Usuario.activo).first()

            if not usuario:
                raise HTTPException(status_code=422, detail="El responsable de beneficio no existe o está inactivo")

            data["beneficio_recibe"] = usuario.nombre
            data["beneficio_firma_url"] = usuario.firma_url
        else:
            data["beneficio_recibe_usuario_id"] = None
            data["beneficio_firma_url"] = None

    for key, value in data.items():
        setattr(row, key, value)

    if row.precio_fanega and not row.precio_fanega_letras:
        row.precio_fanega_letras = numero_a_letras_es(row.precio_fanega)

    db.commit()
    row = recibo_query(db).filter(ReciboCafe.id == recibo_id).first()
    return serialize_recibo(row)


@router.delete("/recibos/{recibo_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_recibo(
    recibo_id: int,
    current: Usuario = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    ensure_bootstrap_superadmin(current)

    row = db.query(ReciboCafe).filter(ReciboCafe.id == recibo_id).first()

    if not row:
        raise HTTPException(status_code=404, detail="Recibo no encontrado")
    if bool(getattr(row, "liquidado", False)):
        raise HTTPException(
            status_code=409,
            detail="Un recibo liquidado no se puede eliminar; su información de pago es auditable",
        )

    vinculaciones = (
        db.query(OrdenTrabajoRecibo)
        .filter(OrdenTrabajoRecibo.recibo_id == recibo_id)
        .count()
    )

    if vinculaciones:
        raise HTTPException(
            status_code=409,
            detail="No se puede eliminar el recibo porque está vinculado a una o más OT. Elimine primero las OT asociadas o anule el recibo.",
        )

    db.delete(row)
    db.commit()

    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.get("/ot", response_model=list[OTOut])
def list_ot(current: Usuario = Depends(get_current_user), db: Session = Depends(get_db)):
    query = ot_query(db)
    if current.rol == "operario":
        query = query.filter(OrdenTrabajo.operario_id == current.id)
    rows = query.order_by(OrdenTrabajo.fecha_inicio.desc(), OrdenTrabajo.id.desc()).all()
    return [serialize_ot(r) for r in rows]


@router.post("/ot", response_model=OTOut, status_code=201)
def create_ot(payload: OTCreate, current: Usuario = Depends(require_roles("gerente")), db: Session = Depends(get_db)):
    try:
        row = create_ot_from_payload(db, payload, current)
        db.commit()
    except ValueError as exc:
        db.rollback()
        raise HTTPException(status_code=422, detail=str(exc))
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Código de lote duplicado o conflicto de recibos")
    row = ot_query(db).filter(OrdenTrabajo.id == row.id).first()
    return serialize_ot(row)


@router.post("/ot/unir", response_model=OTOut, status_code=201)
def unir_lotes(payload: OTMergePayload, current: Usuario = Depends(require_roles("gerente")), db: Session = Depends(get_db)):
    ids = list(dict.fromkeys([int(x) for x in payload.ot_ids if int(x) > 0]))
    if len(ids) < 2:
        raise HTTPException(status_code=422, detail="Seleccione al menos dos lotes para unir")

    lotes = ot_query(db).filter(OrdenTrabajo.id.in_(ids)).all()
    if len(lotes) != len(ids):
        raise HTTPException(status_code=404, detail="Uno o más lotes no existen")

    bloqueados = [lote.codigo_lote for lote in lotes if lote.estado in {"vendido", "anulada", "unido"}]
    if bloqueados:
        raise HTTPException(status_code=422, detail=f"No se pueden unir lotes vendidos, anulados o ya unidos: {', '.join(bloqueados)}")

    proceso = (payload.proceso or lotes[0].proceso or "miel").strip().lower()
    codigo = (payload.codigo_lote or "").strip() or next_code(db, OrdenTrabajo, "codigo_lote", "LOT-", 5)
    fecha_inicio = payload.fecha_inicio or date.today()
    total_fanegas = round(sum(float(lote.fanegas_estimadas or 0) for lote in lotes), 3)
    qr_token = uuid.uuid4().hex + uuid.uuid4().hex

    nuevo = OrdenTrabajo(
        codigo_lote=codigo,
        proceso=proceso,
        estado="en_proceso",
        fecha_inicio=fecha_inicio,
        finca_id=lotes[0].finca_id,
        operario_id=payload.operario_id if payload.operario_id is not None else lotes[0].operario_id,
        fanegas_estimadas=total_fanegas,
        objetivo_fanegas=total_fanegas,
        qr_token=qr_token,
        qr_public_url=lote_public_url(qr_token),
        observaciones=payload.observaciones or f"Lote unido a partir de: {', '.join(lote.codigo_lote for lote in lotes)}",
        created_by_id=current.id,
    )
    nuevo.qr_payload = build_ot_qr_payload(nuevo)
    db.add(nuevo)
    db.flush()

    motivo = payload.motivo.strip()
    add_lote_evento(db, nuevo.id, "union", f"Lote creado por unión de {', '.join(lote.codigo_lote for lote in lotes)}. Motivo: {motivo}", current)

    recibos_union: dict[int, dict[str, float]] = {}
    for lote in lotes:
        db.add(OrdenTrabajoUnionOrigen(
            ot_destino_id=nuevo.id,
            ot_origen_id=lote.id,
            codigo_lote_origen=lote.codigo_lote,
            estado_origen=lote.estado,
            proceso_origen=lote.proceso,
            fanegas_origen=float(lote.fanegas_estimadas or 0),
            motivo=motivo,
            created_by_id=current.id,
        ))
        add_lote_evento(db, lote.id, "union", f"Lote unido dentro de {nuevo.codigo_lote}. Motivo: {motivo}", current)

        # El lote original conserva sus seguimientos, documentos y comentarios.
        # Solo cambia a estado unido para bloquear operación futura y pintar su avance en azul.
        lote.estado = "unido"

        for rel in lote.recibos_rel or []:
            if not rel.recibo_id:
                continue
            acc = recibos_union.setdefault(rel.recibo_id, {
                "cajuelas": 0.0,
                "cuartillos": 0.0,
                "fanegas": 0.0,
            })
            acc["cajuelas"] += float(rel.cajuelas_asignadas or 0)
            acc["cuartillos"] += float(rel.cuartillos_asignados or 0)
            acc["fanegas"] += float(rel.fanegas_asignadas or 0)

    for recibo_id, cantidades in recibos_union.items():
        db.add(OrdenTrabajoRecibo(
            ot_id=nuevo.id,
            recibo_id=recibo_id,
            cajuelas_asignadas=round(cantidades["cajuelas"], 2),
            cuartillos_asignados=round(cantidades["cuartillos"], 2),
            fanegas_asignadas=round(cantidades["fanegas"], 3),
        ))

    db.commit()
    row = ot_query(db).filter(OrdenTrabajo.id == nuevo.id).first()
    return serialize_ot(row)


@router.get("/ot/recibos-disponibles", response_model=list[ReciboDisponibleOTOut])
def list_recibos_disponibles_ot(
    current: Usuario = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    rows = (
        recibo_query(db)
        .filter(ReciboCafe.estado != "anulado")
        .order_by(ReciboCafe.fecha.asc(), ReciboCafe.id.asc())
        .all()
    )

    disponibles = []

    for row in rows:
        item = serialize_recibo_disponible_ot(db, row)
        if item.cajuelas_disponibles > 0:
            disponibles.append(item)

    return disponibles


@router.get("/ot/lotes/{codigo_lote}", response_model=OTOut)
def get_ot_by_codigo_lote(
    codigo_lote: str,
    current: Usuario = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    row = (
        ot_query(db)
        .filter(OrdenTrabajo.codigo_lote == codigo_lote)
        .first()
    )

    if not row:
        raise HTTPException(status_code=404, detail="Lote no encontrado")

    return serialize_ot(row)


@router.post("/ot/{ot_id}/solicitar-finalizacion", response_model=OTOut)
def solicitar_finalizacion_ot(
    ot_id: int,
    payload: EstadoOTPayload,
    current: Usuario = Depends(require_roles("operario", "gerente")),
    db: Session = Depends(get_db),
):
    row = db.query(OrdenTrabajo).filter(OrdenTrabajo.id == ot_id).first()

    if not row:
        raise HTTPException(status_code=404, detail="OT no encontrada")
    ensure_lote_not_unido(row, "solicitar aprobación")

    if current.rol == "operario" and row.operario_id != current.id:
        raise HTTPException(status_code=403, detail="Solo el operario asignado puede solicitar finalización")

    row.estado = "pendiente_aprobacion"

    nota = payload.comentario or "Solicitud de finalización enviada para revisión física del gerente."
    row.observaciones = ((row.observaciones or "") + f"\n\n[Solicitud de finalización] {nota}").strip()
    add_lote_evento(db, row.id, "estado", f"Solicitud de aprobación enviada. {nota}", current)

    db.commit()
    db.refresh(row)

    return serialize_ot(row)


@router.post("/ot/{ot_id}/aprobar-finalizacion", response_model=OTOut)
def aprobar_finalizacion_ot(
    ot_id: int,
    payload: EstadoOTPayload,
    current: Usuario = Depends(require_roles("gerente")),
    db: Session = Depends(get_db),
):
    row = db.query(OrdenTrabajo).filter(OrdenTrabajo.id == ot_id).first()

    if not row:
        raise HTTPException(status_code=404, detail="OT no encontrada")
    ensure_lote_not_unido(row, "aprobarlo")

    row.estado = "finalizada"
    row.gerente_cierra_id = current.id
    row.fecha_cierre = utcnow()

    nota = payload.comentario or "Finalización aprobada por gerencia después de revisión."
    row.observaciones = ((row.observaciones or "") + f"\n\n[Aprobación gerente] {nota}").strip()
    add_lote_evento(db, row.id, "supervision", f"Lote aprobado por gerencia. {nota}", current)

    db.commit()
    db.refresh(row)

    return serialize_ot(row)


@router.post("/ot/{ot_id}/devolver-finalizacion", response_model=OTOut)
def devolver_finalizacion_ot(
    ot_id: int,
    payload: EstadoOTPayload,
    current: Usuario = Depends(require_roles("gerente")),
    db: Session = Depends(get_db),
):
    row = db.query(OrdenTrabajo).filter(OrdenTrabajo.id == ot_id).first()

    if not row:
        raise HTTPException(status_code=404, detail="OT no encontrada")
    ensure_lote_not_unido(row, "devolverlo a reproceso")

    row.estado = "en_proceso"

    nota = payload.comentario or "Finalización devuelta para corrección o revisión adicional."
    row.observaciones = ((row.observaciones or "") + f"\n\n[Devuelto por gerencia] {nota}").strip()
    add_lote_evento(db, row.id, "supervision", f"Lote devuelto a reproceso. {nota}", current)

    db.commit()
    db.refresh(row)

    return serialize_ot(row)


@router.post("/seguimientos/{seguimiento_id}/revision")
def revisar_seguimiento_ot(
    seguimiento_id: int,
    payload: SeguimientoRevisionPayload,
    current: Usuario = Depends(require_roles("gerente")),
    db: Session = Depends(get_db),
):
    row = db.query(SeguimientoOT).filter(SeguimientoOT.id == seguimiento_id).first()

    if not row:
        raise HTTPException(status_code=404, detail="Seguimiento no encontrado")

    row.estado_revision = payload.estado_revision
    row.comentario_revision = payload.comentario_revision
    row.revisado_por_id = current.id
    row.revisado_at = utcnow()

    db.commit()

    return {"ok": True}


@router.post("/ot/{ot_id}/resolver-alerta", response_model=OTOut)
def resolver_alerta_ot(
    ot_id: int,
    payload: EstadoOTPayload,
    current: Usuario = Depends(require_roles("gerente")),
    db: Session = Depends(get_db),
):
    row = db.query(OrdenTrabajo).filter(OrdenTrabajo.id == ot_id).first()

    if not row:
        raise HTTPException(status_code=404, detail="OT no encontrada")
    ensure_lote_not_unido(row, "quitar alertas")

    alertas = db.query(SeguimientoOT).filter(
        SeguimientoOT.orden_trabajo_id == ot_id,
        SeguimientoOT.dar_alerta,
        SeguimientoOT.estado_revision == "pendiente",
    ).all()

    comentario = (payload.comentario or "Alerta revisada y retirada por gerencia.").strip()

    for alerta in alertas:
        alerta.estado_revision = "aprobado"
        alerta.comentario_revision = comentario
        alerta.revisado_por_id = current.id
        alerta.revisado_at = utcnow()

    if row.estado == "alerta":
        row.estado = "en_proceso"

    add_lote_evento(db, row.id, "supervision", f"Gerencia quitó la alerta del lote. {comentario}", current)

    db.commit()
    row = ot_query(db).filter(OrdenTrabajo.id == ot_id).first()
    return serialize_ot(row)


@router.patch("/ot/{ot_id}", response_model=OTOut)
def update_ot(ot_id: int, payload: OTUpdate, current: Usuario = Depends(require_roles("gerente")), db: Session = Depends(get_db)):
    row = db.query(OrdenTrabajo).filter(OrdenTrabajo.id == ot_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="OT no encontrada")
    ensure_lote_not_unido(row, "actualizarlo")
    data = payload.model_dump(exclude_unset=True)
    estado_anterior = row.estado
    operario_anterior = row.operario_id
    comentario_gerencial = (data.pop("comentario_gerencial", None) or "").strip() if isinstance(data, dict) else ""

    for key, value in data.items():
        setattr(row, key, value)

    if data.get("estado") == "finalizada":
        row.gerente_cierra_id = current.id
        row.fecha_cierre = utcnow()

    eventos = []
    if "estado" in data and data.get("estado") != estado_anterior:
        eventos.append(f"Estado cambiado de {estado_anterior} a {data.get('estado')}.")
    if "operario_id" in data and data.get("operario_id") != operario_anterior:
        eventos.append("Operario asignado actualizado por gerencia.")
    if comentario_gerencial:
        eventos.append(comentario_gerencial)
    if eventos:
        add_lote_evento(db, row.id, "estado", " ".join(eventos), current)

    db.commit()
    row = ot_query(db).filter(OrdenTrabajo.id == ot_id).first()
    return serialize_ot(row)


@router.delete("/ot/{ot_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_ot(
    ot_id: int,
    current: Usuario = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    ensure_bootstrap_superadmin(current)

    row = ot_query(db).filter(OrdenTrabajo.id == ot_id).first()

    if not row:
        raise HTTPException(status_code=404, detail="OT no encontrada")

    dependencias = []

    cotizaciones = db.query(CotizacionVentaLinea).filter(CotizacionVentaLinea.ot_id == ot_id).count()
    if cotizaciones:
        dependencias.append(f"{cotizaciones} línea(s) de cotización comercial")

    salidas = db.query(SolicitudSalidaVentaLinea).filter(SolicitudSalidaVentaLinea.ot_id == ot_id).count()
    if salidas:
        dependencias.append(f"{salidas} línea(s) de solicitud de salida")

    solicitudes = db.query(SolicitudVenta).filter(SolicitudVenta.ot_id == ot_id).count()
    if solicitudes:
        dependencias.append(f"{solicitudes} solicitud(es) de venta")

    liquidaciones = db.query(SolicitudVentaLineaLiquidacion).filter(SolicitudVentaLineaLiquidacion.ot_id == ot_id).count()
    if liquidaciones:
        dependencias.append(f"{liquidaciones} liquidación(es) comerciales")

    uniones_destino = db.query(OrdenTrabajoUnionOrigen).filter(OrdenTrabajoUnionOrigen.ot_destino_id == ot_id).count()
    uniones_origen = db.query(OrdenTrabajoUnionOrigen).filter(OrdenTrabajoUnionOrigen.ot_origen_id == ot_id).count()
    if uniones_destino or uniones_origen:
        dependencias.append("trazabilidad de unión de lotes")

    if dependencias:
        raise HTTPException(
            status_code=409,
            detail="No se puede eliminar el lote porque tiene asociaciones que romperían la trazabilidad: " + "; ".join(dependencias),
        )

    recibo_ids = [rel.recibo_id for rel in (row.recibos_rel or []) if rel.recibo_id]

    db.delete(row)
    db.flush()

    for recibo_id in recibo_ids:
        recibo = db.query(ReciboCafe).filter(ReciboCafe.id == recibo_id).first()

        if not recibo or recibo.estado == "anulado":
            continue

        disponible = recibo_disponible_cajuelas(db, recibo)
        total = float(recibo.cajuelas or 0) + float(recibo.cuartillos or 0) / 4

        recibo.estado = "en_proceso" if disponible <= 0 else "recibido"

        if total <= 0:
            recibo.estado = "recibido"

    db.commit()

    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.post("/ot/{ot_id}/seguimientos", response_model=OTOut, status_code=201)
def create_seguimiento(ot_id: int, payload: SeguimientoCreate, current: Usuario = Depends(require_roles("gerente", "operario")), db: Session = Depends(get_db)):
    row = db.query(OrdenTrabajo).filter(OrdenTrabajo.id == ot_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="OT no encontrada")
    ensure_lote_not_unido(row, "registrar seguimientos")
    if current.rol == "operario" and row.operario_id != current.id:
        raise HTTPException(status_code=403, detail="Solo puedes registrar seguimiento en tus OT asignadas")
    if payload.client_uuid:
        existing = db.query(SeguimientoOT).filter(SeguimientoOT.client_uuid == payload.client_uuid).first()
        if existing:
            return serialize_ot(ot_query(db).filter(OrdenTrabajo.id == existing.orden_trabajo_id).first())
    marca = payload.marca_tiempo or utcnow()
    estado_resultante = payload.estado_lote_resultante
    alerta_mensaje = (payload.alerta_mensaje or payload.comentario or "").strip() or None

    alerta_whatsapp_url = None
    alerta_correo_url = None
    if payload.dar_alerta:
        mensaje = (
            f"ALERTA NAVIA · Lote {row.codigo_lote} · Actividad: {payload.actividad_realizada.strip()} · "
            f"Reporta: {current.nombre}. {alerta_mensaje or ''}"
        ).strip()
        alerta_whatsapp_url, alerta_correo_url = build_alert_links(db, mensaje)
        estado_resultante = "alerta"

    seguimiento = SeguimientoOT(
        orden_trabajo_id=row.id,
        client_uuid=payload.client_uuid,
        fecha=payload.fecha,
        marca_tiempo=marca,
        actividad_realizada=payload.actividad_realizada.strip(),
        horas_implementadas=payload.horas_implementadas,
        temperatura=payload.temperatura,
        humedad=payload.humedad,
        comentario=payload.comentario,
        estado_lote_resultante=estado_resultante,
        dar_alerta=bool(payload.dar_alerta),
        alerta_mensaje=alerta_mensaje,
        alerta_canal=payload.alerta_canal or "whatsapp",
        alerta_whatsapp_url=alerta_whatsapp_url,
        alerta_correo_url=alerta_correo_url,
        created_by_id=current.id,
    )
    db.add(seguimiento)
    tipo_evento = "alerta" if payload.dar_alerta else "seguimiento"
    add_lote_evento(
        db,
        row.id,
        tipo_evento,
        f"{current.nombre} registró {tipo_evento}: {seguimiento.actividad_realizada}. {seguimiento.comentario or alerta_mensaje or ''}".strip(),
        current,
    )

    if estado_resultante:
        estado_anterior = row.estado
        row.estado = estado_resultante
        if estado_anterior != row.estado:
            add_lote_evento(db, row.id, "estado", f"El lote pasó de {estado_anterior} a {row.estado} por registro operativo.", current)
    elif row.estado == "abierta":
        row.estado = "en_proceso"
        add_lote_evento(db, row.id, "estado", "El lote pasó automáticamente a estado en proceso por registro de seguimiento.", current)
    db.commit()
    row = ot_query(db).filter(OrdenTrabajo.id == ot_id).first()
    return serialize_ot(row)



def is_superadmin(current: Usuario) -> bool:
    return bool(getattr(current, "is_superadmin", False))


def require_superadmin_user(current: Usuario) -> None:
    if not is_superadmin(current):
        raise HTTPException(status_code=403, detail="Solo superadmin puede eliminar registros")


def semana_operativa(fecha: date) -> tuple[date, date]:
    inicio = fecha - timedelta(days=fecha.weekday())
    fin = inicio + timedelta(days=5)
    return inicio, fin


def semana_label(inicio: date, fin: date) -> str:
    meses = [
        "enero", "febrero", "marzo", "abril", "mayo", "junio",
        "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre",
    ]
    if inicio.year == fin.year and inicio.month == fin.month:
        return f"{inicio.day} a {fin.day} {meses[inicio.month - 1]} {inicio.year}"
    return f"{inicio.day} {meses[inicio.month - 1]} {inicio.year} a {fin.day} {meses[fin.month - 1]} {fin.year}"


def parse_bool(value, default: bool = True) -> bool:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    raw = str(value).strip().lower()
    return raw in {"1", "si", "sí", "true", "activo", "activa", "x", "yes"}


def parse_float(value, default=None):
    if value is None or value == "":
        return default
    try:
        return float(str(value).replace(",", "."))
    except Exception:
        return default


def serialize_actividad_finca(row: ActividadFinca) -> ActividadFincaOut:
    return ActividadFincaOut(
        id=row.id,
        nombre=row.nombre,
        tipo=row.tipo,
        descripcion=row.descripcion,
        requiere_insumo=bool(row.requiere_insumo),
        unidad_referencia=row.unidad_referencia,
        activa=bool(row.activa),
        created_at=row.created_at,
    )


def serialize_trabajador_finca(row: TrabajadorFinca) -> TrabajadorFincaOut:
    return TrabajadorFincaOut(
        id=row.id,
        codigo=row.codigo,
        nombre=row.nombre,
        identificacion=row.identificacion,
        telefono=row.telefono,
        puesto=row.puesto,
        jornal_diario=row.jornal_diario,
        activo=bool(row.activo),
        created_at=row.created_at,
    )


def serialize_insumo_finca(row: InsumoFinca) -> InsumoFincaOut:
    return InsumoFincaOut(
        id=row.id,
        codigo=row.codigo,
        codigo_fabricante=getattr(row, "codigo_fabricante", None),
        nombre=row.nombre,
        tipo=row.tipo,
        unidad=row.unidad,
        costo_unitario=row.costo_unitario,
        precio_unitario=getattr(row, "precio_unitario", None) if getattr(row, "precio_unitario", None) is not None else row.costo_unitario,
        impuesto_porcentaje=float(getattr(row, "impuesto_porcentaje", 0) or 0),
        stock_actual=float(getattr(row, "stock_actual", 0) or 0),
        stock_minimo=getattr(row, "stock_minimo", None),
        observaciones=row.observaciones,
        activo=bool(row.activo),
        created_at=row.created_at,
    )


def serialize_proveedor(row: Proveedor) -> ProveedorOut:
    return ProveedorOut(
        id=row.id,
        codigo=row.codigo,
        nombre_legal=row.nombre_legal,
        nombre_comercial=row.nombre_comercial,
        identificacion=row.identificacion,
        tipo_identificacion=row.tipo_identificacion,
        correo=row.correo,
        telefono=row.telefono,
        sitio_web=row.sitio_web,
        direccion=row.direccion,
        provincia=row.provincia,
        canton=row.canton,
        distrito=row.distrito,
        notas=row.notas,
        activo=bool(row.activo),
        contactos=[{
            "id": c.id, "nombre": c.nombre, "puesto": c.puesto, "correo": c.correo,
            "telefono": c.telefono, "principal": bool(c.principal), "activo": bool(c.activo),
            "notas": c.notas, "created_at": c.created_at,
        } for c in (row.contactos or [])],
        facturas_count=len(row.facturas or []),
        created_at=row.created_at,
        updated_at=row.updated_at,
    )


def _replace_provider_contacts(row: Proveedor, contactos) -> None:
    row.contactos.clear()
    for item in contactos or []:
        data = item.model_dump() if hasattr(item, "model_dump") else dict(item)
        row.contactos.append(ProveedorContacto(**data))


def serialize_compra_insumo_linea(row: CompraInsumoLinea) -> dict:
    return {
        "id": row.id,
        "insumo_id": row.insumo_id,
        "codigo_fabricante": row.codigo_fabricante,
        "nombre": row.nombre_snapshot,
        "cantidad": float(row.cantidad or 0),
        "unidad": row.unidad,
        "precio_unitario": float(row.precio_unitario or 0),
        "impuesto_porcentaje": float(row.impuesto_porcentaje or 0),
        "subtotal": float(row.subtotal or 0),
        "impuesto": float(row.impuesto or 0),
        "total": float(row.total or 0),
    }


def serialize_compra_insumo(row: CompraInsumoFactura) -> CompraInsumoOut:
    return CompraInsumoOut(
        id=row.id,
        numero_factura=row.numero_factura,
        proveedor_id=getattr(row, "proveedor_id", None),
        proveedor=(row.proveedor_rel.nombre_comercial or row.proveedor_rel.nombre_legal) if getattr(row, "proveedor_rel", None) else row.proveedor,
        fecha=row.fecha,
        moneda=row.moneda,
        observaciones=row.observaciones,
        subtotal=float(row.subtotal or 0),
        impuesto=float(row.impuesto or 0),
        total=float(row.total or 0),
        lineas=[serialize_compra_insumo_linea(linea) for linea in (row.lineas or [])],
        created_by_nombre=row.created_by.nombre if row.created_by else None,
        created_at=row.created_at,
    )


def serialize_registro_finca(row: RegistroFinca) -> RegistroFincaOut:
    trabajadores = []
    for item in row.trabajadores or []:
        trabajadores.append({
            "id": item.id,
            "trabajador_id": item.trabajador_id,
            "nombre": item.nombre_snapshot,
            "horas": float(item.horas or 0),
            "jornal": float(item.jornal or 0) if item.jornal is not None else None,
            "costo": float(item.costo or 0),
        })

    insumos = []
    for item in row.insumos or []:
        insumos.append({
            "id": item.id,
            "insumo_id": item.insumo_id,
            "nombre": item.nombre_snapshot,
            "cantidad": float(item.cantidad or 0),
            "unidad": item.unidad,
            "costo_unitario": float(item.costo_unitario or 0) if item.costo_unitario is not None else None,
            "costo_total": float(item.costo_total or 0),
            "comentario": item.comentario,
        })

    total_horas = round(sum(float(t.get("horas") or 0) for t in trabajadores), 2)
    costo_mano_obra = round(sum(float(t.get("costo") or 0) for t in trabajadores), 2)
    costo_insumos = round(sum(float(i.get("costo_total") or 0) for i in insumos), 2)

    return RegistroFincaOut(
        id=row.id,
        fecha=row.fecha,
        semana_inicio=row.semana_inicio,
        semana_fin=row.semana_fin,
        semana_label=semana_label(row.semana_inicio, row.semana_fin),
        finca_id=row.finca_id,
        finca_nombre=row.finca.nombre if row.finca else None,
        cliente_nombre=row.finca.cliente.nombre_completo if row.finca and row.finca.cliente else None,
        actividad_id=row.actividad_id,
        actividad_nombre=row.actividad.nombre if row.actividad else None,
        actividad_tipo=row.actividad.tipo if row.actividad else None,
        descripcion=row.descripcion,
        estado=row.estado,
        observaciones=row.observaciones,
        trabajadores=trabajadores,
        insumos=insumos,
        total_trabajadores=len(trabajadores),
        total_horas=total_horas,
        costo_mano_obra=costo_mano_obra,
        costo_insumos=costo_insumos,
        costo_total=round(costo_mano_obra + costo_insumos, 2),
        created_by_id=row.created_by_id,
        created_by_nombre=row.created_by.nombre if row.created_by else None,
        created_at=row.created_at,
    )


def registro_finca_query(db: Session):
    return db.query(RegistroFinca).options(
        selectinload(RegistroFinca.finca).selectinload(Finca.cliente),
        selectinload(RegistroFinca.actividad),
        selectinload(RegistroFinca.created_by),
        selectinload(RegistroFinca.trabajadores).selectinload(RegistroFincaTrabajador.trabajador),
        selectinload(RegistroFinca.insumos).selectinload(RegistroFincaInsumo.insumo),
    )



def restaurar_stock_registro(db: Session, row: RegistroFinca) -> None:
    for item in row.insumos or []:
        if item.insumo_id:
            insumo = db.query(InsumoFinca).filter(InsumoFinca.id == item.insumo_id).first()
            if insumo:
                insumo.stock_actual = round(float(getattr(insumo, "stock_actual", 0) or 0) + float(item.cantidad or 0), 4)


def descontar_stock_insumo(insumo: InsumoFinca, cantidad: float) -> None:
    disponible = float(getattr(insumo, "stock_actual", 0) or 0)
    if cantidad > disponible:
        raise HTTPException(
            status_code=422,
            detail=f"Stock insuficiente para {insumo.nombre}. Disponible: {round(disponible, 2)} {insumo.unidad}; solicitado: {round(cantidad, 2)}.",
        )
    insumo.stock_actual = round(disponible - cantidad, 4)

def apply_registro_finca_items(db: Session, row: RegistroFinca, payload) -> None:
    if getattr(payload, "trabajadores", None) is not None:
        row.trabajadores.clear()
        db.flush()
        for item in payload.trabajadores or []:
            trabajador = db.query(TrabajadorFinca).filter(TrabajadorFinca.id == item.trabajador_id).first()
            if not trabajador:
                raise HTTPException(status_code=404, detail=f"Trabajador {item.trabajador_id} no encontrado")
            jornal = item.jornal if item.jornal is not None else trabajador.jornal_diario
            horas = float(item.horas or 0)
            costo = round((float(jornal or 0) / 8) * horas, 2) if jornal is not None else 0
            row.trabajadores.append(RegistroFincaTrabajador(
                trabajador_id=trabajador.id,
                nombre_snapshot=trabajador.nombre,
                horas=horas,
                jornal=jornal,
                costo=costo,
            ))

    if getattr(payload, "insumos", None) is not None:
        restaurar_stock_registro(db, row)
        row.insumos.clear()
        db.flush()
        for item in payload.insumos or []:
            insumo = db.query(InsumoFinca).filter(InsumoFinca.id == item.insumo_id).first()
            if not insumo:
                raise HTTPException(status_code=404, detail=f"Insumo {item.insumo_id} no encontrado")
            costo_unitario = item.costo_unitario if item.costo_unitario is not None else (getattr(insumo, "precio_unitario", None) if getattr(insumo, "precio_unitario", None) is not None else insumo.costo_unitario)
            cantidad = float(item.cantidad or 0)
            descontar_stock_insumo(insumo, cantidad)
            row.insumos.append(RegistroFincaInsumo(
                insumo_id=insumo.id,
                nombre_snapshot=insumo.nombre,
                cantidad=cantidad,
                unidad=item.unidad or insumo.unidad,
                costo_unitario=costo_unitario,
                costo_total=round(cantidad * float(costo_unitario or 0), 2),
                comentario=item.comentario,
            ))


@router.get("/gestion-fincas/actividades", response_model=list[ActividadFincaOut])
def list_actividades_finca(current: Usuario = Depends(require_roles("administrativo", "supervisor_finca", "admin")), db: Session = Depends(get_db)):
    rows = db.query(ActividadFinca).order_by(ActividadFinca.activa.desc(), ActividadFinca.nombre.asc()).all()
    return [serialize_actividad_finca(row) for row in rows]


@router.post("/gestion-fincas/actividades", response_model=ActividadFincaOut, status_code=201)
def create_actividad_finca(payload: ActividadFincaCreate, current: Usuario = Depends(require_roles("administrativo", "admin")), db: Session = Depends(get_db)):
    row = ActividadFinca(**payload.model_dump())
    db.add(row)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Ya existe una actividad con ese nombre")
    db.refresh(row)
    return serialize_actividad_finca(row)


@router.patch("/gestion-fincas/actividades/{actividad_id}", response_model=ActividadFincaOut)
def update_actividad_finca(actividad_id: int, payload: ActividadFincaUpdate, current: Usuario = Depends(require_roles("administrativo", "admin")), db: Session = Depends(get_db)):
    row = db.query(ActividadFinca).filter(ActividadFinca.id == actividad_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Actividad no encontrada")
    for key, value in payload.model_dump(exclude_unset=True).items():
        setattr(row, key, value.strip() if isinstance(value, str) else value)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Ya existe una actividad con ese nombre")
    db.refresh(row)
    return serialize_actividad_finca(row)


@router.delete("/gestion-fincas/actividades/{actividad_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_actividad_finca(actividad_id: int, current: Usuario = Depends(get_current_user), db: Session = Depends(get_db)):
    require_superadmin_user(current)
    row = db.query(ActividadFinca).filter(ActividadFinca.id == actividad_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Actividad no encontrada")
    usados = db.query(RegistroFinca).filter(RegistroFinca.actividad_id == actividad_id).count()
    if usados:
        raise HTTPException(status_code=400, detail="La actividad ya tiene registros. Desactívela en lugar de eliminarla.")
    db.delete(row)
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.get("/gestion-fincas/trabajadores", response_model=list[TrabajadorFincaOut])
def list_trabajadores_finca(current: Usuario = Depends(require_roles("administrativo", "supervisor_finca", "admin")), db: Session = Depends(get_db)):
    rows = db.query(TrabajadorFinca).order_by(TrabajadorFinca.activo.desc(), TrabajadorFinca.nombre.asc()).all()
    return [serialize_trabajador_finca(row) for row in rows]


@router.post("/gestion-fincas/trabajadores", response_model=TrabajadorFincaOut, status_code=201)
def create_trabajador_finca(payload: TrabajadorFincaCreate, current: Usuario = Depends(require_roles("administrativo", "admin")), db: Session = Depends(get_db)):
    data = payload.model_dump()
    data["codigo"] = (data.get("codigo") or "").strip() or next_code(db, TrabajadorFinca, "codigo", "TR-", 3)
    row = TrabajadorFinca(**data)
    db.add(row)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Código de trabajador duplicado")
    db.refresh(row)
    return serialize_trabajador_finca(row)


@router.patch("/gestion-fincas/trabajadores/{trabajador_id}", response_model=TrabajadorFincaOut)
def update_trabajador_finca(trabajador_id: int, payload: TrabajadorFincaUpdate, current: Usuario = Depends(require_roles("administrativo", "admin")), db: Session = Depends(get_db)):
    row = db.query(TrabajadorFinca).filter(TrabajadorFinca.id == trabajador_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Trabajador no encontrado")
    data = payload.model_dump(exclude_unset=True)
    ensure_immutable_code(row.codigo, data.get("codigo"), "código del trabajador")
    data.pop("codigo", None)
    for key, value in data.items():
        setattr(row, key, value.strip() if isinstance(value, str) else value)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Código de trabajador duplicado")
    db.refresh(row)
    return serialize_trabajador_finca(row)


@router.delete("/gestion-fincas/trabajadores/{trabajador_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_trabajador_finca(trabajador_id: int, current: Usuario = Depends(get_current_user), db: Session = Depends(get_db)):
    require_superadmin_user(current)
    row = db.query(TrabajadorFinca).filter(TrabajadorFinca.id == trabajador_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Trabajador no encontrado")
    usados = db.query(RegistroFincaTrabajador).filter(RegistroFincaTrabajador.trabajador_id == trabajador_id).count()
    if usados:
        raise HTTPException(status_code=400, detail="El trabajador ya tiene registros. Desactívelo en lugar de eliminarlo.")
    db.delete(row)
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.get("/proveedores", response_model=list[ProveedorOut])
def list_proveedores(
    activos: Optional[bool] = None,
    current: Usuario = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    query = db.query(Proveedor).options(selectinload(Proveedor.contactos), selectinload(Proveedor.facturas))
    if activos is not None:
        query = query.filter(Proveedor.activo == activos)
    rows = query.order_by(Proveedor.activo.desc(), Proveedor.nombre_comercial.asc().nullslast(), Proveedor.nombre_legal.asc()).all()
    return [serialize_proveedor(row) for row in rows]


@router.post("/proveedores", response_model=ProveedorOut, status_code=201)
def create_proveedor(
    payload: ProveedorCreate,
    current: Usuario = Depends(require_roles("administrativo", "admin")),
    db: Session = Depends(get_db),
):
    data = payload.model_dump(exclude={"contactos"})
    data["codigo"] = (data.get("codigo") or "").strip() or next_code(db, Proveedor, "codigo", "PROV-", 4)
    for key, value in list(data.items()):
        if isinstance(value, str):
            data[key] = value.strip() or None
    row = Proveedor(**data)
    _replace_provider_contacts(row, payload.contactos)
    db.add(row)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Código o identificación de proveedor duplicada")
    row = db.query(Proveedor).options(selectinload(Proveedor.contactos), selectinload(Proveedor.facturas)).filter(Proveedor.id == row.id).first()
    return serialize_proveedor(row)


@router.get("/proveedores/{proveedor_id}", response_model=ProveedorOut)
def get_proveedor(proveedor_id: int, current: Usuario = Depends(get_current_user), db: Session = Depends(get_db)):
    row = db.query(Proveedor).options(selectinload(Proveedor.contactos), selectinload(Proveedor.facturas)).filter(Proveedor.id == proveedor_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Proveedor no encontrado")
    return serialize_proveedor(row)


@router.patch("/proveedores/{proveedor_id}", response_model=ProveedorOut)
def update_proveedor(
    proveedor_id: int, payload: ProveedorUpdate,
    current: Usuario = Depends(require_roles("administrativo", "admin")),
    db: Session = Depends(get_db),
):
    row = db.query(Proveedor).options(selectinload(Proveedor.contactos), selectinload(Proveedor.facturas)).filter(Proveedor.id == proveedor_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Proveedor no encontrado")
    data = payload.model_dump(exclude_unset=True)
    contactos = data.pop("contactos", None)
    ensure_immutable_code(row.codigo, data.get("codigo"), "código del proveedor")
    data.pop("codigo", None)
    for key, value in data.items():
        if isinstance(value, str):
            value = value.strip() or None
        setattr(row, key, value)
    if contactos is not None:
        _replace_provider_contacts(row, contactos)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Identificación de proveedor duplicada")
    row = db.query(Proveedor).options(selectinload(Proveedor.contactos), selectinload(Proveedor.facturas)).filter(Proveedor.id == proveedor_id).first()
    return serialize_proveedor(row)


@router.delete("/proveedores/{proveedor_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_proveedor(
    proveedor_id: int,
    current: Usuario = Depends(require_roles("administrativo", "admin")),
    db: Session = Depends(get_db),
):
    row = db.query(Proveedor).filter(Proveedor.id == proveedor_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Proveedor no encontrado")
    facturas = db.query(CompraInsumoFactura).filter(CompraInsumoFactura.proveedor_id == proveedor_id).count()
    if facturas:
        raise HTTPException(status_code=409, detail=f"No se puede eliminar: el proveedor tiene {facturas} factura(s) relacionada(s). Desactívelo.")
    db.delete(row)
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.get("/insumos/compras", response_model=list[CompraInsumoOut])
def list_compras_insumos(current: Usuario = Depends(get_current_user), db: Session = Depends(get_db)):
    rows = db.query(CompraInsumoFactura).options(
        selectinload(CompraInsumoFactura.lineas),
        selectinload(CompraInsumoFactura.created_by),
        selectinload(CompraInsumoFactura.proveedor_rel),
    ).order_by(CompraInsumoFactura.fecha.desc(), CompraInsumoFactura.id.desc()).all()
    return [serialize_compra_insumo(row) for row in rows]


@router.get("/insumos/compras/{factura_id}", response_model=CompraInsumoOut)
def get_compra_insumo(factura_id: int, current: Usuario = Depends(get_current_user), db: Session = Depends(get_db)):
    row = db.query(CompraInsumoFactura).options(
        selectinload(CompraInsumoFactura.lineas),
        selectinload(CompraInsumoFactura.created_by),
        selectinload(CompraInsumoFactura.proveedor_rel),
    ).filter(CompraInsumoFactura.id == factura_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Factura no encontrada")
    return serialize_compra_insumo(row)


@router.post("/insumos/compras", response_model=CompraInsumoOut, status_code=201)
def create_compra_insumos(payload: CompraInsumoCreate, current: Usuario = Depends(require_roles("administrativo", "admin")), db: Session = Depends(get_db)):
    proveedor = None
    if payload.proveedor_id:
        proveedor = db.query(Proveedor).filter(Proveedor.id == payload.proveedor_id, Proveedor.activo).first()
        if not proveedor:
            raise HTTPException(status_code=404, detail="Proveedor no encontrado o inactivo")
    proveedor_snapshot = (proveedor.nombre_comercial or proveedor.nombre_legal) if proveedor else ((payload.proveedor or "").strip() or None)
    factura = CompraInsumoFactura(
        numero_factura=(payload.numero_factura or "").strip() or None,
        proveedor_id=proveedor.id if proveedor else None,
        proveedor=proveedor_snapshot,
        fecha=payload.fecha,
        moneda=(payload.moneda or "CRC").strip().upper(),
        observaciones=payload.observaciones,
        created_by_id=current.id,
    )
    db.add(factura)
    db.flush()

    subtotal_total = 0.0
    impuesto_total = 0.0

    for item in payload.lineas:
        insumo = None
        if item.insumo_id:
            insumo = db.query(InsumoFinca).filter(InsumoFinca.id == item.insumo_id).first()
            if not insumo:
                raise HTTPException(status_code=404, detail=f"Insumo {item.insumo_id} no encontrado")
        else:
            codigo = (item.codigo or "").strip()
            if codigo:
                insumo = db.query(InsumoFinca).filter(InsumoFinca.codigo == codigo).first()
            if not insumo:
                insumo = db.query(InsumoFinca).filter(func.lower(InsumoFinca.nombre) == item.nombre.strip().lower()).first()
            if not insumo:
                insumo = InsumoFinca(
                    codigo=codigo or next_code(db, InsumoFinca, "codigo", "INS-", 3),
                    codigo_fabricante=(item.codigo_fabricante or "").strip() or None,
                    nombre=item.nombre.strip(),
                    tipo=item.tipo or "general",
                    unidad=item.unidad or "unidad",
                    costo_unitario=float(item.precio_unitario or 0),
                    precio_unitario=float(item.precio_unitario or 0),
                    impuesto_porcentaje=float(item.impuesto_porcentaje or 0),
                    stock_actual=0,
                    activo=True,
                )
                db.add(insumo)
                db.flush()

        cantidad = float(item.cantidad or 0)
        precio = float(item.precio_unitario or 0)
        impuesto_pct = float(item.impuesto_porcentaje or 0)
        subtotal = round(cantidad * precio, 2)
        impuesto = round(subtotal * (impuesto_pct / 100), 2)
        total = round(subtotal + impuesto, 2)

        insumo.codigo_fabricante = (item.codigo_fabricante or getattr(insumo, "codigo_fabricante", None) or "").strip() or None
        insumo.tipo = item.tipo or insumo.tipo or "general"
        insumo.unidad = item.unidad or insumo.unidad or "unidad"
        insumo.costo_unitario = precio
        insumo.precio_unitario = precio
        insumo.impuesto_porcentaje = impuesto_pct
        insumo.stock_actual = round(float(getattr(insumo, "stock_actual", 0) or 0) + cantidad, 4)

        db.add(CompraInsumoLinea(
            factura_id=factura.id,
            insumo_id=insumo.id,
            codigo_fabricante=insumo.codigo_fabricante,
            nombre_snapshot=insumo.nombre,
            cantidad=cantidad,
            unidad=insumo.unidad,
            precio_unitario=precio,
            impuesto_porcentaje=impuesto_pct,
            subtotal=subtotal,
            impuesto=impuesto,
            total=total,
        ))
        subtotal_total += subtotal
        impuesto_total += impuesto

    factura.subtotal = round(subtotal_total, 2)
    factura.impuesto = round(impuesto_total, 2)
    factura.total = round(subtotal_total + impuesto_total, 2)

    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Conflicto al registrar compra de insumos")

    row = db.query(CompraInsumoFactura).options(
        selectinload(CompraInsumoFactura.lineas),
        selectinload(CompraInsumoFactura.created_by),
        selectinload(CompraInsumoFactura.proveedor_rel),
    ).filter(CompraInsumoFactura.id == factura.id).first()
    return serialize_compra_insumo(row)


@router.get("/gestion-fincas/insumos", response_model=list[InsumoFincaOut])
def list_insumos_finca(current: Usuario = Depends(get_current_user), db: Session = Depends(get_db)):
    rows = db.query(InsumoFinca).order_by(InsumoFinca.activo.desc(), InsumoFinca.nombre.asc()).all()
    return [serialize_insumo_finca(row) for row in rows]


@router.post("/gestion-fincas/insumos", response_model=InsumoFincaOut, status_code=201)
def create_insumo_finca(payload: InsumoFincaCreate, current: Usuario = Depends(require_roles("administrativo", "admin")), db: Session = Depends(get_db)):
    data = payload.model_dump()
    data["codigo"] = (data.get("codigo") or "").strip() or next_code(db, InsumoFinca, "codigo", "INS-", 3)
    row = InsumoFinca(**data)
    db.add(row)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Código o nombre de insumo duplicado")
    db.refresh(row)
    return serialize_insumo_finca(row)


@router.patch("/gestion-fincas/insumos/{insumo_id}", response_model=InsumoFincaOut)
def update_insumo_finca(insumo_id: int, payload: InsumoFincaUpdate, current: Usuario = Depends(require_roles("administrativo", "admin")), db: Session = Depends(get_db)):
    row = db.query(InsumoFinca).filter(InsumoFinca.id == insumo_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Insumo no encontrado")
    data = payload.model_dump(exclude_unset=True)
    ensure_immutable_code(row.codigo, data.get("codigo"), "código del insumo")
    data.pop("codigo", None)
    for key, value in data.items():
        setattr(row, key, value.strip() if isinstance(value, str) else value)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Código o nombre de insumo duplicado")
    db.refresh(row)
    return serialize_insumo_finca(row)


@router.delete("/gestion-fincas/insumos/{insumo_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_insumo_finca(insumo_id: int, current: Usuario = Depends(get_current_user), db: Session = Depends(get_db)):
    require_superadmin_user(current)
    row = db.query(InsumoFinca).filter(InsumoFinca.id == insumo_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Insumo no encontrado")
    usados = db.query(RegistroFincaInsumo).filter(RegistroFincaInsumo.insumo_id == insumo_id).count()
    if usados:
        raise HTTPException(status_code=400, detail="El insumo ya tiene registros. Desactívelo en lugar de eliminarlo.")
    db.delete(row)
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.get("/gestion-fincas/registros", response_model=list[RegistroFincaOut])
def list_registros_finca(current: Usuario = Depends(require_roles("gerente", "administrativo", "supervisor_finca")), db: Session = Depends(get_db)):
    rows = registro_finca_query(db).order_by(RegistroFinca.semana_inicio.desc(), RegistroFinca.fecha.desc(), RegistroFinca.id.desc()).all()
    return [serialize_registro_finca(row) for row in rows]


@router.post("/gestion-fincas/registros", response_model=RegistroFincaOut, status_code=201)
def create_registro_finca(payload: RegistroFincaCreate, current: Usuario = Depends(require_roles("gerente", "administrativo", "supervisor_finca")), db: Session = Depends(get_db)):
    finca = db.query(Finca).filter(
        Finca.id == payload.finca_id,
        Finca.activa,
        Finca.gestion_fincas_habilitada,
    ).first()
    if not finca:
        raise HTTPException(status_code=404, detail="Finca no encontrada, inactiva o no habilitada para gestión de fincas")
    actividad = db.query(ActividadFinca).filter(ActividadFinca.id == payload.actividad_id, ActividadFinca.activa).first()
    if not actividad:
        raise HTTPException(status_code=404, detail="Actividad no encontrada o inactiva")
    inicio, fin = semana_operativa(payload.fecha)
    row = RegistroFinca(
        fecha=payload.fecha,
        semana_inicio=inicio,
        semana_fin=fin,
        finca_id=finca.id,
        actividad_id=actividad.id,
        descripcion=payload.descripcion,
        estado=payload.estado or "registrado",
        observaciones=payload.observaciones,
        created_by_id=current.id,
    )
    db.add(row)
    db.flush()
    apply_registro_finca_items(db, row, payload)
    db.commit()
    row = registro_finca_query(db).filter(RegistroFinca.id == row.id).first()
    return serialize_registro_finca(row)


@router.get("/gestion-fincas/registros/{registro_id}", response_model=RegistroFincaOut)
def get_registro_finca(registro_id: int, current: Usuario = Depends(require_roles("gerente", "administrativo", "supervisor_finca")), db: Session = Depends(get_db)):
    row = registro_finca_query(db).filter(RegistroFinca.id == registro_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Registro no encontrado")
    return serialize_registro_finca(row)


@router.patch("/gestion-fincas/registros/{registro_id}", response_model=RegistroFincaOut)
def update_registro_finca(registro_id: int, payload: RegistroFincaUpdate, current: Usuario = Depends(require_roles("gerente", "administrativo", "supervisor_finca")), db: Session = Depends(get_db)):
    row = registro_finca_query(db).filter(RegistroFinca.id == registro_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Registro no encontrado")
    data = payload.model_dump(exclude_unset=True, exclude={"trabajadores", "insumos"})
    if "fecha" in data and data["fecha"]:
        inicio, fin = semana_operativa(data["fecha"])
        row.semana_inicio = inicio
        row.semana_fin = fin
    if "finca_id" in data and data["finca_id"]:
        finca = db.query(Finca).filter(
            Finca.id == data["finca_id"],
            Finca.activa,
            Finca.gestion_fincas_habilitada,
        ).first()
        if not finca:
            raise HTTPException(status_code=404, detail="Finca no encontrada, inactiva o no habilitada para gestión de fincas")
    if "actividad_id" in data and data["actividad_id"]:
        actividad = db.query(ActividadFinca).filter(ActividadFinca.id == data["actividad_id"], ActividadFinca.activa).first()
        if not actividad:
            raise HTTPException(status_code=404, detail="Actividad no encontrada o inactiva")
    for key, value in data.items():
        setattr(row, key, value.strip() if isinstance(value, str) else value)
    apply_registro_finca_items(db, row, payload)
    db.commit()
    row = registro_finca_query(db).filter(RegistroFinca.id == registro_id).first()
    return serialize_registro_finca(row)


@router.delete("/gestion-fincas/registros/{registro_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_registro_finca(registro_id: int, current: Usuario = Depends(get_current_user), db: Session = Depends(get_db)):
    require_superadmin_user(current)
    row = registro_finca_query(db).filter(RegistroFinca.id == registro_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Registro no encontrado")
    restaurar_stock_registro(db, row)
    db.delete(row)
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


def read_xlsx_rows(file_bytes: bytes) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Falta dependencia openpyxl en backend: {exc}")
    wb = load_workbook(filename=BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h or "").strip().lower().replace(" ", "_") for h in rows[0]]
    data = []
    for raw in rows[1:]:
        item = {headers[i]: raw[i] if i < len(raw) else None for i in range(len(headers)) if headers[i]}
        if any(v not in (None, "") for v in item.values()):
            data.append(item)
    return data


@router.post("/gestion-fincas/import/{catalogo}")
async def import_gestion_fincas_catalogo(catalogo: str, file: UploadFile = File(...), current: Usuario = Depends(require_roles("gerente", "administrativo")), db: Session = Depends(get_db)):
    if not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Debe subir un archivo Excel .xlsx")
    rows = read_xlsx_rows(await file.read())
    creados = 0
    actualizados = 0
    errores: list[str] = []

    for idx, item in enumerate(rows, start=2):
        try:
            if catalogo == "trabajadores":
                nombre = str(item.get("nombre") or "").strip()
                if not nombre:
                    raise ValueError("Falta nombre")
                codigo = str(item.get("codigo") or "").strip() or None
                row = None
                if codigo:
                    row = db.query(TrabajadorFinca).filter(TrabajadorFinca.codigo == codigo).first()
                if not row and item.get("identificacion"):
                    row = db.query(TrabajadorFinca).filter(TrabajadorFinca.identificacion == str(item.get("identificacion")).strip()).first()
                if not row:
                    row = TrabajadorFinca()
                    db.add(row)
                    creados += 1
                else:
                    actualizados += 1
                row.codigo = codigo
                row.nombre = nombre
                row.identificacion = str(item.get("identificacion") or "").strip() or None
                row.telefono = str(item.get("telefono") or "").strip() or None
                row.puesto = str(item.get("puesto") or "").strip() or None
                row.jornal_diario = parse_float(item.get("jornal_diario"), row.jornal_diario)
                row.activo = parse_bool(item.get("activo"), True)
            elif catalogo == "actividades":
                nombre = str(item.get("nombre") or "").strip()
                if not nombre:
                    raise ValueError("Falta nombre")
                row = db.query(ActividadFinca).filter(ActividadFinca.nombre == nombre).first()
                if not row:
                    row = ActividadFinca(nombre=nombre)
                    db.add(row)
                    creados += 1
                else:
                    actualizados += 1
                row.tipo = str(item.get("tipo") or row.tipo or "mantenimiento").strip().lower()
                row.descripcion = str(item.get("descripcion") or "").strip() or None
                row.requiere_insumo = parse_bool(item.get("requiere_insumo"), False)
                row.unidad_referencia = str(item.get("unidad_referencia") or "").strip() or None
                row.activa = parse_bool(item.get("activa"), True)
            elif catalogo == "insumos":
                nombre = str(item.get("nombre") or "").strip()
                if not nombre:
                    raise ValueError("Falta nombre")
                codigo = str(item.get("codigo") or "").strip() or None
                row = db.query(InsumoFinca).filter(or_(InsumoFinca.nombre == nombre, InsumoFinca.codigo == codigo)).first() if codigo else db.query(InsumoFinca).filter(InsumoFinca.nombre == nombre).first()
                if not row:
                    row = InsumoFinca(nombre=nombre)
                    db.add(row)
                    creados += 1
                else:
                    actualizados += 1
                row.codigo = codigo
                row.tipo = str(item.get("tipo") or row.tipo or "general").strip().lower()
                row.unidad = str(item.get("unidad") or row.unidad or "unidad").strip()
                row.costo_unitario = parse_float(item.get("costo_unitario"), row.costo_unitario)
                row.observaciones = str(item.get("observaciones") or "").strip() or None
                row.activo = parse_bool(item.get("activo"), True)
            else:
                raise HTTPException(status_code=400, detail="Catálogo inválido. Use trabajadores, actividades o insumos")
        except Exception as exc:
            errores.append(f"Fila {idx}: {exc}")

    db.commit()
    return {"ok": True, "catalogo": catalogo, "creados": creados, "actualizados": actualizados, "errores": errores}




BACKUP_MODELS = [
    ConfiguracionSistema,
    ReciboConsecutivo,
    Usuario,
    Cliente,
    Finca,
    ReciboCafe,
    OrdenTrabajo,
    OrdenTrabajoUnionOrigen,
    OrdenTrabajoRecibo,
    ActividadFinca,
    TrabajadorFinca,
    InsumoFinca,
    Proveedor,
    ProveedorContacto,
    CompraInsumoFactura,
    CompraInsumoLinea,
    RegistroFinca,
    RegistroFincaTrabajador,
    RegistroFincaInsumo,
    SeguimientoOT,
    DocumentoLote,
    ComentarioLote,
    CotizacionVenta,
    CotizacionVentaLinea,
    SolicitudSalidaVenta,
    SolicitudSalidaVentaLinea,
    VentaLeadLote,
    VentaSeguimiento,
    SolicitudVenta,
    SolicitudVentaLinea,
    SolicitudVentaLineaLiquidacion,
    SolicitudVentaSeguimiento,
    SolicitudVentaDocumento,
]

RESET_DELETE_MODELS = [
    SolicitudVentaDocumento,
    SolicitudVentaSeguimiento,
    SolicitudVentaLineaLiquidacion,
    SolicitudVentaLinea,
    SolicitudVenta,
    VentaSeguimiento,
    VentaLeadLote,
    SolicitudSalidaVentaLinea,
    SolicitudSalidaVenta,
    CotizacionVentaLinea,
    CotizacionVenta,
    ComentarioLote,
    DocumentoLote,
    SeguimientoOT,
    OrdenTrabajoUnionOrigen,
    OrdenTrabajoRecibo,
    OrdenTrabajo,
    ReciboCafe,
    RegistroFincaInsumo,
    RegistroFincaTrabajador,
    RegistroFinca,
    CompraInsumoLinea,
    CompraInsumoFactura,
    InsumoFinca,
    TrabajadorFinca,
    ActividadFinca,
    Finca,
    Cliente,
]


def serialize_backup_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def model_to_backup_dict(row) -> dict:
    return {column.name: serialize_backup_value(getattr(row, column.name)) for column in row.__table__.columns}


def parse_backup_value(column, value):
    if value is None:
        return None
    try:
        python_type = column.type.python_type
    except Exception:
        python_type = None

    if python_type is datetime and isinstance(value, str):
        return datetime.fromisoformat(value)
    if python_type is date and isinstance(value, str):
        return date.fromisoformat(value)
    return value


@router.get("/respaldos/exportar")
def exportar_respaldo(current: Usuario = Depends(require_roles("gerente")), db: Session = Depends(get_db)):
    payload = {
        "schema": "navia_backup_v1",
        "exported_at": utcnow().isoformat(),
        "exported_by": {
            "id": current.id,
            "nombre": current.nombre,
            "correo": current.correo,
        },
        "tables": {},
    }

    for model in BACKUP_MODELS:
        rows = db.query(model).order_by(model.id.asc()).all()
        payload["tables"][model.__tablename__] = [model_to_backup_dict(row) for row in rows]

    filename = f"navia-respaldo-{utcnow().strftime('%Y%m%d-%H%M%S')}.json"
    return Response(
        content=json.dumps(payload, ensure_ascii=False, indent=2),
        media_type="application/json",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/respaldos/importar")
async def importar_respaldo(file: UploadFile = File(...), current: Usuario = Depends(require_roles("gerente")), db: Session = Depends(get_db)):
    raw = await file.read()
    if len(raw) > 30 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="El respaldo no debe superar 30 MB")

    try:
        payload = json.loads(raw.decode("utf-8"))
    except Exception:
        raise HTTPException(status_code=422, detail="El archivo no es un JSON válido de respaldo NAVIA")

    if payload.get("schema") != "navia_backup_v1" or not isinstance(payload.get("tables"), dict):
        raise HTTPException(status_code=422, detail="El archivo no corresponde a un respaldo NAVIA compatible")

    imported = 0
    updated = 0

    try:
        for model in BACKUP_MODELS:
            table = model.__tablename__
            rows = payload["tables"].get(table, [])
            if not isinstance(rows, list):
                continue

            columns = {column.name: column for column in model.__table__.columns}
            for item in rows:
                if not isinstance(item, dict):
                    continue
                row_id = item.get("id")
                values = {
                    key: parse_backup_value(columns[key], value)
                    for key, value in item.items()
                    if key in columns
                }
                if row_id is None:
                    continue

                existing = db.query(model).filter(model.id == row_id).first()
                if existing:
                    for key, value in values.items():
                        setattr(existing, key, value)
                    updated += 1
                else:
                    db.add(model(**values))
                    imported += 1

        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail=f"El respaldo genera conflictos con datos existentes: {exc.orig}")
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"No se pudo importar el respaldo: {exc}")

    return {"ok": True, "importados": imported, "actualizados": updated}


@router.post("/respaldos/borrado-completo")
def borrado_completo_app(current: Usuario = Depends(get_current_user), db: Session = Depends(get_db)):
    ensure_bootstrap_superadmin(current)

    deleted = {}
    try:
        for model in RESET_DELETE_MODELS:
            count = db.query(model).delete(synchronize_session=False)
            deleted[model.__tablename__] = int(count or 0)

        db.query(SesionUsuario).filter(SesionUsuario.usuario_id != current.id).delete(synchronize_session=False)
        db.query(Usuario).filter(func.lower(Usuario.correo) != ADMIN_EMAIL).delete(synchronize_session=False)
        db.commit()
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"No se pudo limpiar la aplicación: {exc}")

    return {"ok": True, "message": "Aplicación limpiada. Se conservó el superadmin de bootstrap y la configuración base.", "deleted": deleted}

@router.get("/dashboard", response_model=DashboardOut)
@router.get("/overview", response_model=DashboardOut)
def dashboard(current: Usuario = Depends(get_current_user), db: Session = Depends(get_db)):
    total_recibos = db.query(ReciboCafe).count()
    total_ots = db.query(OrdenTrabajo).count()
    ots_activas = db.query(OrdenTrabajo).filter(OrdenTrabajo.estado.in_(["abierta", "en_proceso", "pendiente_aprobacion"])).count()
    total_cajuelas = db.query(func.coalesce(func.sum(ReciboCafe.cajuelas), 0)).scalar() or 0
    total_fanegas = round(float(total_cajuelas or 0) / 20, 2)

    by_month = defaultdict(float)
    for r in db.query(ReciboCafe).order_by(ReciboCafe.fecha.asc()).all():
        by_month[r.fecha.strftime("%Y-%m")] += estimated_fanegas(r.cajuelas, r.cuartillos)

    estado_rows = db.query(OrdenTrabajo.estado, func.count(OrdenTrabajo.id)).group_by(OrdenTrabajo.estado).all()
    sugerencias = []
    if total_recibos == 0:
        sugerencias.append("Registre el primer recibo de café para iniciar la trazabilidad.")
    if total_recibos and total_ots == 0:
        sugerencias.append("Ya hay café recibido: puede crear una OT y asignar el proceso miel, natural o semilavado.")
    if ots_activas:
        sugerencias.append("Revise las OT activas y solicite a los operarios registrar el diagnóstico diario.")
    sugerencias.append("El sistema está preparado para trabajo offline: los registros locales se sincronizan al recuperar conexión.")

    return DashboardOut(
        kpis=[
            DashboardKPI(key="recibos", label="Recibos", value=total_recibos),
            DashboardKPI(key="fanegas", label="Fanegas estimadas", value=total_fanegas),
            DashboardKPI(key="ots", label="OT totales", value=total_ots),
            DashboardKPI(key="ots_activas", label="OT activas", value=ots_activas),
        ],
        recibos_por_mes=[ChartPoint(label=k, value=round(v, 2)) for k, v in sorted(by_month.items())[-8:]],
        ots_por_estado=[ChartPoint(label=e or "sin_estado", value=float(c)) for e, c in estado_rows],
        sugerencias=sugerencias,
    )


def safe_filename(name: str) -> str:
    raw = Path(name or "archivo").name
    cleaned = "".join(ch if ch.isalnum() or ch in {".", "-", "_"} else "_" for ch in raw).strip("._")
    return cleaned or "archivo"


ALLOWED_DOCUMENT_EXTENSIONS = {
    ".pdf", ".png", ".jpg", ".jpeg", ".webp",
    ".doc", ".docx", ".xls", ".xlsx", ".csv", ".txt",
}
BLOCKED_CONTENT_TYPES = {
    "text/html", "application/xhtml+xml", "image/svg+xml",
    "application/javascript", "text/javascript",
}


def save_upload_file(file: UploadFile, folder: Path) -> tuple[str, str, int | None, str | None]:
    folder.mkdir(parents=True, exist_ok=True)
    original = safe_filename(file.filename or "archivo")
    extension = Path(original).suffix.lower()
    content_type = (file.content_type or "application/octet-stream").split(";", 1)[0].strip().lower()

    if extension not in ALLOWED_DOCUMENT_EXTENSIONS or content_type in BLOCKED_CONTENT_TYPES:
        raise HTTPException(
            status_code=422,
            detail="Tipo de archivo no permitido. Use PDF, imagen, Word, Excel, CSV o TXT.",
        )

    stored = f"{uuid.uuid4().hex}{extension}"
    path = folder / stored
    max_bytes = MAX_REQUEST_MB * 1024 * 1024
    size = 0

    try:
        with path.open("wb") as buffer:
            while True:
                chunk = file.file.read(1024 * 1024)
                if not chunk:
                    break
                size += len(chunk)
                if size > max_bytes:
                    raise HTTPException(
                        status_code=413,
                        detail=f"El archivo no debe superar {MAX_REQUEST_MB} MB",
                    )
                buffer.write(chunk)
    except Exception:
        path.unlink(missing_ok=True)
        raise

    public = "/uploads/" + "/".join(path.relative_to(upload_root()).parts)
    return public, original, size, content_type




def calc_line_totals(cantidad: float, precio: float, impuesto_porcentaje: float) -> tuple[float, float, float]:
    subtotal = round(float(cantidad or 0) * float(precio or 0), 2)
    impuesto = round(subtotal * float(impuesto_porcentaje or 0) / 100, 2)
    total = round(subtotal + impuesto, 2)
    return subtotal, impuesto, total


def lote_quintales_base(db: Session, ot_id: int) -> float:
    """Cantidad comercial base del lote.

    La recepción se registra en cajuelas/fanegas; para OT, cotización y venta
    se usa el equivalente operativo en quintales: 20 cajuelas = 1 fanega/qq.
    """
    total_cajuelas = (
        db.query(
            func.coalesce(
                func.sum(OrdenTrabajoRecibo.cajuelas_asignadas + (OrdenTrabajoRecibo.cuartillos_asignados / 4)),
                0,
            )
        )
        .filter(OrdenTrabajoRecibo.ot_id == ot_id)
        .scalar()
        or 0
    )
    if float(total_cajuelas or 0) > 0:
        return round(float(total_cajuelas or 0) / 20, 3)

    ot = db.query(OrdenTrabajo).filter(OrdenTrabajo.id == ot_id).first()
    return round(float(ot.fanegas_estimadas or 0), 3) if ot else 0.0


def lote_quintales_vendidos(db: Session, ot_id: int) -> float:
    vendido = (
        db.query(func.coalesce(func.sum(SolicitudSalidaVentaLinea.cantidad_quintales), 0))
        .join(SolicitudSalidaVenta, SolicitudSalidaVenta.id == SolicitudSalidaVentaLinea.solicitud_id)
        .filter(SolicitudSalidaVentaLinea.ot_id == ot_id)
        .filter(SolicitudSalidaVenta.estado != "anulada")
        .scalar()
        or 0
    )
    return round(float(vendido or 0), 3)


def lote_quintales_disponibles(db: Session, ot_id: int) -> float:
    base = lote_quintales_base(db, ot_id)
    vendido = lote_quintales_vendidos(db, ot_id)
    return round(max(base - vendido, 0), 3)


def validate_lote_quantity_available(db: Session, ot: OrdenTrabajo | None, cantidad_quintales: float, contexto: str = "cotización") -> None:
    cantidad = float(cantidad_quintales or 0)
    if cantidad <= 0:
        raise HTTPException(status_code=400, detail=f"La cantidad de {contexto} debe ser mayor a cero")

    if not ot:
        return

    disponible = lote_quintales_disponibles(db, ot.id)
    if cantidad > disponible + 0.0001:
        raise HTTPException(
            status_code=422,
            detail=(
                f"El lote {ot.codigo_lote} solo tiene {disponible} qq disponibles. "
                f"No se puede registrar {cantidad} qq en la {contexto}."
            ),
        )


def sincronizar_estado_comercial_lote(db: Session, ot: OrdenTrabajo | None) -> bool:
    """Sincroniza el estado comercial del lote según su saldo real.

    - Si el lote aprobado se consumió completo en salidas, pasa a vendido.
    - Si conserva saldo, permanece en finalizada para seguir apareciendo en Lotes Aprobados.
    """
    if not ot or str(ot.estado or "").lower() not in {"finalizada", "vendido"}:
        return False
    db.flush()
    disponible = lote_quintales_disponibles(db, ot.id)
    nuevo_estado = "vendido" if disponible <= 0.0001 else "finalizada"
    if ot.estado != nuevo_estado:
        ot.estado = nuevo_estado
        ot.updated_at = utcnow()
        return True
    return False


def cotizacion_query(db: Session):
    return db.query(CotizacionVenta).options(
        selectinload(CotizacionVenta.cliente),
        selectinload(CotizacionVenta.created_by),
        selectinload(CotizacionVenta.lineas).selectinload(CotizacionVentaLinea.ot),
    )


def solicitud_salida_query(db: Session):
    return db.query(SolicitudSalidaVenta).options(
        selectinload(SolicitudSalidaVenta.cliente),
        selectinload(SolicitudSalidaVenta.created_by),
        selectinload(SolicitudSalidaVenta.cotizacion),
        selectinload(SolicitudSalidaVenta.lineas).selectinload(SolicitudSalidaVentaLinea.ot),
    )


def solicitud_venta_query(db: Session):
    return db.query(SolicitudVenta).options(
        selectinload(SolicitudVenta.cliente),
        selectinload(SolicitudVenta.ot),
        selectinload(SolicitudVenta.created_by),
        selectinload(SolicitudVenta.lineas)
            .selectinload(SolicitudVentaLinea.liquidaciones)
            .selectinload(SolicitudVentaLineaLiquidacion.ot),
        selectinload(SolicitudVenta.lineas)
            .selectinload(SolicitudVentaLinea.liquidaciones)
            .selectinload(SolicitudVentaLineaLiquidacion.salida),
        selectinload(SolicitudVenta.lineas)
            .selectinload(SolicitudVentaLinea.liquidaciones)
            .selectinload(SolicitudVentaLineaLiquidacion.created_by),
        selectinload(SolicitudVenta.seguimientos).selectinload(SolicitudVentaSeguimiento.created_by),
        selectinload(SolicitudVenta.documentos).selectinload(SolicitudVentaDocumento.created_by),
    )


def solicitud_linea_pendiente(linea: SolicitudVentaLinea) -> float:
    liquidado = sum(float(liq.cantidad_quintales or 0) for liq in (linea.liquidaciones or []))
    return round(max(float(linea.cantidad_quintales or 0) - liquidado, 0), 3)


@router.get("/cotizaciones-venta", response_model=list[CotizacionVentaOut])
def list_cotizaciones_venta(
    ot_id: Optional[int] = None,
    cliente_id: Optional[int] = None,
    estado: Optional[str] = None,
    current: Usuario = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    query = cotizacion_query(db)
    if cliente_id:
        query = query.filter(CotizacionVenta.cliente_id == cliente_id)
    if estado:
        query = query.filter(CotizacionVenta.estado == estado)
    if ot_id:
        query = query.join(CotizacionVentaLinea).filter(CotizacionVentaLinea.ot_id == ot_id)
    rows = query.order_by(CotizacionVenta.created_at.desc()).all()
    return [serialize_cotizacion_venta(row) for row in rows]


@router.get("/cotizaciones-venta/{codigo}", response_model=CotizacionVentaOut)
def get_cotizacion_venta(codigo: str, current: Usuario = Depends(get_current_user), db: Session = Depends(get_db)):
    row = cotizacion_query(db).filter(CotizacionVenta.codigo == codigo).first()
    if not row:
        raise HTTPException(status_code=404, detail="Cotización no encontrada")
    return serialize_cotizacion_venta(row)


@router.post("/cotizaciones-venta", response_model=CotizacionVentaOut, status_code=201)
def create_cotizacion_venta(
    payload: CotizacionVentaCreate,
    current: Usuario = Depends(require_roles("gerente", "administrativo")),
    db: Session = Depends(get_db),
):
    if not payload.lineas:
        raise HTTPException(status_code=400, detail="La cotización requiere al menos una línea")
    row = CotizacionVenta(
        codigo=next_code(db, CotizacionVenta, "codigo", "COT-", 5),
        cliente_id=payload.cliente_id or None,
        estado="borrador",
        fecha=payload.fecha,
        validez_dias=payload.validez_dias,
        moneda=(payload.moneda or "USD").upper(),
        condiciones=payload.condiciones,
        observaciones=payload.observaciones,
        created_by_id=current.id,
    )
    db.add(row)
    db.flush()
    cliente_cotizacion = db.query(Cliente).filter(Cliente.id == row.cliente_id).first() if row.cliente_id else None
    cliente_nombre_cotizacion = cliente_cotizacion.nombre_completo if cliente_cotizacion else "cliente sin registrar"
    cotizacion_eventos: list[tuple[int, float]] = []
    for item in payload.lineas:
        ot = db.query(OrdenTrabajo).filter(OrdenTrabajo.id == item.ot_id).first() if item.ot_id else None
        if item.ot_id and not ot:
            raise HTTPException(status_code=404, detail=f"Lote {item.ot_id} no encontrado")
        validate_lote_quantity_available(db, ot, item.cantidad_quintales, "cotización")
        subtotal, impuesto, total = calc_line_totals(item.cantidad_quintales, item.precio_unitario, item.impuesto_porcentaje)
        db.add(CotizacionVentaLinea(
            cotizacion_id=row.id,
            ot_id=item.ot_id,
            descripcion=item.descripcion or (f"Café {ot.proceso} · {ot.codigo_lote}" if ot else "Café"),
            cantidad_quintales=item.cantidad_quintales,
            precio_unitario=item.precio_unitario,
            impuesto_porcentaje=item.impuesto_porcentaje,
            subtotal=subtotal,
            impuesto=impuesto,
            total=total,
        ))
        if item.ot_id:
            cotizacion_eventos.append((item.ot_id, float(item.cantidad_quintales or 0)))
    for ot_id, cantidad in cotizacion_eventos:
        add_lote_evento(
            db,
            ot_id,
            "comercial",
            f"{current.nombre} creó la cotización {row.codigo} por {cantidad} qq para {cliente_nombre_cotizacion}.",
            current,
        )
    db.commit()
    row = cotizacion_query(db).filter(CotizacionVenta.id == row.id).first()
    return serialize_cotizacion_venta(row)


@router.patch("/cotizaciones-venta/{cotizacion_id}", response_model=CotizacionVentaOut)
def update_cotizacion_venta(
    cotizacion_id: int,
    payload: CotizacionVentaUpdate,
    current: Usuario = Depends(require_roles("gerente")),
    db: Session = Depends(get_db),
):
    row = db.query(CotizacionVenta).filter(CotizacionVenta.id == cotizacion_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Cotización no encontrada")

    data = payload.model_dump(exclude_unset=True)
    lineas = data.pop("lineas", None)

    for key, value in data.items():
        if key == "moneda" and value:
            value = str(value).upper()
        setattr(row, key, value)

    cotizacion_update_eventos: list[tuple[int, float]] = []

    if lineas is not None:
        if not lineas:
            raise HTTPException(status_code=400, detail="La cotización requiere al menos una línea")
        db.query(CotizacionVentaLinea).filter(CotizacionVentaLinea.cotizacion_id == row.id).delete(synchronize_session=False)
        db.flush()
        for item in lineas:
            item_data = item if isinstance(item, dict) else item.model_dump()
            ot_id = item_data.get("ot_id")
            ot = db.query(OrdenTrabajo).filter(OrdenTrabajo.id == ot_id).first() if ot_id else None
            if ot_id and not ot:
                raise HTTPException(status_code=404, detail=f"Lote {ot_id} no encontrado")
            cantidad = float(item_data.get("cantidad_quintales") or 0)
            validate_lote_quantity_available(db, ot, cantidad, "cotización")
            precio = float(item_data.get("precio_unitario") or 0)
            impuesto_porcentaje = float(item_data.get("impuesto_porcentaje") or 0)
            subtotal, impuesto, total = calc_line_totals(cantidad, precio, impuesto_porcentaje)
            descripcion = (item_data.get("descripcion") or (f"Café {ot.proceso} · {ot.codigo_lote}" if ot else "Café")).strip()
            db.add(CotizacionVentaLinea(
                cotizacion_id=row.id,
                ot_id=ot_id,
                descripcion=descripcion,
                cantidad_quintales=cantidad,
                precio_unitario=precio,
                impuesto_porcentaje=impuesto_porcentaje,
                subtotal=subtotal,
                impuesto=impuesto,
                total=total,
            ))
            if ot_id:
                cotizacion_update_eventos.append((int(ot_id), cantidad))

    row.updated_at = utcnow()
    if not cotizacion_update_eventos:
        cotizacion_update_eventos = [(linea.ot_id, float(linea.cantidad_quintales or 0)) for linea in (row.lineas or []) if linea.ot_id]
    for ot_id, cantidad in cotizacion_update_eventos:
        add_lote_evento(
            db,
            ot_id,
            "comercial",
            f"{current.nombre} actualizó la cotización {row.codigo}. Estado: {row.estado}. Cantidad relacionada: {cantidad} qq.",
            current,
        )
    db.commit()
    row = cotizacion_query(db).filter(CotizacionVenta.id == row.id).first()
    return serialize_cotizacion_venta(row)


@router.get("/solicitudes-salida", response_model=list[SolicitudSalidaVentaOut])
def list_solicitudes_salida(
    ot_id: Optional[int] = None,
    cliente_id: Optional[int] = None,
    current: Usuario = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    query = solicitud_salida_query(db)
    if cliente_id:
        query = query.filter(SolicitudSalidaVenta.cliente_id == cliente_id)
    if ot_id:
        query = query.join(SolicitudSalidaVentaLinea).filter(SolicitudSalidaVentaLinea.ot_id == ot_id)
    rows = query.order_by(SolicitudSalidaVenta.created_at.desc()).all()
    return [serialize_solicitud_salida_venta(row) for row in rows]


@router.post("/solicitudes-salida", response_model=SolicitudSalidaVentaOut, status_code=201)
def create_solicitud_salida(
    payload: SolicitudSalidaVentaCreate,
    current: Usuario = Depends(require_roles("gerente", "administrativo")),
    db: Session = Depends(get_db),
):
    if not payload.lineas:
        raise HTTPException(status_code=400, detail="La solicitud de salida requiere al menos una línea de lote")

    cliente = db.query(Cliente).filter(Cliente.id == payload.cliente_id).first() if payload.cliente_id else None
    cliente_nombre = cliente.nombre_completo if cliente else "cliente sin registrar"
    total_por_lote: dict[int, float] = defaultdict(float)
    ots_por_id: dict[int, OrdenTrabajo] = {}

    for item in payload.lineas:
        ot = db.query(OrdenTrabajo).filter(OrdenTrabajo.id == item.ot_id).first()
        if not ot:
            raise HTTPException(status_code=404, detail=f"Lote {item.ot_id} no encontrado")
        if ot.estado != "finalizada":
            raise HTTPException(status_code=400, detail=f"El lote {ot.codigo_lote} no está disponible como lote aprobado")
        total_por_lote[ot.id] += float(item.cantidad_quintales or 0)
        ots_por_id[ot.id] = ot

    for ot_id, cantidad_total in total_por_lote.items():
        validate_lote_quantity_available(db, ots_por_id[ot_id], cantidad_total, "solicitud de salida")

    subtotal_total = 0.0
    impuesto_total = 0.0
    total_total = 0.0
    row = SolicitudSalidaVenta(
        codigo=next_code(db, SolicitudSalidaVenta, "codigo", "SAL-", 5),
        cotizacion_id=payload.cotizacion_id or None,
        cliente_id=payload.cliente_id or None,
        estado="vendida",
        fecha=payload.fecha,
        moneda=(payload.moneda or "USD").upper(),
        tipo_envio=payload.tipo_envio,
        direccion_entrega=payload.direccion_entrega,
        observaciones=payload.observaciones,
        created_by_id=current.id,
    )
    db.add(row)
    db.flush()

    for item in payload.lineas:
        ot = db.query(OrdenTrabajo).filter(OrdenTrabajo.id == item.ot_id).first()
        if not ot:
            raise HTTPException(status_code=404, detail=f"Lote {item.ot_id} no encontrado")
        if ot.estado != "finalizada":
            raise HTTPException(status_code=400, detail=f"El lote {ot.codigo_lote} no está disponible como lote aprobado")
        validate_lote_quantity_available(db, ot, item.cantidad_quintales, "solicitud de salida")
        disponible_previo = lote_quintales_disponibles(db, ot.id)
        descripcion = item.descripcion or f"Café {ot.proceso} · {ot.codigo_lote}"
        subtotal, impuesto, total = calc_line_totals(item.cantidad_quintales, item.precio_unitario, item.impuesto_porcentaje)
        subtotal_total += subtotal
        impuesto_total += impuesto
        total_total += total
        db.add(SolicitudSalidaVentaLinea(
            solicitud_id=row.id,
            ot_id=ot.id,
            descripcion=descripcion,
            cantidad_quintales=item.cantidad_quintales,
            precio_unitario=item.precio_unitario,
            impuesto_porcentaje=item.impuesto_porcentaje,
            subtotal=subtotal,
            impuesto=impuesto,
            total=total,
        ))
        disponible_posterior = round(max(disponible_previo - float(item.cantidad_quintales or 0), 0), 3)
        if disponible_posterior <= 0.0001:
            ot.estado = "vendido"
        else:
            ot.estado = "finalizada"
        ot.updated_at = utcnow()
        db.add(ComentarioLote(
            ot_id=ot.id,
            tipo="comercial",
            comentario=(
                f"Se vendieron {item.cantidad_quintales} qq del lote al cliente {cliente_nombre} en la solicitud de salida {row.codigo}. "
                f"Disponible anterior: {disponible_previo} qq; disponible posterior: {disponible_posterior} qq."
            ),
            created_by_id=current.id,
        ))

    for ot_id in total_por_lote:
        sincronizar_estado_comercial_lote(db, ots_por_id.get(ot_id) or db.query(OrdenTrabajo).filter(OrdenTrabajo.id == ot_id).first())

    row.subtotal = round(subtotal_total, 2)
    row.impuesto = round(impuesto_total, 2)
    row.total = round(total_total, 2)

    if payload.cotizacion_id:
        cot = db.query(CotizacionVenta).filter(CotizacionVenta.id == payload.cotizacion_id).first()
        if cot:
            cot.estado = "cerrada"
            cot.updated_at = utcnow()

    db.commit()
    row = solicitud_salida_query(db).filter(SolicitudSalidaVenta.id == row.id).first()
    return serialize_solicitud_salida_venta(row)


@router.get("/ot/lotes-aprobados", response_model=list[OTOut])
def list_lotes_aprobados(current: Usuario = Depends(get_current_user), db: Session = Depends(get_db)):
    rows = (
        ot_query(db)
        .filter(OrdenTrabajo.estado.in_(["finalizada", "vendido"]))
        .order_by(OrdenTrabajo.fecha_cierre.desc().nullslast(), OrdenTrabajo.updated_at.desc())
        .all()
    )
    changed = False
    for row in rows:
        changed = sincronizar_estado_comercial_lote(db, row) or changed
    if changed:
        db.commit()
    aprobados = [row for row in rows if row.estado == "finalizada" and lote_quintales_disponibles(db, row.id) > 0.0001]
    return [serialize_ot(r) for r in aprobados]


@router.get("/ot/lotes-vendidos", response_model=list[OTOut])
def list_lotes_vendidos(current: Usuario = Depends(get_current_user), db: Session = Depends(get_db)):
    rows = (
        ot_query(db)
        .filter(OrdenTrabajo.estado.in_(["finalizada", "vendido"]))
        .order_by(OrdenTrabajo.updated_at.desc())
        .all()
    )
    changed = False
    for row in rows:
        changed = sincronizar_estado_comercial_lote(db, row) or changed
    if changed:
        db.commit()
    vendidos = [row for row in rows if row.estado == "vendido" or lote_quintales_disponibles(db, row.id) <= 0.0001]
    return [serialize_ot(r) for r in vendidos]


@router.post("/ot/{ot_id}/documentos", status_code=201)
def upload_lote_documento(
    ot_id: int,
    request: Request,
    titulo: str = Form("Documento vinculado"),
    tipo: str = Form("documento"),
    descripcion: Optional[str] = Form(None),
    file: UploadFile = File(...),
    current: Usuario = Depends(require_roles("gerente", "administrativo", "operario", "supervisor_finca")),
    db: Session = Depends(get_db),
):
    row = db.query(OrdenTrabajo).filter(OrdenTrabajo.id == ot_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Lote no encontrado")
    ensure_lote_not_unido(row, "adjuntar documentos")
    if current.rol == "operario" and row.operario_id != current.id:
        raise HTTPException(status_code=403, detail="Solo puedes adjuntar documentos a tus lotes asignados")

    file_url, original, size, content_type = save_upload_file(file, upload_root() / "lotes" / str(ot_id))
    if file_url.startswith("/uploads/"):
        file_url = f"{APP_BASE_URL}{file_url}"
    doc = DocumentoLote(
        ot_id=ot_id,
        titulo=(titulo or original).strip() or original,
        tipo=(tipo or "documento").strip().lower(),
        descripcion=descripcion,
        file_url=file_url,
        file_name=original,
        content_type=content_type,
        size_bytes=size,
        created_by_id=current.id,
    )
    db.add(doc)
    add_lote_evento(
        db,
        row.id,
        "documento",
        f"{current.nombre} adjuntó documento al lote: {doc.titulo} ({doc.file_name}).",
        current,
    )
    db.commit()
    db.refresh(doc)
    return serialize_lote_documento(doc)


@router.delete("/ot/documentos/{documento_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_lote_documento(
    documento_id: int,
    current: Usuario = Depends(require_roles("gerente", "administrativo")),
    db: Session = Depends(get_db),
):
    doc = db.query(DocumentoLote).filter(DocumentoLote.id == documento_id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Documento no encontrado")
    db.delete(doc)
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.post("/ot/{ot_id}/comentarios", status_code=201)
def create_lote_comentario(
    ot_id: int,
    payload: ComentarioLoteCreate,
    current: Usuario = Depends(require_roles("gerente", "administrativo", "operario", "supervisor_finca")),
    db: Session = Depends(get_db),
):
    row = db.query(OrdenTrabajo).filter(OrdenTrabajo.id == ot_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Lote no encontrado")
    ensure_lote_not_unido(row, "agregar notas o registros")
    comentario = ComentarioLote(
        ot_id=ot_id,
        tipo=(payload.tipo or "nota").strip().lower(),
        comentario=payload.comentario.strip(),
        created_by_id=current.id,
    )
    db.add(comentario)
    db.commit()
    db.refresh(comentario)
    return serialize_lote_comentario(comentario)


@router.post("/ot/{ot_id}/ventas", status_code=201)
def create_venta_lead(
    ot_id: int,
    payload: VentaLeadCreate,
    current: Usuario = Depends(require_roles("gerente", "administrativo")),
    db: Session = Depends(get_db),
):
    row = db.query(OrdenTrabajo).filter(OrdenTrabajo.id == ot_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Lote no encontrado")
    if row.estado not in {"finalizada", "vendido"}:
        raise HTTPException(status_code=400, detail="Solo se puede comercializar un lote aprobado")
    if payload.cliente_id:
        cliente = db.query(Cliente).filter(Cliente.id == payload.cliente_id, Cliente.activo).first()
        if not cliente:
            raise HTTPException(status_code=404, detail="Cliente no encontrado o inactivo")
    lead = VentaLeadLote(
        ot_id=ot_id,
        cliente_id=payload.cliente_id,
        cantidad_quintales=payload.cantidad_quintales,
        precio_unitario=payload.precio_unitario,
        moneda=(payload.moneda or "USD").upper(),
        notas=payload.notas,
        created_by_id=current.id,
    )
    db.add(lead)
    db.commit()
    lead = db.query(VentaLeadLote).options(selectinload(VentaLeadLote.cliente), selectinload(VentaLeadLote.created_by), selectinload(VentaLeadLote.seguimientos)).filter(VentaLeadLote.id == lead.id).first()
    return serialize_venta_lead(lead)


@router.patch("/ventas-leads/{lead_id}")
def update_venta_lead(
    lead_id: int,
    payload: VentaLeadUpdate,
    current: Usuario = Depends(require_roles("gerente", "administrativo")),
    db: Session = Depends(get_db),
):
    lead = db.query(VentaLeadLote).filter(VentaLeadLote.id == lead_id).first()
    if not lead:
        raise HTTPException(status_code=404, detail="Lead no encontrado")
    data = payload.model_dump(exclude_unset=True)
    for key, value in data.items():
        if key == "moneda" and value:
            value = str(value).upper()
        setattr(lead, key, value)
    db.commit()
    lead = db.query(VentaLeadLote).options(selectinload(VentaLeadLote.cliente), selectinload(VentaLeadLote.created_by), selectinload(VentaLeadLote.seguimientos).selectinload(VentaSeguimiento.created_by)).filter(VentaLeadLote.id == lead_id).first()
    return serialize_venta_lead(lead)


@router.post("/ventas-leads/{lead_id}/seguimientos", status_code=201)
def create_venta_seguimiento(
    lead_id: int,
    payload: VentaSeguimientoCreate,
    current: Usuario = Depends(require_roles("gerente", "administrativo")),
    db: Session = Depends(get_db),
):
    lead = db.query(VentaLeadLote).filter(VentaLeadLote.id == lead_id).first()
    if not lead:
        raise HTTPException(status_code=404, detail="Lead no encontrado")
    lead.estado = "seguimiento" if lead.estado == "lead" else lead.estado
    row = VentaSeguimiento(
        lead_id=lead_id,
        fecha=payload.fecha,
        canal=(payload.canal or "nota").strip().lower(),
        asunto=payload.asunto,
        comentario=payload.comentario.strip(),
        created_by_id=current.id,
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return serialize_venta_seguimiento(row)


@router.post("/ventas-leads/{lead_id}/solicitar-salida")
def solicitar_salida_lead(
    lead_id: int,
    payload: ComentarioLoteCreate,
    current: Usuario = Depends(require_roles("gerente", "administrativo")),
    db: Session = Depends(get_db),
):
    lead = db.query(VentaLeadLote).filter(VentaLeadLote.id == lead_id).first()
    if not lead:
        raise HTTPException(status_code=404, detail="Lead no encontrado")
    lead.estado = "solicitud_salida"
    db.add(VentaSeguimiento(
        lead_id=lead.id,
        fecha=date.today(),
        canal="salida",
        asunto="Solicitud de salida",
        comentario=payload.comentario.strip(),
        created_by_id=current.id,
    ))
    db.commit()
    return {"ok": True, "estado": lead.estado}


@router.post("/ventas-leads/{lead_id}/marcar-vendido")
def marcar_lote_vendido(
    lead_id: int,
    payload: ComentarioLoteCreate,
    current: Usuario = Depends(require_roles("gerente")),
    db: Session = Depends(get_db),
):
    lead = db.query(VentaLeadLote).filter(VentaLeadLote.id == lead_id).first()
    if not lead:
        raise HTTPException(status_code=404, detail="Lead no encontrado")
    ot = db.query(OrdenTrabajo).filter(OrdenTrabajo.id == lead.ot_id).first()
    if not ot:
        raise HTTPException(status_code=404, detail="Lote no encontrado")
    lead.estado = "vendido"
    ot.estado = "vendido"
    db.add(ComentarioLote(ot_id=ot.id, tipo="comercial", comentario=payload.comentario.strip(), created_by_id=current.id))
    db.commit()
    return serialize_ot(ot_query(db).filter(OrdenTrabajo.id == ot.id).first())


@router.get("/solicitudes-venta", response_model=list[SolicitudVentaOut])
def list_solicitudes_venta(current: Usuario = Depends(get_current_user), db: Session = Depends(get_db)):
    rows = solicitud_venta_query(db).order_by(SolicitudVenta.created_at.desc()).all()
    return [serialize_solicitud_venta(r) for r in rows]


@router.get("/solicitudes-venta/compatibles", response_model=list[SolicitudVentaOut])
def solicitudes_compatibles(
    ot_id: int,
    current: Usuario = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    ot = db.query(OrdenTrabajo).filter(OrdenTrabajo.id == ot_id).first()
    if not ot:
        raise HTTPException(status_code=404, detail="Lote no encontrado")
    disponible = lote_quintales_disponibles(db, ot.id)
    if disponible <= 0:
        return []
    rows = solicitud_venta_query(db).filter(SolicitudVenta.estado.in_(["pendiente", "parcial", "asignada"])).order_by(SolicitudVenta.created_at.asc()).all()
    compatibles = []
    for row in rows:
        lineas = row.lineas or []
        if not lineas and float(row.cantidad_quintales or 0) > 0:
            proceso = row.proceso_preferido
            if not proceso or proceso == "flexible" or proceso == ot.proceso:
                compatibles.append(row)
                continue
        for linea in lineas:
            proceso = (linea.proceso_preferido or "flexible").lower()
            if solicitud_linea_pendiente(linea) <= 0:
                continue
            if proceso in {"", "flexible"} or proceso == (ot.proceso or "").lower():
                compatibles.append(row)
                break
    return [serialize_solicitud_venta(r) for r in compatibles]


@router.get("/solicitudes-venta/{solicitud_id}/lineas/{linea_id}/lotes-compatibles")
def lotes_compatibles_para_solicitud_linea(
    solicitud_id: int,
    linea_id: int,
    current: Usuario = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    linea = (
        db.query(SolicitudVentaLinea)
        .options(selectinload(SolicitudVentaLinea.liquidaciones))
        .filter(SolicitudVentaLinea.id == linea_id, SolicitudVentaLinea.solicitud_id == solicitud_id)
        .first()
    )
    if not linea:
        raise HTTPException(status_code=404, detail="Línea de solicitud no encontrada")
    pendiente = solicitud_linea_pendiente(linea)
    if pendiente <= 0:
        return []
    query = ot_query(db).filter(OrdenTrabajo.estado == "finalizada")
    proceso = (linea.proceso_preferido or "flexible").lower()
    if proceso not in {"", "flexible"}:
        query = query.filter(OrdenTrabajo.proceso == proceso)
    rows = []
    for ot in query.order_by(OrdenTrabajo.fecha_cierre.asc().nullslast(), OrdenTrabajo.updated_at.asc()).all():
        disponible = lote_quintales_disponibles(db, ot.id)
        if disponible <= 0:
            continue
        rows.append({
            "id": ot.id,
            "codigo_lote": ot.codigo_lote,
            "proceso": ot.proceso,
            "finca_nombre": ot.finca.nombre if ot.finca else None,
            "quintales_disponibles": disponible,
            "cantidad_sugerida": round(min(disponible, pendiente), 3),
            "fecha_cierre": ot.fecha_cierre,
        })
    return rows


@router.post("/solicitudes-venta", response_model=SolicitudVentaOut, status_code=201)
def create_solicitud_venta(
    payload: SolicitudVentaCreate,
    current: Usuario = Depends(require_roles("gerente", "administrativo")),
    db: Session = Depends(get_db),
):
    if payload.cliente_id:
        cliente = db.query(Cliente).filter(Cliente.id == payload.cliente_id, Cliente.activo).first()
        if not cliente:
            raise HTTPException(status_code=404, detail="Cliente no encontrado o inactivo")

    lineas_payload = list(payload.lineas or [])
    if not lineas_payload and payload.cantidad_quintales:
        lineas_payload.append(type("LegacyLinea", (), {
            "descripcion": None,
            "proceso_preferido": payload.proceso_preferido,
            "cantidad_quintales": payload.cantidad_quintales,
            "precio_objetivo": payload.precio_objetivo,
            "moneda": payload.moneda,
            "observaciones": None,
        })())
    if not lineas_payload:
        raise HTTPException(status_code=400, detail="La solicitud requiere al menos una línea de café")

    total_quintales = round(sum(float(item.cantidad_quintales or 0) for item in lineas_payload), 3)
    procesos = {str(item.proceso_preferido or "flexible").lower() for item in lineas_payload}
    proceso_padre = next(iter(procesos)) if len(procesos) == 1 else "flexible"
    row = SolicitudVenta(
        codigo=next_code(db, SolicitudVenta, "codigo", "SV-", 5),
        cliente_id=payload.cliente_id,
        estado="pendiente",
        cantidad_quintales=total_quintales,
        proceso_preferido=None if proceso_padre == "flexible" else proceso_padre,
        precio_objetivo=payload.precio_objetivo,
        moneda=(payload.moneda or "USD").upper(),
        observaciones=payload.observaciones,
        created_by_id=current.id,
    )
    db.add(row)
    db.flush()
    for idx, item in enumerate(lineas_payload, start=1):
        proceso = (item.proceso_preferido or "flexible").lower()
        descripcion = (item.descripcion or f"Línea {idx}: café {proceso}").strip()
        db.add(SolicitudVentaLinea(
            solicitud_id=row.id,
            descripcion=descripcion,
            proceso_preferido=None if proceso == "flexible" else proceso,
            cantidad_quintales=float(item.cantidad_quintales or 0),
            precio_objetivo=item.precio_objetivo,
            moneda=(item.moneda or payload.moneda or "USD").upper(),
            observaciones=item.observaciones,
        ))
    db.commit()
    row = solicitud_venta_query(db).filter(SolicitudVenta.id == row.id).first()
    return serialize_solicitud_venta(row)


@router.get("/solicitudes-venta/{solicitud_id}", response_model=SolicitudVentaOut)
def get_solicitud_venta(
    solicitud_id: int,
    current: Usuario = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    row = solicitud_venta_query(db).filter(SolicitudVenta.id == solicitud_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Solicitud no encontrada")
    return serialize_solicitud_venta(row)


@router.patch("/solicitudes-venta/{solicitud_id}", response_model=SolicitudVentaOut)
def update_solicitud_venta(
    solicitud_id: int,
    payload: SolicitudVentaUpdate,
    current: Usuario = Depends(require_roles("gerente")),
    db: Session = Depends(get_db),
):
    row = solicitud_venta_query(db).filter(SolicitudVenta.id == solicitud_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Solicitud no encontrada")
    data = payload.model_dump(exclude_unset=True)
    lineas = data.pop("lineas", None)
    for key, value in data.items():
        if key == "moneda" and value:
            value = str(value).upper()
        setattr(row, key, value)

    if lineas is not None:
        tiene_liquidaciones = any((linea.liquidaciones or []) for linea in (row.lineas or []))
        if tiene_liquidaciones:
            raise HTTPException(status_code=400, detail="No se pueden reemplazar líneas que ya tienen liquidaciones. Cree una nueva solicitud o agregue una nota de ajuste.")
        if not lineas:
            raise HTTPException(status_code=400, detail="La solicitud requiere al menos una línea de café")
        db.query(SolicitudVentaLinea).filter(SolicitudVentaLinea.solicitud_id == row.id).delete(synchronize_session=False)
        db.flush()
        total = 0.0
        procesos = set()
        for idx, item in enumerate(lineas, start=1):
            item_data = item if isinstance(item, dict) else item.model_dump()
            proceso = str(item_data.get("proceso_preferido") or "flexible").lower()
            cantidad = float(item_data.get("cantidad_quintales") or 0)
            total += cantidad
            procesos.add(proceso)
            db.add(SolicitudVentaLinea(
                solicitud_id=row.id,
                descripcion=(item_data.get("descripcion") or f"Línea {idx}: café {proceso}").strip(),
                proceso_preferido=None if proceso == "flexible" else proceso,
                cantidad_quintales=cantidad,
                precio_objetivo=item_data.get("precio_objetivo"),
                moneda=(item_data.get("moneda") or row.moneda or "USD").upper(),
                observaciones=item_data.get("observaciones"),
            ))
        row.cantidad_quintales = round(total, 3)
        row.proceso_preferido = None if len(procesos) != 1 or "flexible" in procesos else next(iter(procesos))

    row.updated_at = utcnow()
    db.commit()
    row = solicitud_venta_query(db).filter(SolicitudVenta.id == solicitud_id).first()
    return serialize_solicitud_venta(row)


@router.delete("/solicitudes-venta/{solicitud_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_solicitud_venta(
    solicitud_id: int,
    current: Usuario = Depends(require_roles("gerente")),
    db: Session = Depends(get_db),
):
    row = solicitud_venta_query(db).filter(SolicitudVenta.id == solicitud_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Solicitud no encontrada")
    if row.estado in {"vendida", "salida_solicitada"} or any((linea.liquidaciones or []) for linea in (row.lineas or [])):
        raise HTTPException(status_code=400, detail="No se puede eliminar una solicitud con ventas o liquidaciones asociadas")
    db.delete(row)
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.post("/solicitudes-venta/{solicitud_id}/asignar-lote", response_model=SolicitudVentaOut)
def asignar_lote_solicitud(
    solicitud_id: int,
    payload: AssignSolicitudVentaPayload,
    current: Usuario = Depends(require_roles("gerente", "administrativo")),
    db: Session = Depends(get_db),
):
    solicitud = solicitud_venta_query(db).filter(SolicitudVenta.id == solicitud_id).first()
    if not solicitud:
        raise HTTPException(status_code=404, detail="Solicitud no encontrada")
    ot = db.query(OrdenTrabajo).filter(OrdenTrabajo.id == payload.ot_id, OrdenTrabajo.estado == "finalizada").first()
    if not ot:
        raise HTTPException(status_code=404, detail="Lote aprobado no encontrado")
    solicitud.ot_id = ot.id
    solicitud.estado = "asignada"
    if payload.comentario:
        db.add(SolicitudVentaSeguimiento(
            solicitud_id=solicitud.id,
            fecha=date.today(),
            canal="asignacion",
            asunto=f"Lote asignado {ot.codigo_lote}",
            comentario=payload.comentario,
            created_by_id=current.id,
        ))
    db.commit()
    solicitud = solicitud_venta_query(db).filter(SolicitudVenta.id == solicitud_id).first()
    return serialize_solicitud_venta(solicitud)


@router.post("/solicitudes-venta/{solicitud_id}/liquidar", response_model=SolicitudSalidaVentaOut, status_code=201)
def liquidar_solicitud_venta(
    solicitud_id: int,
    payload: SolicitudVentaLiquidarPayload,
    current: Usuario = Depends(require_roles("gerente", "administrativo")),
    db: Session = Depends(get_db),
):
    solicitud = solicitud_venta_query(db).filter(SolicitudVenta.id == solicitud_id).first()
    if not solicitud:
        raise HTTPException(status_code=404, detail="Solicitud no encontrada")
    if not payload.lineas:
        raise HTTPException(status_code=400, detail="Debe seleccionar al menos un lote para liquidar")

    cliente_id = payload.cliente_id or solicitud.cliente_id
    cliente = db.query(Cliente).filter(Cliente.id == cliente_id).first() if cliente_id else None
    cliente_nombre = cliente.nombre_completo if cliente else "cliente sin registrar"

    lineas_por_id = {linea.id: linea for linea in (solicitud.lineas or [])}
    total_por_linea: dict[int, float] = defaultdict(float)
    total_por_lote: dict[int, float] = defaultdict(float)
    ots_por_id: dict[int, OrdenTrabajo] = {}

    for item in payload.lineas:
        linea = lineas_por_id.get(item.solicitud_linea_id)
        if not linea:
            raise HTTPException(status_code=404, detail=f"Línea {item.solicitud_linea_id} no pertenece a la solicitud")
        cantidad = float(item.cantidad_quintales or 0)
        if cantidad <= 0:
            raise HTTPException(status_code=400, detail="La cantidad a liquidar debe ser mayor a cero")
        ot = db.query(OrdenTrabajo).filter(OrdenTrabajo.id == item.ot_id, OrdenTrabajo.estado == "finalizada").first()
        if not ot:
            raise HTTPException(status_code=404, detail=f"Lote {item.ot_id} no encontrado o no aprobado")
        proceso_linea = (linea.proceso_preferido or "flexible").lower()
        if proceso_linea not in {"", "flexible"} and proceso_linea != (ot.proceso or "").lower():
            raise HTTPException(status_code=422, detail=f"El lote {ot.codigo_lote} no coincide con el proceso solicitado ({proceso_linea})")
        total_por_linea[linea.id] += cantidad
        total_por_lote[ot.id] += cantidad
        ots_por_id[ot.id] = ot

    for linea_id, cantidad_total in total_por_linea.items():
        linea = lineas_por_id[linea_id]
        pendiente = solicitud_linea_pendiente(linea)
        if cantidad_total > pendiente + 0.0001:
            raise HTTPException(status_code=422, detail=f"La línea {linea.descripcion} solo tiene {pendiente} qq pendientes")
        if cantidad_total + 0.0001 < pendiente:
            raise HTTPException(
                status_code=422,
                detail=(
                    f"No se puede liquidar parcialmente la línea {linea.descripcion}. "
                    f"Debe cubrir {pendiente} qq y se seleccionaron {round(cantidad_total, 3)} qq."
                ),
            )

    for ot_id, cantidad_total in total_por_lote.items():
        validate_lote_quantity_available(db, ots_por_id[ot_id], cantidad_total, "liquidación de solicitud")

    row = SolicitudSalidaVenta(
        codigo=next_code(db, SolicitudSalidaVenta, "codigo", "SAL-", 5),
        cotizacion_id=payload.cotizacion_id or None,
        cliente_id=cliente_id or None,
        estado="vendida",
        fecha=payload.fecha,
        moneda=(payload.moneda or solicitud.moneda or "USD").upper(),
        tipo_envio=payload.tipo_envio,
        direccion_entrega=payload.direccion_entrega,
        observaciones=payload.observaciones,
        created_by_id=current.id,
    )
    db.add(row)
    db.flush()

    subtotal_total = 0.0
    impuesto_total = 0.0
    total_total = 0.0

    for item in payload.lineas:
        linea = lineas_por_id.get(item.solicitud_linea_id)
        if not linea:
            raise HTTPException(status_code=404, detail=f"Línea {item.solicitud_linea_id} no pertenece a la solicitud")
        pendiente = solicitud_linea_pendiente(linea)
        cantidad = float(item.cantidad_quintales or 0)
        if cantidad > pendiente + 0.0001:
            raise HTTPException(status_code=422, detail=f"La línea {linea.descripcion} solo tiene {pendiente} qq pendientes")
        ot = db.query(OrdenTrabajo).filter(OrdenTrabajo.id == item.ot_id, OrdenTrabajo.estado == "finalizada").first()
        if not ot:
            raise HTTPException(status_code=404, detail=f"Lote {item.ot_id} no encontrado o no aprobado")
        proceso_linea = (linea.proceso_preferido or "flexible").lower()
        if proceso_linea not in {"", "flexible"} and proceso_linea != (ot.proceso or "").lower():
            raise HTTPException(status_code=422, detail=f"El lote {ot.codigo_lote} no coincide con el proceso solicitado ({proceso_linea})")
        validate_lote_quantity_available(db, ot, cantidad, "liquidación de solicitud")
        disponible_previo = lote_quintales_disponibles(db, ot.id)
        descripcion = item.descripcion or f"{linea.descripcion} · lote {ot.codigo_lote}"
        subtotal, impuesto, total = calc_line_totals(cantidad, item.precio_unitario, item.impuesto_porcentaje)
        subtotal_total += subtotal
        impuesto_total += impuesto
        total_total += total
        salida_linea = SolicitudSalidaVentaLinea(
            solicitud_id=row.id,
            ot_id=ot.id,
            descripcion=descripcion,
            cantidad_quintales=cantidad,
            precio_unitario=item.precio_unitario,
            impuesto_porcentaje=item.impuesto_porcentaje,
            subtotal=subtotal,
            impuesto=impuesto,
            total=total,
        )
        db.add(salida_linea)
        db.flush()
        db.add(SolicitudVentaLineaLiquidacion(
            solicitud_id=solicitud.id,
            linea_id=linea.id,
            salida_id=row.id,
            salida_linea_id=salida_linea.id,
            ot_id=ot.id,
            cantidad_quintales=cantidad,
            created_by_id=current.id,
        ))
        disponible_posterior = round(max(disponible_previo - cantidad, 0), 3)
        ot.estado = "vendido" if disponible_posterior <= 0.0001 else "finalizada"
        ot.updated_at = utcnow()
        db.add(ComentarioLote(
            ot_id=ot.id,
            tipo="comercial",
            comentario=(
                f"Se vendieron {cantidad} qq del lote al cliente {cliente_nombre} en la solicitud de salida {row.codigo}, "
                f"liquidando la solicitud de venta {solicitud.codigo}. "
                f"Disponible anterior: {disponible_previo} qq; disponible posterior: {disponible_posterior} qq."
            ),
            created_by_id=current.id,
        ))

    for ot_id in total_por_lote:
        sincronizar_estado_comercial_lote(db, ots_por_id.get(ot_id) or db.query(OrdenTrabajo).filter(OrdenTrabajo.id == ot_id).first())

    row.subtotal = round(subtotal_total, 2)
    row.impuesto = round(impuesto_total, 2)
    row.total = round(total_total, 2)
    db.flush()

    pendientes = []
    for linea in solicitud.lineas or []:
        liquidado = (
            db.query(func.coalesce(func.sum(SolicitudVentaLineaLiquidacion.cantidad_quintales), 0))
            .filter(SolicitudVentaLineaLiquidacion.linea_id == linea.id)
            .scalar()
            or 0
        )
        pendientes.append(round(max(float(linea.cantidad_quintales or 0) - float(liquidado or 0), 0), 3))
    solicitud.estado = "vendida" if all(p <= 0.0001 for p in pendientes) else "parcial"
    solicitud.updated_at = utcnow()
    db.add(SolicitudVentaSeguimiento(
        solicitud_id=solicitud.id,
        fecha=payload.fecha,
        canal="liquidacion",
        asunto=f"Liquidación en salida {row.codigo}",
        comentario=payload.observaciones or f"Se generó la solicitud de salida {row.codigo} desde la solicitud de venta.",
        created_by_id=current.id,
    ))
    db.commit()
    row = solicitud_salida_query(db).filter(SolicitudSalidaVenta.id == row.id).first()
    return serialize_solicitud_salida_venta(row)


@router.post("/solicitudes-venta/{solicitud_id}/seguimientos", status_code=201)
def create_solicitud_seguimiento(
    solicitud_id: int,
    payload: SolicitudVentaSeguimientoCreate,
    current: Usuario = Depends(require_roles("gerente", "administrativo")),
    db: Session = Depends(get_db),
):
    solicitud = db.query(SolicitudVenta).filter(SolicitudVenta.id == solicitud_id).first()
    if not solicitud:
        raise HTTPException(status_code=404, detail="Solicitud no encontrada")
    row = SolicitudVentaSeguimiento(
        solicitud_id=solicitud_id,
        fecha=payload.fecha,
        canal=(payload.canal or "nota").strip().lower(),
        asunto=payload.asunto,
        comentario=payload.comentario.strip(),
        created_by_id=current.id,
    )
    db.add(row)

    ot_ids: set[int] = set()
    if solicitud.ot_id:
        ot_ids.add(solicitud.ot_id)
    for linea in solicitud.lineas or []:
        for liq in linea.liquidaciones or []:
            if liq.ot_id:
                ot_ids.add(liq.ot_id)

    for ot_id in ot_ids:
        add_lote_evento(
            db,
            ot_id,
            "comercial",
            f"{current.nombre} registró seguimiento comercial en la solicitud {solicitud.codigo}: {row.canal} · {row.asunto or 'sin asunto'} · {row.comentario}",
            current,
        )

    db.commit()
    db.refresh(row)
    return serialize_solicitud_seguimiento(row)


@router.post("/solicitudes-venta/{solicitud_id}/documentos", status_code=201)
def upload_solicitud_documento(
    solicitud_id: int,
    request: Request,
    titulo: str = Form("Documento vinculado"),
    tipo: str = Form("documento"),
    descripcion: Optional[str] = Form(None),
    file: UploadFile = File(...),
    current: Usuario = Depends(require_roles("gerente", "administrativo")),
    db: Session = Depends(get_db),
):
    solicitud = db.query(SolicitudVenta).filter(SolicitudVenta.id == solicitud_id).first()
    if not solicitud:
        raise HTTPException(status_code=404, detail="Solicitud no encontrada")
    file_url, original, size, content_type = save_upload_file(file, upload_root() / "solicitudes_venta" / str(solicitud_id))
    if file_url.startswith("/uploads/"):
        file_url = f"{APP_BASE_URL}{file_url}"
    doc = SolicitudVentaDocumento(
        solicitud_id=solicitud_id,
        titulo=(titulo or original).strip() or original,
        tipo=(tipo or "documento").strip().lower(),
        descripcion=descripcion,
        file_url=file_url,
        file_name=original,
        content_type=content_type,
        size_bytes=size,
        created_by_id=current.id,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    return serialize_solicitud_documento(doc)


@router.post("/sync/push", response_model=SyncResult)
def sync_push(payload: SyncPush, current: Usuario = Depends(require_roles("gerente", "operario", "administrativo")), db: Session = Depends(get_db)):
    result = SyncResult()
    for item in payload.recibos:
        try:
            row = create_recibo_from_payload(db, item, current)
            db.commit()
            result.recibos.append(SyncResultItem(client_uuid=item.client_uuid, id=row.id, status="ok"))
        except Exception as exc:
            db.rollback()
            result.recibos.append(SyncResultItem(client_uuid=item.client_uuid, status="error", detail=str(exc)))
    for item in payload.ots:
        try:
            row = create_ot_from_payload(db, item, current)
            db.commit()
            result.ots.append(SyncResultItem(client_uuid=item.client_uuid, id=row.id, status="ok"))
        except Exception as exc:
            db.rollback()
            result.ots.append(SyncResultItem(client_uuid=item.client_uuid, status="error", detail=str(exc)))
    for raw in payload.seguimientos:
        try:
            ot_id = int(raw.get("ot_id") or raw.get("orden_trabajo_id") or 0)
            if not ot_id:
                raise ValueError("Falta ot_id")
            body = SeguimientoCreate(**{k: v for k, v in raw.items() if k not in {"ot_id", "orden_trabajo_id"}})
            ot = db.query(OrdenTrabajo).filter(OrdenTrabajo.id == ot_id).first()
            if not ot:
                raise ValueError("OT no encontrada")
            seg = SeguimientoOT(
                orden_trabajo_id=ot.id,
                client_uuid=body.client_uuid,
                fecha=body.fecha,
                actividad_realizada=body.actividad_realizada,
                horas_implementadas=body.horas_implementadas,
                temperatura=body.temperatura,
                humedad=body.humedad,
                comentario=body.comentario,
                created_by_id=current.id,
            )
            db.add(seg)
            if ot.estado == "abierta":
                ot.estado = "en_proceso"
            db.commit()
            result.seguimientos.append(SyncResultItem(client_uuid=body.client_uuid, id=seg.id, status="ok"))
        except Exception as exc:
            db.rollback()
            result.seguimientos.append(SyncResultItem(client_uuid=raw.get("client_uuid"), status="error", detail=str(exc)))
    return result

@router.get("/public/lotes/{qr_token}", response_model=PublicLoteOut)
def get_public_lote(
    qr_token: str,
    db: Session = Depends(get_db),
):
    row = (
        ot_query(db)
        .filter(OrdenTrabajo.qr_token == qr_token)
        .first()
    )

    if not row:
        raise HTTPException(status_code=404, detail="Lote no encontrado")

    return serialize_public_lote(row)

@router.get("/public/lotes/{qr_token}/qr.svg")
def get_public_lote_qr_svg(
    qr_token: str,
    db: Session = Depends(get_db),
):
    row = db.query(OrdenTrabajo).filter(OrdenTrabajo.qr_token == qr_token).first()

    if not row:
        raise HTTPException(status_code=404, detail="Lote no encontrado")

    url = row.qr_public_url or lote_public_url(row.qr_token)
    svg = make_qr_svg(url)

    return Response(content=svg, media_type="image/svg+xml")
